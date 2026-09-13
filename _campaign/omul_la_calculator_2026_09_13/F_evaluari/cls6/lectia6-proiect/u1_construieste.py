"""U1 - fac proiectul lectiei lectia6-proiect ca un elev de a VI-a, EXACT cum cere lectia (Ex.1 plan + Ex.2 construiesc + Ex.3 cuprins cu linkuri).
python-pptx pentru continut; lxml pentru <p:transition> (pune_tranzitie, copiat din lectia5-tranzitii -> _l5_u1_construieste.py),
<p:timing> (clasa Timing din lectia4-animatii -> _l4_u1_construieste.py) si <a:hlinkClick action="ppaction://hlinksldjump"> (link intern pe TEXT,
cum cere Ex.3: „selecteaza textul, Insert -> Link -> Place in This Document”).

Ce cere lectia si ce pun (sursa = innerText.txt):
  Ex.1 (r. 774-780): subiect + >=5 idei + continut pe diapozitive + >=3 imagini cu sursa notata  -> plan_ex1.txt
  Structura atomului 1 (r. 82-92): 1 Titlu (Subiect+Autor+Clasa/Data), 2 Cuprins, 3-5 Continut (titlu+continut+imagine), 6 Concluzie, 7 Multumiri+Bibliografie
  Ex.2 (r. 804-808): >=6 diapozitive, tema consecventa, >=3 formatari (bold, lista numerotata, text colorat), >=1 tranzitie + 1 animatie
  Lista de verificare atom 3 (r. 257-271): titlu cu autor si data, >=6 diap., >=3 imagini, animatie Entrance, tranzitie uniforma, font >=24pt, 6x6, final + .pptx
  Ex.3 (r. 834): fiecare titlu din cuprins = hyperlink intern
Imaginile: DESENE PROPRII facute cu PIL (sursa permisa fara discutie: „Imagini proprii”, lectia3 r. 402). Nu pun poze cu persoane.
Tema: sablonul implicit python-pptx („Office Theme”) - temele „Ion”/„Facet” nu exista fara PowerPoint -> pas interfata.
Produce: produs_elev/Proiect_Albinele.pptx, produs_elev/img_*.png, plan_ex1.txt; iesire u1_iesire.txt
"""
import copy
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Pt

L = Path(__file__).resolve().parent
_src = (L / "_l4_u1_construieste.py").read_text(encoding="utf-8")
exec(_src.split("prs = Presentation()")[0].split('"""', 2)[2])  # Timing, pune_timing, caseta, P, NS_P
_s5 = (L / "_l5_u1_construieste.py").read_text(encoding="utf-8")
exec("MC = %r\nP14 = %r\nP159 = %r\n" % ("http://schemas.openxmlformats.org/markup-compatibility/2006",
     "http://schemas.microsoft.com/office/powerpoint/2010/main", "http://schemas.microsoft.com/office/powerpoint/2015/09/main")
     + "def efect_xml" + _s5.split("def efect_xml", 1)[1].split("def deck", 1)[0])  # efect_xml, pune_tranzitie

OUT = L / "produs_elev"
OUT.mkdir(exist_ok=True)
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
A = "{%s}" % NS_A
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
log = []

# ---------- Ex.1: planul (pe hartie la clasa; aici fisier text) ----------
PLAN = {
    "subiect": "Albinele - mici, dar importante",
    "idei": ["Cum arată o albină", "Viața în stup", "Cum fac mierea", "De ce avem nevoie de albine", "Cum le putem ajuta"],
    "diapozitive": [
        ("Albinele - mici, dar importante", ["Autor: Elev, clasa a VI-a", "13.11.2026"]),
        ("Cuprins", ["Cum arată o albină", "Viața în stup", "De ce avem nevoie de albine"]),
        ("Cum arată o albină", ["Corp în trei părți", "Șase picioare, patru aripi", "Dungi galbene și negre"]),
        ("Viața în stup", ["O regină, multe lucrătoare", "Lucrătoarele adună nectar", "Mierea se păstrează în faguri"]),
        ("De ce avem nevoie de albine", ["Polenizează florile", "Ajută fructele să crească", "Ne dau miere și ceară"]),
        ("Concluzie", ["Fără albine, mai puține fructe", "Plantează flori pentru albine"]),
        ("Mulțumesc pentru atenție!", ["Bibliografie:", "Manualul de biologie, clasa a V-a", "Imagini: desene proprii"]),
    ],
    "imagini": [("img_albina.png", "desen propriu"), ("img_stup.png", "desen propriu"), ("img_floare.png", "desen propriu")],
}
with open(L / "plan_ex1.txt", "w", encoding="utf-8") as f:
    f.write(f"Subiect: {PLAN['subiect']}\nIdei ({len(PLAN['idei'])}): " + "; ".join(PLAN["idei"]) + "\n")
    for i, (t, r) in enumerate(PLAN["diapozitive"], 1):
        f.write(f"Diap. {i}: {t} | " + " / ".join(r) + "\n")
    f.write("Imagini: " + "; ".join(f"{n} (sursa: {s})" for n, s in PLAN["imagini"]) + "\n")
log.append(f"Ex.1 plan: {len(PLAN['idei'])} idei, {len(PLAN['diapozitive'])} diapozitive, {len(PLAN['imagini'])} imagini cu sursa")


# ---------- imagini proprii (desene) ----------
def desen(nume, fn):
    im = Image.new("RGB", (600, 600), "white")
    fn(ImageDraw.Draw(im))
    im.save(OUT / nume)


desen("img_albina.png", lambda d: ([d.ellipse((150, 220, 450, 400), fill=(245, 190, 20), outline="black", width=6)]
                                   + [d.rectangle((210 + k * 70, 225, 240 + k * 70, 395), fill="black") for k in range(3)]
                                   + [d.ellipse((230, 90, 330, 230), outline=(80, 80, 200), width=6), d.ellipse((300, 90, 400, 230), outline=(80, 80, 200), width=6)]))
desen("img_stup.png", lambda d: [d.regular_polygon((150 + (k % 3) * 150, 180 + (k // 3) * 150 + (75 if k % 2 else 0), 80), 6, fill=(240, 180, 40), outline="black")
                                 for k in range(6)])
desen("img_floare.png", lambda d: ([d.ellipse((300 + dx - 70, 250 + dy - 70, 300 + dx + 70, 250 + dy + 70), fill=(230, 80, 140))
                                    for dx, dy in ((0, -110), (0, 110), (-110, 0), (110, 0))]
                                   + [d.ellipse((230, 180, 370, 320), fill=(250, 210, 0)), d.rectangle((290, 360, 310, 590), fill="green")]))

# ---------- Ex.2: construiesc ----------
prs = Presentation()              # tema implicita „Office Theme” (o singura tema -> consecventa)
LAY_TITLU, LAY_CONT = prs.slide_layouts[0], prs.slide_layouts[1]
ALBASTRU = RGBColor(0x1F, 0x4E, 0x9A)
sl = []


def titlu(s, text, size=40):
    s.shapes.title.text = text
    r = s.shapes.title.text_frame.paragraphs[0].runs[0]
    r.font.bold = True            # formatare 1: titluri bold (Ctrl+B)
    r.font.size = Pt(size)
    return s.shapes.title


def lista_numerotata(ph, randuri, size=28, cuvant_colorat=None):
    tf = ph.text_frame
    for i, t in enumerate(randuri):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        for old in pPr.findall(A + "buNone") + pPr.findall(A + "buChar"):
            pPr.remove(old)
        etree.SubElement(pPr, A + "buAutoNum", type="arabicPeriod")   # formatare 2: lista numerotata (butonul 1,2,3 din Home)
        if cuvant_colorat and cuvant_colorat in t:
            a, b = t.split(cuvant_colorat, 1)
            for bucata, culoare in ((a, None), (cuvant_colorat, ALBASTRU), (b, None)):
                if bucata:
                    r = p.add_run()
                    r.text = bucata
                    r.font.size = Pt(size)
                    if culoare:
                        r.font.color.rgb = culoare   # formatare 3: Font Color pe un cuvant-cheie
                        r.font.bold = True
        else:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)


D = PLAN["diapozitive"]
# 1 Titlu
s = prs.slides.add_slide(LAY_TITLU)
tit1 = titlu(s, D[0][0], 44)
s.placeholders[1].text = "\n".join(D[0][1])
for p in s.placeholders[1].text_frame.paragraphs:
    p.runs[0].font.size = Pt(28)
sl.append(s)
# 2 Cuprins
s = prs.slides.add_slide(LAY_CONT)
titlu(s, D[1][0])
lista_numerotata(s.placeholders[1], D[1][1])
sl.append(s)
# 3-5 continut cu imagine
for k, (t, randuri) in enumerate(D[2:5]):
    s = prs.slides.add_slide(LAY_CONT)
    titlu(s, t, 36)
    ph = s.placeholders[1]
    ph.width = Emu(int(prs.slide_width * 0.55))
    lista_numerotata(ph, randuri, cuvant_colorat={0: "trei", 1: "regină", 2: "Polenizează"}[k])
    img = OUT / PLAN["imagini"][k][0]
    s.shapes.add_picture(str(img), Emu(int(prs.slide_width * 0.62)), Emu(int(prs.slide_height * 0.32)), width=Emu(int(prs.slide_width * 0.33)))
    sl.append(s)
# 6 Concluzie, 7 Multumiri + bibliografie
for t, randuri in D[5:7]:
    s = prs.slides.add_slide(LAY_CONT)
    titlu(s, t, 36)
    tf = s.placeholders[1].text_frame
    for i, r in enumerate(randuri):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = r
        p.runs[0].font.size = Pt(28 if i == 0 or t == "Concluzie" else 24)
    sl.append(s)

# tranzitie: Fade pe toate (Apply To All = aceeasi tranzitie scrisa pe fiecare diapozitiv)
for s in sl:
    pune_tranzitie(s, "fade", 700)
# animatie: titlul diapozitivului 1, Fade (Entrance), La clic
t = Timing()
t.on_click(t.efect(tit1.shape_id, "entr", 10, 0, 500, filt="fade"))
pune_timing(sl[0], t)
log.append("Ex.2: 7 diapozitive (layout Title Slide + Title and Content), titluri bold, liste numerotate, 1 cuvant colorat/diap. 3-5, 3 imagini, Fade pe toate, Fade Entrance pe titlul diap.1")

# ---------- Ex.3: cuprins interactiv (link intern pe text) ----------
cup = sl[1].placeholders[1].text_frame
for i, p in enumerate(cup.paragraphs):
    tinta = sl[2 + i]
    rid = sl[1].part.relate_to(tinta.part, RT.SLIDE)
    for r in p.runs:
        rPr = r._r.get_or_add_rPr()
        etree.SubElement(rPr, A + "hlinkClick", {"{%s}id" % NS_R: rid, "action": "ppaction://hlinksldjump"})
log.append("Ex.3: 3 linkuri interne pe textul cuprinsului -> diap. 3, 4, 5")

dest = OUT / "Proiect_Albinele.pptx"
prs.save(str(dest))                # „fisierul salvat ca .pptx” (lista de verificare, atomul 3)
log.append(f"salvat {dest.name}: {dest.stat().st_size} octeti, {len(prs.slides)} diapozitive")
(L / "u1_iesire.txt").write_text("\n".join(log), encoding="utf-8")
print("\n".join(log))
