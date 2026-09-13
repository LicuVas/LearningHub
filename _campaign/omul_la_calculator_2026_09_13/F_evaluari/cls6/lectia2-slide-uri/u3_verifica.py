"""U3 + U7: a doua cale pe ce am construit si pe ce afirma lectia. Proces Python NOU fata de u1_construieste.py.
Scrie u3_iesire.json si 07_redeschis.json; tipareste un rezumat scurt."""
import json
import re
from pathlib import Path

import fitz
from pptx import Presentation

L = Path(__file__).resolve().parent
T = (L / "innerText.txt").read_text(encoding="utf-8")
LINII = T.splitlines()
REL1 = L.parent / "lectia1-powerpoint-intro" / "innerText.txt"
rez = {}

# --- U7: redeschid fisierele ---
red = {}
for f in sorted((L / "produs_elev").glob("*.pptx")):
    p = Presentation(str(f))
    red[f.name] = {
        "diapozitive": len(p.slides),
        "layouturi": [s.slide_layout.name for s in p.slides],
        "titluri": [s.shapes.title.text if s.shapes.title is not None else "" for s in p.slides],
        "texte_pe_master": [sh.text_frame.text for sh in p.slide_master.shapes if sh.has_text_frame and not sh.is_placeholder],
        "fundal_slide1_xml": ("gradFill" if b"gradFill" in __import__("lxml.etree", fromlist=["x"]).tostring(p.slides[0]._element) else "altul"),
    }
(L / "07_redeschis.json").write_text(json.dumps(red, ensure_ascii=False, indent=1), encoding="utf-8")

# --- a doua cale 1: PDF randat de LibreOffice (proces separat) vs structura pptx ---
pdf = {}
for f in sorted((L / "randat_pptx").glob("*.pdf")):
    d = fitz.open(str(f))
    pag = [d[i].get_text() for i in range(len(d))]
    pdf[f.stem] = {"pagini": len(d),
                   "VG_pe_pagini": [i + 1 for i, t in enumerate(pag) if "VG" in t],
                   "nume_master_pe_pagini": [i + 1 for i, t in enumerate(pag) if "Nume Prenume" in t and "Realizat" not in t.split("Nume Prenume")[0][-12:]],
                   "numere_pagina_gasite": [re.findall(r"^\s*(\d{1,2})\s*$", t, re.M) for t in pag]}
rez["pdf"] = pdf

# --- a doua cale 2: contrastul culorilor din gradientul atomului 7 (WCAG, formula W3C) ---
def lum(h):
    c = [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    c = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]

def cr(a, b):
    la, lb = sorted((lum(a), lum(b)), reverse=True)
    return round((la + 0.05) / (lb + 0.05), 2)

rez["gradient"] = {"60a5fa_vs_3b82f6": cr("60a5fa", "3b82f6"),
                   "text_negru_pe_60a5fa": cr("000000", "60a5fa"), "text_negru_pe_3b82f6": cr("000000", "3b82f6"),
                   "text_alb_pe_60a5fa": cr("ffffff", "60a5fa"), "text_alb_pe_3b82f6": cr("ffffff", "3b82f6")}

# --- a doua cale 3: intrebarea din atom vs atomul care preda termenul ---
def linie(pat, start=0):
    for i, l in enumerate(LINII[start:], start):
        if re.search(pat, l):
            return i + 1
    return None

atomi = {n: linie(rf"^{n}\. Continut") for n in range(1, 8)}
def atom_la(lin):
    k = [n for n, a in atomi.items() if a and a <= lin]
    return max(k) if k else 0

intrebari = [("Care este scurtatura de tastatura pentru adaugarea unui slide nou?", r"Ctrl \+ M - Cel mai rapid"),
             ("Ce este Slide Master?", r"^Definitie: Slide Master este"),
             ("Ce tip de fundal foloseste o trecere progresiva", r"^2\. Gradient Fill"),
             ("Unde gasesti si aplici teme predefinite?", r"^Mergi la tab-ul Design"),
             ("Ce layout folosesti pentru primul slide", r"^Title Slide$"),
             ("Cum schimbi fundalul unui slide?", r"^Metoda 1: Click dreapta pe slide → Format Background"),
             ("Cum stergi rapid un slide?", r"^Metoda 1: Selecteaza slide-ul in panoul din stanga → Apasa Delete"),
             ("Ce layout este ideal pentru o comparatie", r"^Comparison$"),
             ("Cum duplici rapid un slide?", r"^Metoda 2: Selecteaza slide → Ctrl \+ D"),
             ("Cum reordonezi slide-urile cel mai usor?", r"^Metoda 1 - Drag and Drop")]
tab = []
for q, pred in intrebari:
    lq = linie(re.escape(q))
    lp = linie(pred, 110)
    tab.append({"intrebare": q, "linia": lq, "in_atomul": atom_la(lq), "predat_la_linia": lp, "predat_in_atomul": atom_la(lp),
                "inainte_de_predare": atom_la(lq) < atom_la(lp)})
rez["intrebari_vs_predare"] = tab
rez["intrebari_inainte_de_predare"] = sum(1 for x in tab if x["inainte_de_predare"])
rez["intrebari_despre_atomul_propriu"] = sum(1 for x in tab if x["in_atomul"] == x["predat_in_atomul"])

# --- a doua cale 4: Slide Master = optional / nivel standard ---
rez["slide_master"] = {
    "nota_optional": [l.strip() for l in LINII if "Nu este evaluata la nivel minim sau standard" in l][:1],
    "ex2_titlu": [LINII[i + 4].strip() for i, l in enumerate(LINII) if l.startswith("Exercitiul 2 (Nivel standard)")][:1],
    "ex4_cere_master": any("Slide Master cu logo/nume in colt" in l for l in LINII),
    "aparitii_Slide_Master": len(re.findall(r"Slide Master", T)),
}

# --- a doua cale 5: ce a predat deja lectia 1 (repetitie cu ora anterioara) ---
if REL1.is_file():
    t1 = REL1.read_text(encoding="utf-8")
    rez["deja_in_lectia1"] = {k: len(re.findall(re.escape(k), t1)) for k in ["Ctrl+M", "Ctrl + M", "Duplicate Slide", "Delete Slide", "Slide Sorter", "Layout", "Design"]}
    rez["in_lectia2"] = {k: len(re.findall(re.escape(k), T)) for k in ["Ctrl+M", "Ctrl + M", "Duplicate Slide", "Delete Slide", "Slide Sorter", "Layout", "Design"]}

# --- a doua cale 6: termeni englezi de interfata vs diacritice ---
ro = ["ă", "â", "î", "ș", "ț", "ş", "ţ"]
rez["cuvinte_cu_diacritice"] = sorted({w for w in re.findall(r"\w+", T) if any(c in w.lower() for c in ro)})
rez["termeni_en"] = {k: len(re.findall(re.escape(k), T)) for k in ["Slide Master", "Format Background", "Gradient", "Layout", "Title Slide", "Two Content", "Comparison", "Design", "Home", "View", "Insert", "Close Master View", "Header & Footer"]}
rez["indiciu2_slide_sorter_panou"] = [l.strip() for l in LINII if l.strip().startswith("In panoul din stanga (Slide Sorter)")]

(L / "u3_iesire.json").write_text(json.dumps(rez, ensure_ascii=False, indent=1), encoding="utf-8")
print("redeschise:", {k: v["diapozitive"] for k, v in red.items()})
print("pdf:", {k: (v["pagini"], v["VG_pe_pagini"]) for k, v in pdf.items()})
print("gradient:", rez["gradient"])
print("intrebari inainte de predare:", rez["intrebari_inainte_de_predare"], "| despre atomul propriu:", rez["intrebari_despre_atomul_propriu"])
for x in tab:
    print("  ", x["in_atomul"], "<-", x["predat_in_atomul"], x["intrebare"][:50])
print("slide_master:", rez["slide_master"])
print("lectia1 vs 2:", rez.get("deja_in_lectia1"), rez.get("in_lectia2"))
print("diacritice:", rez["cuvinte_cu_diacritice"])
