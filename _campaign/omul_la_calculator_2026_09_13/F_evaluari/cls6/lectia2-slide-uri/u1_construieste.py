"""U1 - fac sarcinile lectiei lectia2-slide-uri ca un elev de a VI-a, cu python-pptx + lxml.

Produce in produs_elev/:
  Vacanta_Ideala.pptx      (Incearca tu: 5 slide-uri noi, sterg 1, schimb ordinea a 2, layout diferit, fundal ultimul slide)
  Pasiunile_Mele.pptx      (Ex.1: 7 slide-uri cu layout-urile cerute, slide 3 = duplicat al slide-ului 2,
                            fundal gradient pe slide 1 cu culorile din atomul 7)
  Slide_Master_Ex2.pptx    (Ex.2: initiala in coltul dreapta sus pusa PE MASTER + numar de slide; 3 slide-uri de proba)
  Ghid_Turistic.pptx       (Ex.4: 10 slide-uri, >=5 layout-uri, 3 tipuri de fundal, 1 duplicat, master cu nume)
Iesire: u1_iesire.txt
"""
import copy
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

L = Path(__file__).resolve().parent
OUT = L / "produs_elev"
IMG = OUT / "imagini"
OUT.mkdir(exist_ok=True)
IMG.mkdir(exist_ok=True)
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
P = "{%s}" % NS["p"]
A = "{%s}" % NS["a"]
log = []

LAY = {"Title Slide": 0, "Title and Content": 1, "Two Content": 3, "Comparison": 4,
       "Title Only": 5, "Blank": 6, "Picture with Caption": 8}


def imagine(nume, culoare, text):
    p = IMG / nume
    im = Image.new("RGB", (800, 600), culoare)
    ImageDraw.Draw(im).text((40, 280), text, fill="white")
    im.save(p)
    return p


def lay(prs, nume):
    l = prs.slide_layouts[LAY[nume]]
    assert l.name == nume, (l.name, nume)
    return l


def pune_text(s, titlu, corpuri=()):
    if s.shapes.title is not None:
        s.shapes.title.text = titlu
    ph = [x for x in s.placeholders if x.placeholder_format.idx != 0]
    for x, t in zip(ph, corpuri):
        if t is None:
            continue
        if isinstance(t, Path):
            x.insert_picture(str(t))
        else:
            x.text = t


def muta(prs, de_la, la):
    """reordonare = drag and drop in panoul din stanga (pozitii 0-based)"""
    lst = prs.slides._sldIdLst
    el = list(lst)[de_la]
    lst.remove(el)
    lst.insert(la, el)


def dubleaza(prs, idx):
    """Duplicate Slide: copie imediat DUPA original (asa scrie Microsoft: 'Dublarea este inserata imediat dupa original')"""
    src = prs.slides[idx]
    nou = prs.slides.add_slide(src.slide_layout)
    for sh in list(nou.shapes):
        sh._element.getparent().remove(sh._element)
    for sh in src.shapes:
        el = copy.deepcopy(sh._element)
        nou.shapes._spTree.append(el)
        # imaginile: relatia trebuie copiata
        for blip in el.iter(A + "blip"):
            rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            part = src.part.related_part(rid)
            new_rid = nou.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image")
            blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", new_rid)
    muta(prs, len(prs.slides) - 1, idx + 1)
    return prs.slides[idx + 1]


def fundal_solid(slide, hexcol):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(hexcol)


def fundal_gradient(slide, c1, c2, unghi_grade=45):
    """Format Background -> Gradient fill, Linear Diagonal - Top Left to Bottom Right, 2 stops"""
    cSld = slide._element.find(P + "cSld")
    for old in cSld.findall(P + "bg"):
        cSld.remove(old)
    bg = etree.SubElement(cSld, P + "bg")
    cSld.remove(bg)
    cSld.insert(0, bg)
    bgPr = etree.SubElement(bg, P + "bgPr")
    g = etree.SubElement(bgPr, A + "gradFill", rotWithShape="1")
    gs = etree.SubElement(g, A + "gsLst")
    for pos, c in ((0, c1), (100000, c2)):
        e = etree.SubElement(gs, A + "gs", pos=str(pos))
        etree.SubElement(e, A + "srgbClr", val=c)
    etree.SubElement(g, A + "lin", ang=str(unghi_grade * 60000), scaled="0")
    etree.SubElement(bgPr, A + "effectLst")


def fundal_imagine(slide, prs, img):
    """Picture or texture fill: imagine pe tot slide-ul, ca fundal (emulat: blipFill in p:bg)"""
    rid = slide.part.get_or_add_image_part(str(img))[1]
    cSld = slide._element.find(P + "cSld")
    bg = etree.Element(P + "bg")
    cSld.insert(0, bg)
    bgPr = etree.SubElement(bg, P + "bgPr")
    bf = etree.SubElement(bgPr, A + "blipFill", dpi="0", rotWithShape="1")
    etree.SubElement(bf, A + "blip", {"{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed": rid})
    st = etree.SubElement(bf, A + "stretch")
    etree.SubElement(st, A + "fillRect")
    etree.SubElement(bgPr, A + "effectLst")


def textbox_pe_master(prs, text, colt="dreapta_sus"):
    """Ex.2: View -> Slide Master -> primul slide (cel mare) -> Insert -> Text Box.
    python-pptx nu are add_textbox pe master: fac caseta pe un slide de lucru si o mut in spTree-ul masterului."""
    tmp = prs.slides.add_slide(prs.slide_layouts[6])
    w, h = Emu(1400000), Emu(500000)
    left = prs.slide_width - w - Emu(200000) if "dreapta" in colt else Emu(200000)
    top = Emu(150000) if "sus" in colt else prs.slide_height - h - Emu(150000)
    tb = tmp.shapes.add_textbox(left, top, w, h)
    tb.text_frame.text = text
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    tb.text_frame.paragraphs[0].runs[0].font.bold = True
    el = tb._element
    prs.slide_master.shapes._spTree.append(el)
    # sterg slide-ul de lucru
    lst = prs.slides._sldIdLst
    last = list(lst)[-1]
    prs.part.drop_rel(last.rId)
    lst.remove(last)


def numar_slide(slide, prs):
    """Insert -> Header & Footer -> Slide number -> Apply to All: in fisier inseamna un placeholder sldNum cu campul slidenum
    pe fiecare slide. Emulat cu o caseta care contine campul <a:fld type='slidenum'> (se recalculeaza la randare)."""
    tb = slide.shapes.add_textbox(prs.slide_width - Emu(1300000), prs.slide_height - Emu(600000), Emu(1100000), Emu(400000))
    p = tb.text_frame.paragraphs[0]._p
    fld = etree.SubElement(p, A + "fld", id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
    etree.SubElement(fld, A + "rPr", lang="ro-RO")
    t = etree.SubElement(fld, A + "t")
    t.text = "‹#›"


def subsol(slide, prs, text):
    tb = slide.shapes.add_textbox(Emu(3000000), prs.slide_height - Emu(600000), Emu(3500000), Emu(400000))
    tb.text_frame.text = text


def rezumat(prs):
    return [(i + 1, s.slide_layout.name, (s.shapes.title.text if s.shapes.title is not None else "")) for i, s in enumerate(prs.slides)]


# ---------------- Incearca tu: Vacanta Ideala ----------------
prs = Presentation()
for nume, t in [("Title Slide", "Vacanta Ideala"), ("Title and Content", "Unde merg"), ("Two Content", "Munte sau mare"),
                ("Title Only", "Ce iau cu mine"), ("Blank", ""), ("Comparison", "Inainte / dupa")]:
    s = prs.slides.add_slide(lay(prs, nume))
    pune_text(s, t)
log.append(f"Incearca tu: am adaugat {len(prs.slides)} slide-uri (1 de titlu + 5 noi), layout-uri: {[s.slide_layout.name for s in prs.slides]}")
# sterg unul (Blank, pozitia 5)
lst = prs.slides._sldIdLst
el = list(lst)[4]
prs.part.drop_rel(el.rId)
lst.remove(el)
muta(prs, 3, 1)  # schimb ordinea a doua slide-uri
fundal_solid(prs.slides[-1], "FFD166")
prs.save(OUT / "Vacanta_Ideala.pptx")
log.append(f"Incearca tu dupa stergere + mutare: {rezumat(prs)}; fundal galben pe ultimul. "
           "Tema predefinita: NU se poate aplica din python-pptx -> pas interfata (Proiectare/Design).")
log.append("Incearca tu: lista din atomul 1 are 4 metode de adaugare (New Slide, Ctrl+M, click dreapta -> New Slide, Duplicate Slide); "
           "sarcina cere 5 slide-uri 'cate o metoda diferita pentru fiecare' -> a 5-a metoda nu exista in lectie pana la acel moment "
           "(Copy/Paste apare abia la atomul 3).")

# ---------------- Ex.1: Pasiunile Mele ----------------
i1 = imagine("pasiune1_desen.png", (37, 99, 235), "Pasiunea 1: desenul")
i2 = imagine("pasiune2_fotbal.png", (22, 163, 74), "Pasiunea 2: fotbalul")
i5 = imagine("poza_preferata.png", (180, 83, 9), "Poza preferata (fara chip)")
prs = Presentation()
s1 = prs.slides.add_slide(lay(prs, "Title Slide")); pune_text(s1, "Pasiunile Mele", ["Nume Prenume"])
s2 = prs.slides.add_slide(lay(prs, "Title and Content")); pune_text(s2, "Pasiunea #1", ["Desenez in fiecare zi."])
s2.shapes.add_picture(str(i1), Emu(5200000), Emu(2600000), width=Emu(3200000))
s4 = prs.slides.add_slide(lay(prs, "Two Content")); pune_text(s4, "De ce imi plac?", ["Ma relaxeaza", "Ma misc"])
s5 = prs.slides.add_slide(lay(prs, "Picture with Caption"))
ph5 = {x.placeholder_format.type: x for x in s5.placeholders}
s5.shapes.title.text = "Poza mea preferata"
pic_ph = [x for x in s5.placeholders if "Picture" in x.name]
pic_ph[0].insert_picture(str(i5))
s6 = prs.slides.add_slide(lay(prs, "Comparison"))
pune_text(s6, "Inainte si dupa", ["Inainte de pasiune", "Stateam mult la telefon", "Dupa ce am descoperit pasiunea", "Desenez si ies afara"])
s7 = prs.slides.add_slide(lay(prs, "Title Only")); pune_text(s7, "Multumesc pentru atentie!")
s3 = dubleaza(prs, 1)
s3.shapes.title.text = "Pasiunea #2"
body = [x for x in s3.placeholders if x.placeholder_format.idx != 0][0]
body.text = "Joc fotbal cu prietenii."
fundal_gradient(prs.slides[0], "60A5FA", "3B82F6", 45)
prs.save(OUT / "Pasiunile_Mele.pptx")
log.append(f"Ex.1: {rezumat(prs)}")
log.append("Ex.1: slide 3 = duplicat al slide-ului 2 (copie inserata DUPA original), apoi titlul/textul schimbate; imaginea slide-ului 3 "
           "a ramas cea de la slide 2 -> elevul trebuie sa stie sa inlocuiasca imaginea (click dreapta -> Change Picture), neexplicat in lectie.")
log.append("Ex.1: Comparison are 4 casete de continut + titlu; Picture with Caption are caseta de imagine + titlu + text.")

# ---------------- Ex.2: Slide Master ----------------
prs = Presentation()
textbox_pe_master(prs, "VG", "dreapta_sus")
prs.slide_master.placeholders  # doar ca sa verific ca exista
for t, n in [("Title Slide", "Test 1"), ("Title and Content", "Test 2"), ("Two Content", "Test 3")]:
    s = prs.slides.add_slide(lay(prs, t)); pune_text(s, n)
    numar_slide(s, prs)
    subsol(s, prs, "Prezentare Nume Prenume - 2026")
# fontul titlurilor pe master: schimb fontul in titleStyle
ts = prs.slide_master._element.find(P + "txStyles/" + P + "titleStyle")
lvl = ts.find(A + "lvl1pPr")
rpr = lvl.find(A + "defRPr")
lat = rpr.find(A + "latin")
if lat is None:
    lat = etree.SubElement(rpr, A + "latin")
lat.set("typeface", "Georgia")
prs.save(OUT / "Slide_Master_Ex2.pptx")
log.append("Ex.2: caseta 'VG' pusa in spTree-ul masterului (nu pe slide-uri); font Georgia in titleStyle; "
           "numar + subsol = Header & Footer -> Apply to All (emulat pe fiecare slide, pas interfata). Randarea arata daca 'VG' apare pe toate.")

# ---------------- Ex.4: Ghid turistic ----------------
prs = Presentation()
textbox_pe_master(prs, "Nume Prenume", "dreapta_sus")
continut = [("Title Slide", "Descopera Piatra Neamt", ["Ghid turistic"]),
            ("Picture with Caption", "Harta", [imagine("harta.png", (71, 85, 105), "Harta"), None, "Unde se afla orasul"]),
            ("Title and Content", "Istoric pe scurt", ["Oras vechi, pe Bistrita"]),
            ("Title and Content", "Obiectiv 1: Cetatuia", ["Telegondola"]),
            ("Title and Content", "Obiectiv 3: Muzeul Cucuteni", ["Ceramica veche"]),
            ("Two Content", "Mancare traditionala", ["Placinte", "Branza"]),
            ("Title and Content", "Evenimente", ["Festivaluri de vara"]),
            ("Comparison", "Informatii practice", ["Transport", "Autobuz, tren", "Cazare", "Pensiuni"]),
            ("Title Only", "De ce sa vizitezi?", [])]
for t, n, c in continut:
    s = prs.slides.add_slide(lay(prs, t)); pune_text(s, n, c)
d = dubleaza(prs, 3)
d.shapes.title.text = "Obiectiv 2: Curtea Domneasca"
for s in prs.slides:
    numar_slide(s, prs)
    subsol(s, prs, "Realizat de Nume Prenume, clasa a VI-a")
fundal_solid(prs.slides[2], "E0F2FE")
fundal_gradient(prs.slides[0], "60A5FA", "3B82F6", 45)
fundal_imagine(prs.slides[9], prs, imagine("fundal_munte.png", (30, 64, 100), "fundal"))
prs.save(OUT / "Ghid_Turistic.pptx")
lays = sorted({s.slide_layout.name for s in prs.slides})
log.append(f"Ex.4: {len(prs.slides)} slide-uri, {len(lays)} layout-uri diferite {lays}; {rezumat(prs)}")
log.append("Ex.4: fundaluri: gradient (1), solid (3), imagine (10); Slide Master cu nume; numar + subsol pe toate. "
           "Tema: pas interfata. 'Top 3 obiective, cate un slide pentru fiecare' + 9 continuturi sugerate = 10 slide-uri exact.")

(L / "u1_iesire.txt").write_text("\n".join(log), encoding="utf-8")
print("\n".join(x[:200] for x in log[:12]))
