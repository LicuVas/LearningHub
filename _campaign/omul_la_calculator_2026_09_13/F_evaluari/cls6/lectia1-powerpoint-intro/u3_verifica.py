"""U3 + U7 + U12 - verificari pe a doua cale, redeschiderea artefactelor si sensibilitatea timpului.
Scrie: u3_iesire.json, 07_redeschis.json, u12_sensibilitate.txt"""
import html
import json
import re
import zipfile
from pathlib import Path

import fitz  # PyMuPDF
from pptx import Presentation

L = Path(__file__).resolve().parent
LESSON = Path(r"C:/00/Projects/LearningHub/content/tic/cls6/m1-prezentari/lectia1-powerpoint-intro.html")
src = LESSON.read_text(encoding="utf-8")
rez = {}

# ---- 1. Intrebarea din atomul N - unde e predat termenul? (cheia vs textul atomilor) ----
poz = [m.start() for m in re.finditer(r'<div class="atom" id="atom-\d"', src)]
fin = src.find('id="lesson-summary"')
start_atomi = poz[0]


def curat(h):
    h = re.sub(r"data-quiz='.*?'>", ">", h, flags=re.S)
    h = re.sub(r"(?s)<[^>]+>", " ", h)
    return re.sub(r"\s+", " ", html.unescape(h))


atomi = [curat(src[poz[i]:(poz[i + 1] if i + 1 < len(poz) else fin)]) for i in range(len(poz))]
inainte = curat(src[:start_atomi])  # sectiunea „Incearca tu” (inclusiv indiciile pliate)
TERMEN = {  # intrebare -> termenul fara de care nu se poate raspunde
    "Ce vizualizare folosesti pentru a vedea toate slide-urile ca miniaturi": "Slide Sorter",
    "Cum iesiti din modul Slide Show": "ESC",
    "Cum adaugi un slide nou RAPID": "Ctrl \\+ M|Ctrl\\+M",
    "Unde gasesti toate instrumentele de editare": "Ribbon",
    "La ce este utila zona de note": "Zona de Note|zona de note",
    "Ce tasta porneste prezentarea": "F5",
    "Ce combinatie de taste salveaza": "Ctrl \\+ S|Ctrl\\+S",
    "Ce extensie au fisierele PowerPoint moderne": "pptx",
    "Ce tab din Ribbon folosesti pentru a insera imagini": "Insert:",
    "Ce este un \\\\?\"?slide": "slide-uri",
}
quiz = []
for i, m in enumerate(re.finditer(r"data-quiz='(.*?)'>", src, re.S)):
    for q in json.loads(m.group(1)):
        for k, term in TERMEN.items():
            if re.search(k, q["question"]):
                in_atom = [j + 1 for j, a in enumerate(atomi) if re.search(term, a)]
                quiz.append({"atom_intrebare": i + 1, "intrebare": q["question"], "termen": term,
                             "predat_in_atomii": in_atom, "apare_in_incearca_tu": bool(re.search(term, inainte)),
                             "predat_inainte_sau_in_atomul_intrebarii": bool(in_atom and min(in_atom) <= i + 1)})
rez["quiz_vs_atomi"] = quiz
rez["quiz_nepredate_la_momentul_intrebarii"] = sum(1 for q in quiz if not q["predat_inainte_sau_in_atomul_intrebarii"])

# ---- 2. Ex.3: 8 slide-uri - pptx vs PDF randat de LibreOffice ----
p3 = Presentation(str(L / "produs_elev/Scoala_mea.pptx"))
pdf3 = fitz.open(str(L / "randat_pptx/Scoala_mea.pdf"))
rez["ex3_slideuri_pptx"] = len(p3.slides)
rez["ex3_pagini_pdf_libreoffice"] = pdf3.page_count
rez["ex3_text_pdf_p1"] = pdf3[0].get_text().strip()[:60]
m = re.search(r"Introducere → (\d) slide-uri de continut → Concluzie", curat(src))
rez["ex3_structura_ceruta"] = f"1 + {m.group(1)} + 1 = {2 + int(m.group(1))}" if m else "negasit"
m2 = re.search(r"slide 1 introducere, (\d)-(\d) continut, (\d) concluzie", curat(src))
rez["ex3_structura_rezolvare"] = f"continut {m2.group(1)}-{m2.group(2)} = {int(m2.group(2)) - int(m2.group(1)) + 1} slide-uri, total {m2.group(3)}" if m2 else "negasit"

# ---- 3. Ex.1: imaginile - python-pptx vs PDF randat ----
p1 = Presentation(str(L / "produs_elev/Despre_Nume_Prenume.pptx"))
rez["ex1_imagini_pptx"] = sum(1 for s in p1.slides for sh in s.shapes if sh.shape_type == 13)
with zipfile.ZipFile(L / "produs_elev/Despre_Nume_Prenume.pptx") as z1:
    rez["ex1_imagini_in_ppt_media_zip"] = sum(1 for n in z1.namelist() if n.startswith("ppt/media/"))
    rez["ex1_referinte_imagine_in_slide_xml"] = sum(z1.read(n).decode("utf-8").count("<p:pic>") for n in z1.namelist()
                                                     if re.match(r"ppt/slides/slide\d+\.xml$", n))

# ---- 4. „Un fisier .pptx e de fapt o arhiva ZIP cu fisiere XML” ----
with zipfile.ZipFile(L / "produs_elev/Scoala_mea.pptx") as z:
    nume = z.namelist()
rez["pptx_este_zip"] = zipfile.is_zipfile(L / "produs_elev/Scoala_mea.pptx")
rez["pptx_fisiere_xml"] = sum(1 for n in nume if n.endswith(".xml"))

# ---- 5. Ctrl+M / F5 / Esc: lectia vs sursa Microsoft bruta ----
t_ro = (L / "surse/raw/scurtaturi_ro.txt").read_text(encoding="utf-8")
rez["ms_ro_ctrl_m"] = "Inserarea unui diapozitiv nou. Ctrl+M" in t_ro
rez["ms_ro_f5"] = "Pornirea expunerii de diapozitive. F5" in t_ro
rez["ms_ro_esc"] = "Încheierea expunerii de diapozitive. Esc" in t_ro
rez["lectia_ctrl_m"] = "Apasa Ctrl + M pentru a adauga instant un slide nou" in curat(src)
rez["file_ro_ms"] = sorted(set(re.findall(r"Deschide(?:ți)? fila (\w+(?: diapozitive)?)", t_ro)))
lectie_file = re.findall(r"<strong>(Home|Insert|Design|Transitions|Animations|Slide Show|Review|View):</strong>", src)
rez["file_lectie"] = lectie_file
rez["nume_ro_in_lectie"] = {w: len(re.findall(w, src)) for w in ["Pornire", "Inserare", "Proiectare", "Diapozitiv nou", "Sortare diapozitive", "Expunere"]}
rez["nume_en_in_lectie"] = {w: len(re.findall(w, src)) for w in ["Home", "Insert", "Design", "Slide Sorter", "Slide Show", "New Slide"]}
(L / "u3_iesire.json").write_text(json.dumps(rez, ensure_ascii=False, indent=1), encoding="utf-8")

# ---- U7: redeschid intr-un proces nou (acesta) ----
red = {}
for f in ["Prima_mea_explorare.pptx", "Despre_Nume_Prenume.pptx", "Ex2_pozitii_rezolvare.pptx", "Ex2_numere.pptx", "Scoala_mea.pptx", "CumFacSandvis.pptx"]:
    p = Presentation(str(L / "produs_elev" / f))
    red[f] = {"slideuri": len(p.slides), "titluri": [s.shapes.title.text if s.shapes.title is not None else "" for s in p.slides],
              "note": sum(1 for s in p.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())}
(L / "07_redeschis.json").write_text(json.dumps(red, ensure_ascii=False, indent=1), encoding="utf-8")

# ---- U12: sensibilitate (aceleasi formule si ritmuri ca G_poarta.py, cls6) ----
S_CLIC, CPM, WPM = 20, 50, 100
pasi = json.loads((L / "03_pasi.json").read_text(encoding="utf-8"))
cuv = json.loads((L / "masuri.json").read_text(encoding="utf-8"))["cuvinte_total"]


def sarcina(grupuri):
    return sum(p["clicuri"] * S_CLIC + p["caractere"] * 60 / CPM + p["cuvinte_citite"] * 60 / WPM for p in pasi if p["grup"] in grupuri) / 60


citire = cuv / WPM
linii = [f"cuvinte lectie {cuv} -> citire {citire:.1f} min (ritm cls6 {WPM} cuv/min, provizoriu)"]
for nume, g in [("toata lectia (incearca tu + 4 exercitii)", {"incearca_tu", "ex1", "ex2", "ex3", "ex4"}),
                ("incearca tu + Ex.1 (nivel minim)", {"incearca_tu", "ex1"}),
                ("incearca tu + Ex.2 (nivel standard)", {"incearca_tu", "ex2"}),
                ("doar incearca tu", {"incearca_tu"})]:
    s = sarcina(g)
    tot = 8 + citire + s
    dublu = 8 + (citire + s) / 2
    linii.append(f"{nume}: 8 + {citire:.1f} + {s:.1f} = {tot:.1f} min | ritm dublu {dublu:.1f} min")
(L / "u12_sensibilitate.txt").write_text("\n".join(linii) + "\n", encoding="utf-8")
print(json.dumps({k: rez[k] for k in ["quiz_nepredate_la_momentul_intrebarii", "ex3_slideuri_pptx", "ex3_pagini_pdf_libreoffice", "ex3_structura_ceruta", "ex3_structura_rezolvare", "ex1_imagini_pptx", "ex1_imagini_in_ppt_media_zip", "ex1_referinte_imagine_in_slide_xml", "pptx_este_zip", "pptx_fisiere_xml", "ms_ro_ctrl_m", "file_ro_ms", "nume_ro_in_lectie", "nume_en_in_lectie"]}, ensure_ascii=False))
print("\n".join(linii))
