"""U1 - fac sarcinile lectiei lectia5-tranzitii ca un elev de a VI-a.
python-pptx pentru continut + lxml pentru <p:transition> (python-pptx nu are API pentru tranzitii) si <p:timing> (animatiile din Ex.3,
cu clasa Timing din lectia4-animatii/u1_construieste.py, copiata aici ca _l4_u1_construieste.py).

Structura XML: ECMA-376 <p:transition spd advClick advTm> + extensiile Microsoft (MS-PPTX 2.2.1, surse/raw/mspptx_en.txt):
atributul dur (namespace powerpoint/2010/main) si morph (namespace powerpoint/2015/09/main), in mc:AlternateContent cu Fallback.
IPOTEZE (nu se afla din sursele descarcate): Push „From Right” = <p:push dir="l"/> (slide-ul nou intra de la dreapta, se misca spre stanga);
Cube = <p14:prism/>. CE NUME VEDE ELEVUL in galerie nu se afla din XML -> pasi „interfata” in 03_pasi.json.

Produce in produs_elev/:
  Tranzitii_incearca_tu.pptx   3 diapozitive, Push pe toate (Apply To All = aceeasi tranzitie scrisa pe fiecare diapozitiv)
  Tranzitii_ex1.pptx            5 diapozitive, Push + From Right + 1 s pe toate, la clic
  Tranzitii_ex2_runda1/2/3.pptx 3 diapozitive: Fade / Cube / Morph (runda 3: acelasi cerc mutat si marit pe diapozitivul 2)
  Tranzitii_ex3.pptx            6 diapozitive, doar After (5/8/8/6/10/4 s), On Mouse Click debifat, Morph pe 4, animatii After/With Previous
Iesire: u1_iesire.txt
"""
import copy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu

L = Path(__file__).resolve().parent
_src = (L / "_l4_u1_construieste.py").read_text(encoding="utf-8")
exec(_src.split("prs = Presentation()")[0].split('"""', 2)[2])  # Timing, pune_timing, caseta, P, NS_P (fara partea care construieste L4)

OUT = L / "produs_elev"
OUT.mkdir(exist_ok=True)
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"
CM = 360000
log = []


def efect_xml(tip, ns_parent):
    if tip == "push_right":
        e = etree.SubElement(ns_parent, P + "push", dir="l")
    elif tip == "fade":
        e = etree.SubElement(ns_parent, P + "fade")
    elif tip == "cube":
        e = etree.SubElement(ns_parent, "{%s}prism" % P14)
    elif tip == "morph":
        e = etree.SubElement(ns_parent, "{%s}morph" % P159, option="byObject")
    else:
        raise ValueError(tip)
    return e


def pune_tranzitie(slide, tip, dur_ms, clic=True, dupa_ms=None):
    """Scrie <p:transition> ca PowerPoint: Choice cu p14:dur (si p159 pentru morph), Fallback ECMA-376 simplu."""
    el = slide._element
    for old in el.findall("{%s}AlternateContent" % MC) + el.findall(P + "transition"):
        el.remove(old)
    req = "p159" if tip == "morph" else "p14"
    ac = etree.Element("{%s}AlternateContent" % MC, nsmap={"mc": MC})
    ch = etree.SubElement(ac, "{%s}Choice" % MC, nsmap={"p14": P14, "p159": P159}, Requires=req)
    fb = etree.SubElement(ac, "{%s}Fallback" % MC)
    for parent, cu_dur in ((ch, True), (fb, False)):
        a = {"spd": "med" if dur_ms <= 1000 else "slow"}
        if cu_dur:
            a["{%s}dur" % P14] = str(dur_ms)
        if not clic:
            a["advClick"] = "0"
        if dupa_ms is not None:
            a["advTm"] = str(dupa_ms)
        tr = etree.SubElement(parent, P + "transition", **a)
        efect_xml(tip if (cu_dur or tip in ("push_right", "fade")) else "fade", tr)
    # ordinea in p:sld: cSld, clrMapOvr, transition, timing, extLst
    anchor = el.find(P + "timing")
    if anchor is None:
        anchor = el.find(P + "extLst")
    if anchor is not None:
        anchor.addprevious(ac)
    else:
        el.append(ac)


def deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    return prs, prs.slide_layouts[6]


def salveaza(prs, nume):
    dest = OUT / nume
    prs.save(str(dest))
    log.append(f"salvat {nume}: {dest.stat().st_size} octeti, {len(prs.slides)} diapozitive")


def automat(slide):
    """Ex.3 cere „fara niciun clic”: primul efect din secventa principala porneste singur (After Previous, delay 0)."""
    el = slide._element
    ms = el.find(".//" + P + "cTn[@nodeType='mainSeq']")
    if ms is None:
        return
    for grp in ms.find(P + "childTnLst").findall(P + "par"):
        grp.find(P + "cTn/" + P + "stCondLst/" + P + "cond").set("delay", "0")
    for c in el.iter(P + "cTn"):
        if c.get("nodeType") == "clickEffect":
            c.set("nodeType", "afterEffect")


# ---- Incearca tu: 3 diapozitive, Push, Apply To All ----
prs, B = deck()
for i, t in enumerate(["Primul diapozitiv", "Al doilea diapozitiv", "Al treilea diapozitiv"]):
    s = prs.slides.add_slide(B)
    caseta(s, 3 * CM, 7 * CM, 26 * CM, 3 * CM, t, 44)
    pune_tranzitie(s, "push_right", 1000)  # durata nu e ceruta; pun 1 s ca Ex.1
salveaza(prs, "Tranzitii_incearca_tu.pptx")

# ---- Ex.1: 5 diapozitive, Push From Right 1 s, Apply To All ----
prs, B = deck()
texte = ["Hobby-ul meu: Fotbalul", "De ce imi place", "Cum practic", "Realizari", "Multumesc!"]
for t in texte:
    s = prs.slides.add_slide(B)
    caseta(s, 3 * CM, 7 * CM, 26 * CM, 3 * CM, t, 44)
    pune_tranzitie(s, "push_right", 1000)
salveaza(prs, "Tranzitii_ex1.pptx")

# ---- Ex.2: trei runde ----
for k, (tip, dur) in enumerate([("fade", 700), ("cube", 1000), ("morph", 2000)], 1):
    prs, B = deck()
    for j in range(3):
        s = prs.slides.add_slide(B)
        caseta(s, 2 * CM, 1 * CM, 28 * CM, 3 * CM, f"Runda {k} - diapozitivul {j + 1}", 36)
        if tip == "morph":
            m = 2 + j * 8
            c = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(m * CM), Emu(6 * CM), Emu((3 + 2 * j) * CM), Emu((3 + 2 * j) * CM))
            c.name = "!!Cerc"  # acelasi nume pe toate diapozitivele (potrivirea obiectelor pentru Morph)
        pune_tranzitie(s, tip, dur)  # duratele 0,7 / 1 / 2 s = presupuse, lectia nu le cere la Ex.2
    salveaza(prs, f"Tranzitii_ex2_runda{k}.pptx")

# ---- Ex.3: 6 diapozitive, complet automat ----
prs, B = deck()
plan = []
# diap. 1: titlu Fly In From Bottom, apoi subtitlu Fade; Fade, After 5 s
s = prs.slides.add_slide(B)
tit = caseta(s, 3 * CM, 4 * CM, 26 * CM, 3 * CM, "Tranzitii si animatii", 48)
sub = caseta(s, 3 * CM, 9 * CM, 26 * CM, 2 * CM, "Prezentare automata", 28)
t = Timing()
t.on_click(t.efect(tit.shape_id, "entr", 2, 4, 500, motion=("#ppt_x", "1+#ppt_h/2")))
t.after_previous(t.efect(sub.shape_id, "entr", 10, 0, 500, filt="fade"))
pune_timing(s, t)
automat(s)
pune_tranzitie(s, "fade", 700, clic=False, dupa_ms=5000)
plan.append(("fade", 5000))
# diap. 2-3: text Wipe + forma „Zoom” (in XML: preset 53, filtru fade - aproximare; conteaza cronologia); Push, After 8 s
for n in (2, 3):
    s = prs.slides.add_slide(B)
    tx = caseta(s, 2 * CM, 2 * CM, 16 * CM, 3 * CM, f"Continut {n - 1}", 36)
    fo = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(20 * CM), Emu(6 * CM), Emu(6 * CM), Emu(4 * CM))
    fo.name = "!!Forma"
    t = Timing()
    t.on_click(t.efect(tx.shape_id, "entr", 22, 4, 500, filt="wipe(down)"))
    t.after_previous(t.efect(fo.shape_id, "entr", 53, 16, 500, filt="fade"))
    pune_timing(s, t)
    automat(s)
    pune_tranzitie(s, "push_right", 1000, clic=False, dupa_ms=8000)
    plan.append(("push_right", 8000))
# diap. 4: „Duplica slide-ul 3”, muta si mareste forma, Morph, After 6 s.
# Dublarea copiaza diapozitivul intreg (surse/s_dublare.txt: „The duplicate is inserted immediately after the original”); copiez XML-ul
# diapozitivului 3 CU tot cu <p:timing> - asa ramane si la elev, lectia nu spune sa stearga animatiile copiate.
s3 = prs.slides[2]
s = prs.slides.add_slide(B)
sp4 = s.shapes._spTree
for shp in list(s3.shapes._spTree)[2:]:
    sp4.append(copy.deepcopy(shp))
for shp in s.shapes:
    if shp.name == "!!Forma":
        shp.left, shp.top, shp.width, shp.height = Emu(4 * CM), Emu(8 * CM), Emu(12 * CM), Emu(8 * CM)
tim3 = s3._element.find(P + "timing")
s._element.append(copy.deepcopy(tim3))
pune_tranzitie(s, "morph", 2000, clic=False, dupa_ms=6000)
plan.append(("morph", 6000))
# diap. 5: lista cu 3 motive, randurile Appear cu delay 0,5 s; Fade, After 10 s
s = prs.slides.add_slide(B)
lst = caseta(s, 3 * CM, 3 * CM, 26 * CM, 10 * CM, "", 30,
             bullets=["Fade e discret", "Push arata directia", "Morph arata miscarea"])
t = Timing()
t.on_click(t.efect(lst.shape_id, "entr", 1, 0, 1, para=0, filt="fade"))  # Appear ~ vizibil instant (dur 1 ms)
for k in (1, 2):
    t.after_previous(t.efect(lst.shape_id, "entr", 1, 0, 1, delay=500, para=k, filt="fade"))
pune_timing(s, t)
automat(s)
pune_tranzitie(s, "fade", 700, clic=False, dupa_ms=10000)
plan.append(("fade", 10000))
# diap. 6: „Multumesc!” Grow & Turn (preset 31 in lista ECMA, aproximat cu fade); Fade, After 4 s
s = prs.slides.add_slide(B)
mt = caseta(s, 6 * CM, 7 * CM, 20 * CM, 3 * CM, "Multumesc!", 54)
t = Timing()
t.on_click(t.efect(mt.shape_id, "entr", 31, 0, 1000, filt="fade"))
pune_timing(s, t)
automat(s)
pune_tranzitie(s, "fade", 700, clic=False, dupa_ms=4000)
plan.append(("fade", 4000))
# ordinea elementelor pe diap. 4: transition trebuie inaintea lui timing
el4 = prs.slides[3]._element
ac4 = el4.find("{%s}AlternateContent" % MC)
el4.find(P + "timing").addprevious(ac4)
salveaza(prs, "Tranzitii_ex3.pptx")
log.append("Ex.3 plan: " + ", ".join(f"{a}/{b / 1000:.0f}s" for a, b in plan))
(L / "u1_iesire.txt").write_text("\n".join(log), encoding="utf-8")
print("\n".join(log))
