"""U3 + U7 + U9 pentru lectia3-text-imagini. Proces Python NOU fata de u1_construieste.py (tiparul lectiei 2).
Scrie: 07_redeschis.json, u3_iesire.json, u9_quiz.json. Tipareste un rezumat scurt (<= 20 randuri)."""
import html
import json
import re
from pathlib import Path

import fitz
from lxml import etree
from pptx import Presentation
from pptx.util import Pt

L = Path(__file__).resolve().parent
F = Path(r"C:\00\Projects\LearningHub\content\tic\cls6\m1-prezentari\lectia3-text-imagini.html")
T = (L / "innerText.txt").read_text(encoding="utf-8")
LINII = T.splitlines()
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
rez = {}

# ---------- U7: redeschid fisierele construite ----------
red = {}
for f in sorted((L / "produs_elev").glob("*.pptx")):
    p = Presentation(str(f))
    info = {"diapozitive": len(p.slides), "pe_diapozitiv": []}
    for s in p.slides:
        texte, marimi, fonturi, imagini, forme, grupuri = [], set(), set(), 0, 0, 0
        xml = etree.tostring(s._element).decode()
        for sh in s.shapes:
            if sh.shape_type == 13:
                imagini += 1
            if sh.shape_type == 6:
                grupuri += 1
            if sh.shape_type == 1:
                forme += 1
            if sh.has_text_frame if hasattr(sh, "has_text_frame") else False:
                for par in sh.text_frame.paragraphs:
                    if par.text.strip():
                        texte.append(par.text)
                    for r in par.runs:
                        if r.font.size:
                            marimi.add(r.font.size.pt)
                        if r.font.name:
                            fonturi.add(r.font.name)
        # 6x6 (regula lectiei): randuri de corp (fara titlul = primul rand) si cuvinte pe rand
        corp = texte[1:]
        info["pe_diapozitiv"].append({
            "randuri_corp": len(corp), "max_cuvinte_rand": max([len(x.split()) for x in corp] or [0]),
            "respecta_6x6": len(corp) <= 6 and max([len(x.split()) for x in corp] or [0]) <= 6,
            "font_min_pt": min(marimi) if marimi else None, "fonturi": sorted(fonturi), "imagini": imagini,
            "forme": forme, "grupuri": grupuri, "elipsa_pe_imagine": 'prst="ellipse"' in xml and "<p:pic" in xml,
            "gradient_pe_text": "gradFill" in xml and "a:rPr" in xml, "umbra": "outerShdw" in xml,
            "numerotare": "buAutoNum" in xml, "marcatori": "buChar" in xml})
    red[f.name] = info
(L / "07_redeschis.json").write_text(json.dumps(red, ensure_ascii=False, indent=1), encoding="utf-8")

# ---------- a doua cale 1: PDF randat de LibreOffice vs structura pptx (acelasi text, acelasi numar de pagini) ----------
pdf = {}
for f in sorted((L / "randat_pptx").glob("*.pdf")):
    d = fitz.open(str(f))
    pag = [d[i].get_text() for i in range(len(d))]
    pdf[f.stem] = {"pagini": len(d), "pptx_diapozitive": red.get(f.stem + ".pptx", {}).get("diapozitive"),
                   "Hobby_pe_p2": "Hobby-urile mele" in (pag[1] if len(pag) > 1 else ""),
                   "email_fictiv_gasit": any("elev@scoala.ro" in t for t in pag),
                   "Delfinul_p1": "Delfinul" in pag[0]}
rez["pdf_vs_pptx"] = pdf

# ---------- a doua cale 2: intrebarea din atom vs atomul care preda raspunsul ----------
def linie(pat, start=0):
    for i, l in enumerate(LINII[start:], start):
        if re.search(pat, l):
            return i + 1
    return None

atomi = {n: linie(rf"^{n}\. Continut") for n in range(1, 11)}
def atom_la(lin):
    k = [n for n, a in atomi.items() if a and lin and a <= lin]
    return max(k) if k else 0

intrebari = [("Cum adaugi o caseta de text noua pe un slide?", r"^Click pe butonul Text Box"),
             ("Care tip de SmartArt folosesti pentru a arata pasi", r"^Pentru pasi intr-un proces \(Pas 1"),
             ("Ce functie folosesti pentru a aduce un obiect IN FATA", r"^Controleaza ordinea straturilor"),
             ("Cum grupezi mai multe obiecte", r"^Combina multiple obiecte intr-unul singur"),
             ("Ce marime MINIMA ar trebui sa aiba fontul", r"Niciodata sub 18pt"),
             ("Care este cea mai buna practica pentru combinatia text-fundal?", r"Alege culori care contrasteaz"),
             ("Care este shortcut-ul pentru formatare BOLD", r"Shortcut: Ctrl\+B"),
             ("Care este regula \"6x6\"", r"^ Regula \"6x6\" pentru liste"),
             ("Cum inserezi o imagine de pe calculatorul tau", r"^Insert → Pictures → This Device"),
             ("Cand tragi de coltul unei imagini", r"^Tragi de coltul imaginii pentru a o redimensiona"),
             ("Unde gasesti optiunea pentru a sterge fundalul", r"^Picture Format → Remove Background"),
             ("Ce este SmartArt in PowerPoint?", r"^Definitie: SmartArt este")]
tab = []
for q, pred in intrebari:
    lq = linie(re.escape(q))
    lp = linie(pred, 95)
    tab.append({"intrebare": q, "linia": lq, "in_atomul": atom_la(lq), "predat_la_linia": lp, "predat_in_atomul": atom_la(lp),
                "inainte_de_predare": atom_la(lq) < atom_la(lp), "despre_atomul_propriu": atom_la(lq) == atom_la(lp)})
rez["intrebari_vs_predare"] = tab
rez["intrebari_inainte_de_predare"] = sum(x["inainte_de_predare"] for x in tab)
rez["intrebari_despre_atomul_propriu"] = sum(x["despre_atomul_propriu"] for x in tab)
rez["intrebari_total"] = len(tab)

# ---------- a doua cale 3: granita BAZA/EXTINDERE vs ce cer exercitiile minim/standard ----------
lin_ext = linie(r"EXTINDERE — Continut avansat")
ex1 = linie(r"^Exercitiul 1 \(Nivel minim\)")
ex2 = linie(r"^Exercitiul 2 \(Nivel standard\)")
ex3 = linie(r"^Exercitiul 3 \(Nivel performanta\)")
bloc = lambda a, b: "\n".join(LINII[a - 1:b - 1])
rez["extindere"] = {
    "banner_linia": lin_ext, "banner_in_dupa_atomul": atom_la(lin_ext),
    "banner_text": LINII[lin_ext + 1].strip() if lin_ext else "",
    "atomi_dupa_banner": [n for n, a in atomi.items() if a and lin_ext and a > lin_ext],
    "atomi_marcati_Aprofundare": [n for n in range(1, 11) if linie(rf"^{n}\. Continut — Aprofundare")],
    "ex1_minim_cere_forme(atom6)": "Forme geometrice decorative" in bloc(ex1, ex2),
    "ex1_minim_cere_6x6_contrast_consistenta(atom10)": all(k in bloc(ex1, ex2) for k in ("regula 6x6", "Contrast ridicat", "Consistenta")),
    "ex2_standard_rezolvare_cere_Align(atom9)": "aliniezi obiectele cu Align" in bloc(ex2, ex3),
    "ex2_standard_cere_regulile(atom10)": "regulile invatate" in bloc(ex2, ex3),
}

# ---------- a doua cale 4: siguranta online vs cerinta Slide 3 (in acelasi exercitiu) ----------
rez["date_personale_ex1"] = {
    "nota_siguranta": [l.strip() for l in LINII[ex1:ex2] if l.startswith("Siguranta online")][:1],
    "slide3": [l.strip() for l in LINII[ex1:ex2] if l.startswith("Slide 3 (Familie)")][:1],
    "slide2_poze_hobby": [l.strip() for l in LINII[ex1:ex2] if l.startswith("Slide 2")][:1],
    "regula_imagine_pe_fiecare_slide": any("Minim 1 imagine per slide" in l for l in LINII[ex1:ex2]),
    "slide5_cere_imagine": any(l.startswith("Slide 5") and "imagin" in l.lower() for l in LINII[ex1:ex2]),
}

# ---------- a doua cale 5: cat citeste elevul inainte de primul exercitiu ----------
cuv = lambda a, b: len(re.findall(r"\w+", re.sub(r"\[/?ASCUNS[^\]]*\]", " ", "\n".join(LINII[a:b]))))
rez["cuvinte"] = {"pana_la_Ex1": cuv(0, ex1 - 1), "Ex1": cuv(ex1 - 1, ex2 - 1), "Ex2": cuv(ex2 - 1, ex3 - 1),
                  "Ex3_si_final": cuv(ex3 - 1, len(LINII)), "atomii_6_10": cuv(lin_ext - 1, ex1 - 1)}

# ---------- a doua cale 6: contrastul exemplelor din lectie (formula W3C) ----------
def lum(h):
    c = [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    c = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]
def cr(a, b):
    la, lb = sorted((lum(a), lum(b)), reverse=True)
    return round((la + 0.05) / (lb + 0.05), 2)
rez["contrast_W3C"] = {"rosu_E74C3C_pe_portocaliu_F39C12(prost)": cr("E74C3C", "F39C12"),
                       "negru_pe_alb(bun)": cr("000000", "FFFFFF"), "text_lectie_a0a0b0_pe_1a1a2e": cr("a0a0b0", "1a1a2e")}

# ---------- U9: data-quiz (R1.1 cea mai lunga, hint) + checklist spec ----------
src = F.read_text(encoding="utf-8")
q9 = {"marime_bytes": F.stat().st_size, "style_bloc": len(re.findall(r"<style[\s>]", src)), "script_src": len(re.findall(r"<script src", src)),
      "init": re.findall(r"(\w+)\.init\('([^']+)'", src), "escaped_quotes": src.count('\\"'), "img": len(re.findall(r"<img", src)),
      "lesson_summary_div": '<div id="lesson-summary" style="display: none;">' in src,
      "todo": len(re.findall(r"MODEL_ANSWER_REQUIRED|TODO|TBD|FIXME|PLACEHOLDER", src)), "atomi": []}
for m in re.finditer(r'<div class="atom" id="(atom-\d+)" data-quiz=\'([^\']*)\'', src):
    qq = json.loads(html.unescape(m.group(2)))
    rows = []
    for it in qq:
        opts = it["options"]
        c = it["correct"]
        ci = "abcd".index(c) if isinstance(c, str) else c
        lens = [len(o) for o in opts]
        rows.append({"q": it["question"][:70], "correct": c, "raspuns": opts[ci][:50], "nr_optiuni": len(opts),
                     "cea_mai_lunga": lens[ci] == max(lens) and lens.count(max(lens)) == 1, "are_hint": bool(it.get("hint"))})
    q9["atomi"].append({"id": m.group(1), "intrebari": rows})
toate = [r for a in q9["atomi"] for r in a["intrebari"]]
q9["total_intrebari"] = len(toate)
q9["corecta_cea_mai_lunga"] = sum(r["cea_mai_lunga"] for r in toate)
q9["fara_hint"] = sum(not r["are_hint"] for r in toate)
(L / "u9_quiz.json").write_text(json.dumps(q9, ensure_ascii=False, indent=1), encoding="utf-8")
(L / "u3_iesire.json").write_text(json.dumps(rez, ensure_ascii=False, indent=1), encoding="utf-8")

print("redeschise:", {k: v["diapozitive"] for k, v in red.items()})
print("6x6 Despre_Mine:", [x["respecta_6x6"] for x in red["Despre_Mine.pptx"]["pe_diapozitiv"]], "| Slide_Prost_Bun:", [x["respecta_6x6"] for x in red["Slide_Prost_Bun.pptx"]["pe_diapozitiv"]])
print("pdf:", {k: (v["pagini"], v["pptx_diapozitive"]) for k, v in pdf.items()})
print("intrebari inainte de predare:", rez["intrebari_inainte_de_predare"], "/", len(tab), "| despre atomul propriu:", rez["intrebari_despre_atomul_propriu"])
print("  ", [(x["in_atomul"], x["predat_in_atomul"]) for x in tab])
print("extindere:", rez["extindere"])
print("date personale:", rez["date_personale_ex1"]["regula_imagine_pe_fiecare_slide"], rez["date_personale_ex1"]["slide5_cere_imagine"])
print("cuvinte:", rez["cuvinte"])
print("contrast:", rez["contrast_W3C"])
print("u9:", {k: q9[k] for k in ("marime_bytes", "style_bloc", "script_src", "escaped_quotes", "img", "total_intrebari", "corecta_cea_mai_lunga", "fara_hint", "todo", "lesson_summary_div")})
