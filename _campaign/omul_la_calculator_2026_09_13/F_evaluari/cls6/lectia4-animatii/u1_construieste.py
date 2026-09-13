"""U1 - fac sarcinile lectiei lectia4-animatii ca un elev de a VI-a: python-pptx pentru continut + lxml pentru <p:timing>
(python-pptx nu are API pentru animatii). Structura XML = cea din ECMA-376 / PresentationML (tmRoot > mainSeq > grupuri de clic).
presetID-urile sunt ipoteza mea (lista preset ECMA-376): entr Fade=10, Fly=2 (subtype 8=left, 4=bottom), emph Pulse=26,
exit Fade=10, path custom=0. CE NUME VEDE ELEVUL in galerie nu se afla din XML -> pasi „interfata” in 03_pasi.json.

Produce produs_elev/Animatii_elev.pptx:
  diap. 1  „Incearca tu”: text box „Bine ai venit!” + Fade (La clic)
  diap. 2  Ex.1 minim: titlu „Tipuri de Animatii” (Fade 1s La clic) + 4 text boxes: Fly In from Left, Pulse (clic separat),
           Fade Out = Fade de iesire (clic separat), traseu in cerc  -> 5 animatii pe diapozitiv
  diap. 3  Ex.2 standard: lista de 5 puncte, Fly In from Bottom, By Paragraph; p1 La clic, p2-5 After Previous, Delay 0.5s
  diap. 4  Ex.3 performanta: cerc cu Custom Path in S; forma 2 With Previous; forma 3 After Previous
  diap. 5  Ex.3: buton + text box cu declansator (On Click of buton) + forma cu Fade de iesire After Previous, Delay 2s
Iesire: u1_iesire.txt
"""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

L = Path(__file__).resolve().parent
OUT = L / "produs_elev"
OUT.mkdir(exist_ok=True)
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
P = "{%s}" % NS_P
log = []
DUR_IMPLICIT = 1000  # lectia (atomul 4): „Mediu (1 secunda - DEFAULT)” - neconfirmat de sursa Microsoft; il folosesc ca elevul


class Timing:
    def __init__(self):
        self.id = 2
        self.clicks = []      # fiecare clic: lista de „trepte” (after-previous); fiecare treapta: lista de efecte (with-previous)
        self.triggers = []    # (spid_buton, efect)
        self.bld = []

    def nid(self):
        self.id += 1
        return str(self.id)

    def efect(self, spid, cls, preset, sub, dur, delay=0, kind="entr", filt=None, motion=None, para=None, scale=None):
        return dict(spid=spid, cls=cls, preset=preset, sub=sub, dur=dur, delay=delay, filt=filt, motion=motion, para=para, scale=scale)

    def on_click(self, e):
        self.clicks.append([[e]])

    def with_previous(self, e):
        self.clicks[-1][-1].append(e)

    def after_previous(self, e):
        self.clicks[-1].append([e])

    def trigger(self, btn, e):
        self.triggers.append((btn, e))

    # ---- XML ----
    def _tgt(self, parent, e):
        tg = etree.SubElement(parent, P + "tgtEl")
        sp = etree.SubElement(tg, P + "spTgt", spid=str(e["spid"]))
        if e["para"] is not None:
            tx = etree.SubElement(sp, P + "txEl")
            etree.SubElement(tx, P + "pRg", st=str(e["para"]), end=str(e["para"]))

    def _bhvr(self, parent, e, dur, delay=None, attrs=None, extra=None):
        b = etree.SubElement(parent, P + "cBhvr", **(extra or {}))
        c = etree.SubElement(b, P + "cTn", id=self.nid(), dur=str(dur), fill="hold")
        if delay is not None:
            etree.SubElement(etree.SubElement(c, P + "stCondLst"), P + "cond", delay=str(delay))
        self._tgt(b, e)
        if attrs:
            al = etree.SubElement(b, P + "attrNameLst")
            for a in attrs:
                etree.SubElement(al, P + "attrName").text = a
        return b

    def _set_vis(self, parent, e, val, delay):
        s = etree.SubElement(parent, P + "set")
        self._bhvr(s, e, 1, delay, ["style.visibility"])
        etree.SubElement(etree.SubElement(s, P + "to"), P + "strVal", val=val)

    def _effect_par(self, parent, e, node):
        par = etree.SubElement(parent, P + "par")
        c = etree.SubElement(par, P + "cTn", id=self.nid(), presetID=str(e["preset"]), presetClass=e["cls"],
                             presetSubtype=str(e["sub"]), fill="hold", nodeType=node)
        if e["para"] is not None:
            c.set("grpId", "0")
        etree.SubElement(etree.SubElement(c, P + "stCondLst"), P + "cond", delay=str(e["delay"]))
        ch = etree.SubElement(c, P + "childTnLst")
        d = e["dur"]
        if e["cls"] == "entr":
            self._set_vis(ch, e, "visible", 0)
            if e["filt"]:
                ae = etree.SubElement(ch, P + "animEffect", transition="in", filter=e["filt"])
                self._bhvr(ae, e, d)
            else:  # Fly In: animatie pe ppt_x / ppt_y
                for attr, frm, to in (("ppt_x", e["motion"][0], "#ppt_x"), ("ppt_y", e["motion"][1], "#ppt_y")):
                    an = etree.SubElement(ch, P + "anim", calcmode="lin", valueType="num")
                    self._bhvr(an, e, d, attrs=[attr], extra={"additive": "base"})
                    tl = etree.SubElement(an, P + "tavLst")
                    for tm, v in (("0", frm), ("100000", to)):
                        tav = etree.SubElement(tl, P + "tav", tm=tm)
                        etree.SubElement(etree.SubElement(tav, P + "val"), P + "strVal", val=v)
        elif e["cls"] == "exit":
            ae = etree.SubElement(ch, P + "animEffect", transition="out", filter=e["filt"])
            self._bhvr(ae, e, d)
            self._set_vis(ch, e, "hidden", d - 1)
        elif e["cls"] == "emph":
            asc = etree.SubElement(ch, P + "animScale")
            b = self._bhvr(asc, e, d // 2)
            b.find(P + "cTn").set("autoRev", "1")
            etree.SubElement(asc, P + "by", x=str(e["scale"]), y=str(e["scale"]))
        elif e["cls"] == "path":
            am = etree.SubElement(ch, P + "animMotion", origin="layout", path=e["motion"], pathEditMode="relative")
            self._bhvr(am, e, d, attrs=["ppt_x", "ppt_y"])
        if e["cls"] in ("entr", "exit") or e["para"] is not None:
            key = (e["spid"], e["para"] is not None)
            if key not in [(b[0], b[1]) for b in self.bld]:
                self.bld.append(key)

    def xml(self):
        t = etree.Element(P + "timing", nsmap={"p": NS_P})
        tn = etree.SubElement(t, P + "tnLst")
        root = etree.SubElement(etree.SubElement(tn, P + "par"), P + "cTn", id="1", dur="indefinite", restart="never", nodeType="tmRoot")
        rch = etree.SubElement(root, P + "childTnLst")
        if self.clicks:
            seq = etree.SubElement(rch, P + "seq", concurrent="1", nextAc="seek")
            ms = etree.SubElement(seq, P + "cTn", id="2", dur="indefinite", nodeType="mainSeq")
            mch = etree.SubElement(ms, P + "childTnLst")
            for trepte in self.clicks:
                cp = etree.SubElement(etree.SubElement(mch, P + "par"), P + "cTn", id=self.nid(), fill="hold")
                etree.SubElement(etree.SubElement(cp, P + "stCondLst"), P + "cond", delay="indefinite")
                cch = etree.SubElement(cp, P + "childTnLst")
                offset = 0
                for k, treapta in enumerate(trepte):
                    tp = etree.SubElement(etree.SubElement(cch, P + "par"), P + "cTn", id=self.nid(), fill="hold")
                    etree.SubElement(etree.SubElement(tp, P + "stCondLst"), P + "cond", delay=str(offset))
                    tch = etree.SubElement(tp, P + "childTnLst")
                    for j, e in enumerate(treapta):
                        node = "clickEffect" if (k == 0 and j == 0) else ("withEffect" if j > 0 else "afterEffect")
                        self._effect_par(tch, e, node)
                    offset += max(e["delay"] + e["dur"] for e in treapta)
            for tag in ("prevCondLst", "nextCondLst"):
                cl = etree.SubElement(seq, P + tag)
                cond = etree.SubElement(cl, P + "cond", evt="onPrev" if tag == "prevCondLst" else "onNext", delay="0")
                etree.SubElement(etree.SubElement(cond, P + "tgtEl"), P + "sldTgt")
        for btn, e in self.triggers:
            seq = etree.SubElement(rch, P + "seq", concurrent="1", nextAc="seek")
            ic = etree.SubElement(seq, P + "cTn", id=self.nid(), restart="whenNotActive", fill="hold", evtFilter="cancelBubble", nodeType="interactiveSeq")
            cond = etree.SubElement(etree.SubElement(ic, P + "stCondLst"), P + "cond", evt="onClick", delay="0")
            etree.SubElement(etree.SubElement(cond, P + "tgtEl"), P + "spTgt", spid=str(btn))
            es = etree.SubElement(ic, P + "endSync", evt="end", delay="0")
            etree.SubElement(es, P + "rtn", val="all")
            ich = etree.SubElement(ic, P + "childTnLst")
            cp = etree.SubElement(etree.SubElement(ich, P + "par"), P + "cTn", id=self.nid(), fill="hold")
            etree.SubElement(etree.SubElement(cp, P + "stCondLst"), P + "cond", delay="0")
            tp = etree.SubElement(etree.SubElement(etree.SubElement(cp, P + "childTnLst"), P + "par"), P + "cTn", id=self.nid(), fill="hold")
            etree.SubElement(etree.SubElement(tp, P + "stCondLst"), P + "cond", delay="0")
            self._effect_par(etree.SubElement(tp, P + "childTnLst"), e, "clickEffect")
            nc = etree.SubElement(etree.SubElement(seq, P + "nextCondLst"), P + "cond", evt="onClick", delay="0")
            etree.SubElement(etree.SubElement(nc, P + "tgtEl"), P + "spTgt", spid=str(btn))
        bl = etree.SubElement(t, P + "bldLst")
        for spid, para in self.bld:
            a = {"spid": str(spid), "grpId": "0"}
            if para:
                a["build"] = "p"
            else:
                a["animBg"] = "1"
            etree.SubElement(bl, P + "bldP", **a)
        return t


def pune_timing(slide, tm):
    el = slide._element
    for old in el.findall(P + "timing"):
        el.remove(old)
    x = tm.xml()
    ext = el.find(P + "extLst")
    if ext is not None:
        ext.addprevious(x)
    else:
        el.append(x)


def caseta(slide, x, y, w, h, text, size=28, bullets=None):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    rows = bullets or [text]
    for i, r in enumerate(rows):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " + r) if bullets else r
        p.runs[0].font.size = Pt(size)
    return tb


prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
BLANK = prs.slide_layouts[6]
CM = 360000

# --- diap. 1: Incearca tu ---
s = prs.slides.add_slide(BLANK)
tb = caseta(s, 3 * CM, 7 * CM, 20 * CM, 3 * CM, "Bine ai venit!", 48)
t = Timing()
t.on_click(t.efect(tb.shape_id, "entr", 10, 0, 500, filt="fade"))
pune_timing(s, t)
log.append(f"diap.1 Incearca tu: text box id={tb.shape_id}, Fade La clic")

# --- diap. 2: Ex.1 minim ---
s = prs.slides.add_slide(BLANK)
tit = caseta(s, 2 * CM, 1 * CM, 29 * CM, 3 * CM, "Tipuri de Animatii", 40)
cutii = {}
for i, nume in enumerate(["Entrance", "Emphasis", "Exit", "Motion Paths"]):
    cutii[nume] = caseta(s, (2 + i * 7.5) * CM, 8 * CM, 7 * CM, 2.5 * CM, nume, 24)
t = Timing()
t.on_click(t.efect(tit.shape_id, "entr", 10, 0, 1000, filt="fade"))                                    # Fade, 1s, On Click
t.on_click(t.efect(cutii["Entrance"].shape_id, "entr", 2, 8, DUR_IMPLICIT, motion=("0-#ppt_w/2", "#ppt_y")))  # Fly In from Left
t.on_click(t.efect(cutii["Emphasis"].shape_id, "emph", 26, 0, DUR_IMPLICIT, scale=110000))               # Pulse, clic separat
t.on_click(t.efect(cutii["Exit"].shape_id, "exit", 10, 0, DUR_IMPLICIT, filt="fade"))                   # „Fade Out” = Fade de iesire
t.on_click(t.efect(cutii["Motion Paths"].shape_id, "path", 0, 0, 2000,
                   motion="M 0 0 C 0.1 0 0.1 0.18 0 0.18 C -0.1 0.18 -0.1 0 0 0 E"))                      # „miscare in cerc”
pune_timing(s, t)
log.append("diap.2 Ex.1: 5 animatii (titlu + 4 casete), fiecare pe clic separat")

# --- diap. 3: Ex.2 standard ---
s = prs.slides.add_slide(BLANK)
av = ["atrage atentia", "controleaza ritmul", "evidentiaza idei cheie", "face prezentarea mai clara", "ajuta memorarea"]
lst = caseta(s, 3 * CM, 3 * CM, 26 * CM, 12 * CM, "", 28, bullets=av)
t = Timing()
fly_bottom = ("#ppt_x", "1+#ppt_h/2")
t.on_click(t.efect(lst.shape_id, "entr", 2, 4, DUR_IMPLICIT, motion=fly_bottom, para=0))
for k in range(1, 5):
    t.after_previous(t.efect(lst.shape_id, "entr", 2, 4, DUR_IMPLICIT, delay=500, motion=fly_bottom, para=k))
pune_timing(s, t)
log.append("diap.3 Ex.2: 5 paragrafe, p0 La clic, p1-4 After Previous + Delay 0.5s")

# --- diap. 4: Ex.3 slide 1 ---
s = prs.slides.add_slide(BLANK)
cerc = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(2 * CM), Emu(2 * CM), Emu(2 * CM), Emu(2 * CM))
f2 = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(12 * CM), Emu(12 * CM), Emu(4 * CM), Emu(2 * CM))
f3 = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(20 * CM), Emu(12 * CM), Emu(4 * CM), Emu(2 * CM))
t = Timing()
t.on_click(t.efect(cerc.shape_id, "path", 0, 0, 2000, motion="M 0 0 C 0.3 0 0.3 0.25 0.15 0.25 C 0 0.25 0 0.5 0.3 0.5 E"))  # S
t.with_previous(t.efect(f2.shape_id, "entr", 10, 0, DUR_IMPLICIT, filt="fade"))
t.after_previous(t.efect(f3.shape_id, "entr", 10, 0, DUR_IMPLICIT, filt="fade"))
pune_timing(s, t)
log.append("diap.4 Ex.3: Custom Path S (La clic) + forma2 With Previous + forma3 After Previous")

# --- diap. 5: Ex.3 slide 2 ---
s = prs.slides.add_slide(BLANK)
buton = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(2 * CM), Emu(2 * CM), Emu(6 * CM), Emu(2 * CM))
buton.text_frame.text = "Apasa aici"
mesaj = caseta(s, 10 * CM, 2 * CM, 18 * CM, 3 * CM, "Ai apasat butonul!", 32)
dispare = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(12 * CM), Emu(9 * CM), Emu(5 * CM), Emu(5 * CM))
t = Timing()
# Lectia: Exit „cu Start: After Previous, Delay: 2s ... dupa alt eveniment”; singurul eveniment anterior posibil e declansatorul,
# care e in secventa interactiva, nu in cea principala -> After Previous ca PRIM efect al secventei principale = porneste la intrarea pe diapozitiv.
t.trigger(buton.shape_id, t.efect(mesaj.shape_id, "entr", 10, 0, 500, filt="fade"))
t.on_click(t.efect(dispare.shape_id, "exit", 10, 0, DUR_IMPLICIT, delay=2000, filt="fade"))
pune_timing(s, t)
sl = s._element
for c in list(sl.iter(P + "cTn")):
    if c.get("nodeType") == "clickEffect" and c.get("presetClass") == "exit":
        c.set("nodeType", "afterEffect")
        grp = c
        for _ in range(6):  # efect cTn -> par -> childTnLst -> cTn treapta -> par -> childTnLst -> cTn grup
            grp = grp.getparent()
        grp.find(P + "stCondLst").find(P + "cond").set("delay", "0")
log.append("diap.5 Ex.3: buton (declansator) -> mesaj; cerc cu Fade de iesire After Previous, Delay 2s")

dest = OUT / "Animatii_elev.pptx"
prs.save(str(dest))
log.append(f"salvat {dest.name}: {dest.stat().st_size} octeti, {len(prs.slides)} diapozitive")
(L / "u1_iesire.txt").write_text("\n".join(log), encoding="utf-8")
print("\n".join(log))
