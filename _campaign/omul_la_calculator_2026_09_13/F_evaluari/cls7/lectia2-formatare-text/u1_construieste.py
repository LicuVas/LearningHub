"""U1 — fac exercitiile lectiei 2 (Formatarea textului) ca elevul, cu python-docx, pe A4.
Folosesc EXACT ce cere lectia (formatare directa pe caractere; Format Painter = rezultatul lui: aceeasi formatare pe toate 4).
Contra-proba (nu ceruta de lectie): Ex.3 cu stilul Heading 2 modificat o singura data, ca sa compar ce ramane in structura.
Iesire: produs_elev/*.docx + u1_iesire.json; tipareste <= 20 de randuri."""
import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE  # noqa: F401
from docx.shared import Mm, Pt, RGBColor

L = Path(__file__).resolve().parent
OUT = L / "produs_elev"
OUT.mkdir(exist_ok=True)
ALBASTRU = RGBColor(0x00, 0x70, 0xC0)  # „Albastru” din paleta standard Word (culoarea exacta nu e ceruta de lectie)


def a4(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    return doc


def run(p, text, **fmt):
    r = p.add_run(text)
    for k, v in fmt.items():
        if k == "font":
            r.font.name = v
        elif k == "size":
            r.font.size = Pt(v)
        elif k == "color":
            r.font.color.rgb = v
        else:
            setattr(r.font, k, v)
    return r


# ---- Tema intr-un singur document (asa preda elevul: un fisier pe ora) ----
doc = a4(Document())
# Ex.1 — textul EXACT din lectie (fara diacritice, cum il da lectia elevului)
t1 = ("Calculatorul este un dispozitiv electronic care proceseaza date. Primul calculator electronic, ENIAC, "
      "a fost construit in 1945. Astazi, calculatoarele sunt prezente in viata de zi cu zi: telefoane, tablete, laptopuri.")
p = doc.add_paragraph()
i_calc = t1.index("Calculatorul")
i_disp = t1.index("dispozitiv electronic")
i_eniac = t1.index("ENIAC")
run(p, "Calculatorul", bold=True)
run(p, t1[len("Calculatorul"):i_disp])
run(p, "dispozitiv electronic", underline=True)
run(p, t1[i_disp + len("dispozitiv electronic"):i_eniac])
run(p, "ENIAC", italic=True)
run(p, t1[i_eniac + len("ENIAC"):])
doc.add_paragraph()
# Ex.2 — propozitia de 3 ori, trei fonturi
t2 = "Fontul potrivit face textul mai usor de citit."
for f in ("Times New Roman", "Arial", "Courier New"):
    run(doc.add_paragraph(), t2, font=f)
doc.add_paragraph()
# Ex.3 — 4 subtitluri, rand gol intre ele; primul formatat de mana, restul prin Format Painter (= aceeasi formatare directa)
SUB = ["Introducere", "Capitolul 1", "Capitolul 2", "Concluzii"]
for k, s in enumerate(SUB):
    run(doc.add_paragraph(), s, font="Arial", size=16, bold=True, color=ALBASTRU)
    if k < 3:
        doc.add_paragraph()
doc.save(OUT / "Tema_Formatare_Text.docx")

# ---- Contra-proba: Ex.3 cu stil de titlu (ce ar face un tehnoredactor) ----
d2 = a4(Document())
h = d2.styles["Heading 2"]
h.font.name, h.font.size, h.font.bold, h.font.color.rgb = "Arial", Pt(16), True, ALBASTRU
for k, s in enumerate(SUB):
    d2.add_paragraph(s, style="Heading 2")
    if k < 3:
        d2.add_paragraph("Un rand de continut sub subtitlu.")
d2.save(OUT / "Ex3_contraproba_stil_Heading2.docx")

# ---- Scenariul „schimbam culoarea subtitlurilor in verde” pe ambele variante: cate operatii? ----
rez = {
    "ex1_runs": [(r.text, bool(r.bold), bool(r.italic), bool(r.underline)) for r in doc.paragraphs[0].runs],
    "ex2_fonturi": [p.runs[0].font.name for p in doc.paragraphs[2:5]],
    "ex3_subtitluri": [(p.text, p.style.name, p.runs[0].font.name, p.runs[0].font.size.pt, bool(p.runs[0].bold),
                        str(p.runs[0].font.color.rgb)) for p in doc.paragraphs if p.text in SUB],
    "ex3_contraproba": [(p.text, p.style.name) for p in d2.paragraphs if p.text in SUB],
    "modificare_ulterioara_operatii": {"manual_format_painter": "1 formatare + 1 dublu-clic + 3 selectii (de refacut la fiecare schimbare)",
                                       "stil": "1 modificare a stilului Heading 2 -> toate 4 se schimba"},
    "a4_mm": [round(doc.sections[0].page_width.mm), round(doc.sections[0].page_height.mm)],
}
(L / "u1_iesire.json").write_text(json.dumps(rez, ensure_ascii=False, indent=1), encoding="utf-8")
print("Ex1:", rez["ex1_runs"][0], rez["ex1_runs"][2], rez["ex1_runs"][4])
print("Ex2:", rez["ex2_fonturi"])
for x in rez["ex3_subtitluri"]:
    print("Ex3:", x)
print("Contraproba:", rez["ex3_contraproba"])
print("A4:", rez["a4_mm"])
