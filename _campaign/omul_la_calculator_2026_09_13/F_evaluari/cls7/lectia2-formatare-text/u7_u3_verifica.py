"""U7 + U3: redeschid artefactele intr-un proces nou si verific pe a doua cale (PDF randat de LibreOffice, citit cu PyMuPDF).
Plus: cheia chestionarului 9 (din HTML-ul lectiei) contra paginii Microsoft descarcate; 1 pt in mm calculat.
Iesire: 07_redeschis.json, u3_iesire.json; tipareste <= 20 randuri."""
import html
import json
import re
from pathlib import Path

import fitz
from docx import Document

L = Path(__file__).resolve().parent
LECTIE = Path(r"C:\00\Projects\LearningHub\content\tic\cls7\m1-word-fundamente\lectia2-formatare-text.html")

# ---- U7: docx redeschis ----
d = Document(str(L / "produs_elev" / "Tema_Formatare_Text.docx"))
p0 = d.paragraphs[0]
red = {
    "fisier": "produs_elev/Tema_Formatare_Text.docx",
    "paragrafe_cu_text": len([p for p in d.paragraphs if p.text.strip()]),
    "Calculatorul_bold": [r.bold for r in p0.runs if r.text == "Calculatorul"],
    "Concluzii": [(p.style.name, p.runs[0].font.name, p.runs[0].font.size.pt) for p in d.paragraphs if p.text == "Concluzii"],
    "pagina_mm": [round(d.sections[0].page_width.mm), round(d.sections[0].page_height.mm)],
}
(L / "07_redeschis.json").write_text(json.dumps(red, ensure_ascii=False, indent=1), encoding="utf-8")

# ---- U3 cale 2: PDF randat de LibreOffice ----
pdf = fitz.open(str(L / "randat_docx" / "Tema_Formatare_Text.pdf"))
spans = [s for b in pdf[0].get_text("dict")["blocks"] for ln in b.get("lines", []) for s in ln["spans"]]


def span(txt):
    for s in spans:
        if txt in s["text"]:
            return {"text": s["text"].strip(), "font": s["font"], "size": round(s["size"], 1),
                    "bold": bool(s["flags"] & 16) or "Bold" in s["font"], "italic": bool(s["flags"] & 2) or "Italic" in s["font"],
                    "color": f"{s['color']:06X}"}
    return None


fonturi_ex2 = [s["font"] for s in spans if "Fontul potrivit" in s["text"]]
toc_manual = pdf.get_toc()
toc_stil = fitz.open(str(L / "randat_docx" / "Ex3_contraproba_stil_Heading2.pdf")).get_toc()

src = LECTIE.read_text(encoding="utf-8")
quiz = []
for m in re.finditer(r"data-quiz='([^']*)'", src):
    quiz += json.loads(html.unescape(m.group(1)))
q9 = [q for q in quiz if "Paste Special" in q["question"]][0]
cheie9 = q9["options"]["abc".index(q9["correct"])]
ks = (L / "surse" / "raw").glob("support_microsoft_com_en_us_office_keyboard_shortcuts*.txt")
ks_txt = next(ks).read_text(encoding="utf-8").split("\n")
ctrl_alt_v = [ks_txt[i - 1] for i, x in enumerate(ks_txt) if x.strip() == "Ctrl+Alt+V"]
subscript = [ks_txt[i + 1] for i, x in enumerate(ks_txt) if x.strip() == "Apply subscript formatting."]

u3 = {
    "calculatorul_pdf": span("Calculatorul"), "eniac_pdf": span("ENIAC"), "dispozitiv_pdf": span("dispozitiv"),
    "introducere_pdf": span("Introducere"), "fonturi_ex2_pdf": fonturi_ex2,
    "pdf_semne_de_carte_ex3_manual": len(toc_manual), "pdf_semne_de_carte_ex3_stil": len(toc_stil),
    "quiz9_cheie_lectie": cheie9, "microsoft_ctrl_alt_v_rand_anterior": ctrl_alt_v,
    "microsoft_subscript": subscript, "lectie_subscript": "Ctrl+=",
    "1pt_mm_calculat": round(25.4 / 72, 4), "lectie_1pt_mm": "aproximativ 0,35",
}
(L / "u3_iesire.json").write_text(json.dumps(u3, ensure_ascii=False, indent=1), encoding="utf-8")
print("U7:", red)
for k, v in u3.items():
    print(k, "=", v)
