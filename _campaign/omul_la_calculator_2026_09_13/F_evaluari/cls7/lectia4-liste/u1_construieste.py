"""U1 — fac exercitiile lectiei 4 (Liste cu marcatori si numerotate) ca elevul, cu python-docx, pe A4.
Listele sunt LISTE REALE Word (numbering.xml: abstractNum + num, paragrafe cu numPr/ilvl), nu text "1." tastat.
Formatele pe niveluri sunt cele pe care LECTIA le spune (atomul 4 si rezolvarea Ex.2: nivel 1 = 1,2,3, nivel 2 = a,b,c, nivel 3 = i,ii,iii;
lista cu marcatori: nivel 1 punct plin, nivel 2 cerc gol).

Produse (produs_elev/):
  Tema_Liste.docx              Ex.1 + Ex.2 + Ex.3 facute cum cere ENUNTUL (Ex.3: reguli I-V, sub-puncte cu marcatori)
  Ex3_dupa_rezolvare_Tab.docx  Ex.3 facut dupa REZOLVARE: Define New Number Format schimba doar nivelul 1 in I, II, III;
                               Tab muta sub-punctul pe nivelul 2, care ramane in formatul listei numerotate (a, b, c - cum spune
                               chiar lectia la Ex.2). Nimic nu transforma nivelul 2 in marcatori.
  Renumerotare_lista_reala.docx / Renumerotare_manual.docx   provocarea „De ce?": sterg pasul 2 din 4
  Provocare_Continue.docx      3 niveluri + paragraf intre Etapa 1 si Etapa 2; Etapa 2 in aceeasi lista (Continue Numbering)
                               vs o lista repornita (Restart at 1)
Iesire: u1_iesire.json; tipareste <= 20 randuri."""
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Cm, Mm, Pt

L = Path(__file__).resolve().parent
OUT = L / "produs_elev"
OUT.mkdir(exist_ok=True)
_next = {"abs": 900, "num": 900}


def doc_nou():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = s.right_margin = Cm(2.54)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    return d


def lista(d, niveluri, start_override=None, abstract=None):
    """niveluri: [(numFmt, lvlText, font_sau_None)] -> numId. abstract=id existent => alt num pe acelasi abstractNum."""
    numbering = d.part.numbering_part.element
    if abstract is None:
        _next["abs"] += 1
        abstract = _next["abs"]
        lv = ""
        for i, (fmt, txt, font) in enumerate(niveluri):
            rpr = f'<w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}"/></w:rPr>' if font else ""
            lv += (f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{txt}"/>'
                   f'<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="{720 * (i + 1)}" w:hanging="360"/></w:pPr>{rpr}</w:lvl>')
        an = parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract}"><w:multiLevelType w:val="hybridMultilevel"/>{lv}</w:abstractNum>')
        nums = numbering.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num")
        if nums:
            nums[0].addprevious(an)
        else:
            numbering.append(an)
    _next["num"] += 1
    ov = f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="{start_override}"/></w:lvlOverride>' if start_override else ""
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{_next["num"]}"><w:abstractNumId w:val="{abstract}"/>{ov}</w:num>'))
    return _next["num"], abstract


def item(d, text, num_id, ilvl=0):
    p = d.add_paragraph(text)
    p._p.get_or_add_pPr().insert(0, parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'))
    return p


def titlu(d, text, size, centrat=False):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold, r.font.size = True, Pt(size)          # lectia cere Bold + dimensiune (formatare manuala), nu stil de titlu
    if centrat:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER   # Ctrl+E
    return p


MARCATORI = [("bullet", "•", None), ("bullet", "o", "Courier New"), ("bullet", "▪", None)]
NUMERE = [("decimal", "%1.", None), ("lowerLetter", "%2.", None), ("lowerRoman", "%3.", None)]
ROMANE_MARCATORI = [("upperRoman", "%1.", None), ("bullet", "•", None)]
ROMANE_DUPA_TAB = [("upperRoman", "%1.", None), ("lowerLetter", "%2.", None)]

REGULI = [("Respectă-ți colegii", ["Ascultă-l pe cel care vorbește", "Cere cuvântul ridicând mâna"]),
          ("Fii punctual", ["Intră în clasă înainte de sunet", "Anunță dacă întârzii"]),
          ("Păstrează curățenia", ["Nu lăsa gunoaie pe bancă", "Șterge tabla la sfârșitul orei", "Păstrează-ți locul ordonat"]),
          ("Participă activ la ore", ["Pune întrebări când nu înțelegi", "Ajută-ți colegii la exerciții"]),
          ("Ai grijă de materialele școlii", ["Închide calculatorul corect", "Pune scaunul la loc"])]

res = {}
# ---------------- Tema: Ex.1, Ex.2, Ex.3 (dupa enunt) ----------------
d = doc_nou()
titlu(d, "Lista de cumpărături", 16)
n1, _ = lista(d, MARCATORI)
for cat, prod in [("Fructe", ["mere", "banane", "portocale", "căpșuni"]), ("Lactate", ["lapte", "iaurt", "brânză"]),
                  ("Panificație", ["pâine", "cornuri", "covrigi"])]:
    item(d, cat, n1, 0)
    for x in prod:
        item(d, x, n1, 1)
d.add_page_break()
titlu(d, "Clătite", 18, centrat=True)
d.add_paragraph("Ingrediente")
n2b, _ = lista(d, MARCATORI)
for x in ["2 ouă", "250 ml lapte", "150 g făină", "o lingură de zahăr", "un praf de sare", "ulei pentru tigaie"]:
    item(d, x, n2b, 0)
d.add_paragraph("Mod de preparare")
n2n, _ = lista(d, NUMERE)
pasi = [("Pregătește ingredientele", []), ("Prepară aluatul", ["Amestecă făina cu zahărul", "Adaugă ouăle pe rând"]),
        ("Lasă aluatul să stea 10 minute", []), ("Coace clătitele", ["Încinge tigaia cu puțin ulei", "Toarnă un polonic de aluat"]),
        ("Servește-le cu gem", [])]
for p_, subs in pasi:
    item(d, p_, n2n, 0)
    for s in subs:
        item(d, s, n2n, 1)
d.add_page_break()
titlu(d, "Regulamentul clasei", 20, centrat=True)
n3, _ = lista(d, ROMANE_MARCATORI)
for r_, subs in REGULI:
    item(d, r_, n3, 0)
    for s in subs:
        item(d, s, n3, 1)
d.save(OUT / "Tema_Liste.docx")

# ---------------- Ex.3 dupa rezolvare (Tab pastreaza formatul numerotat pe nivelul 2) ----------------
d = doc_nou()
titlu(d, "Regulamentul clasei", 20, centrat=True)
n3t, _ = lista(d, ROMANE_DUPA_TAB)
for r_, subs in REGULI:
    item(d, r_, n3t, 0)
    for s in subs:
        item(d, s, n3t, 1)
d.save(OUT / "Ex3_dupa_rezolvare_Tab.docx")

# ---------------- Provocarea „De ce?": sterg pasul 2 din 4 ----------------
PASI = ["Deschide fișierul", "Scrie textul", "Salvează documentul", "Închide programul"]
fara2 = [p_ for i, p_ in enumerate(PASI) if i != 1]
d = doc_nou()
nr, _ = lista(d, NUMERE)
for x in fara2:
    item(d, x, nr, 0)
d.save(OUT / "Renumerotare_lista_reala.docx")
d = doc_nou()
for i, x in enumerate(PASI):
    if i != 1:
        d.add_paragraph(f"{i + 1}. {x}")          # „1. 2. 3." tastat ca text: nimic nu se renumeroteaza
d.save(OUT / "Renumerotare_manual.docx")

# ---------------- Provocarea: 3 niveluri + Continue Numbering vs Restart at 1 ----------------
d = doc_nou()
TREI = [("decimal", "Etapa %1", None), ("lowerLetter", "%2.", None), ("bullet", "•", None)]
nc, ab = lista(d, TREI)
item(d, "Documentarea", nc, 0)
item(d, "Caut surse", nc, 1)
item(d, "Două cărți din bibliotecă", nc, 2)
item(d, "Notez ideile", nc, 1)
d.add_paragraph("Între etape, profesorul verifică notițele.")
item(d, "Redactarea", nc, 0)                     # aceeasi lista = Continue Numbering
d.add_paragraph("Varianta repornită:")
nr1, _ = lista(d, None, start_override=1, abstract=ab)
item(d, "Redactarea", nr1, 0)                    # lista noua cu startOverride = Restart at 1
d.save(OUT / "Provocare_Continue.docx")

res["fisiere"] = sorted(p.name for p in OUT.glob("*.docx"))
(L / "u1_iesire.json").write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
print("\n".join(res["fisiere"]))
