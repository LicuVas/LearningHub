"""U7 (redeschid in proces nou) + U3 (a doua cale) + U16 (diacritice). Scrie 07_redeschis.json, u3_iesire.json."""
import json
import re
from pathlib import Path

import docx
import fitz

L = Path(__file__).resolve().parent
P = L / "produs_elev"
R = L / "randat_docx"

# ---- U7
e1 = docx.Document(str(P / "Exercitiu1_Interfata.docx"))
st = docx.Document(str(P / "Tema_Word_Structurata.docx"))
sec = e1.sections[0]
red = {
    "Exercitiu1_Interfata.docx": {
        "paragraf_1": e1.paragraphs[0].text,
        "paragrafe_cu_text": len([p for p in e1.paragraphs if p.text.strip()]),
        "stil_titlu": e1.paragraphs[0].style.name,
        "pagina_mm": [round(sec.page_width.mm), round(sec.page_height.mm)],
    },
    "Tema_Word_Structurata.docx": {
        "titluri": [(p.style.name, p.text) for p in st.paragraphs if p.style.name.startswith("Heading")],
    },
    "Tema_Word.pdf_exista": (R / "Tema_Word.pdf").is_file(),
}
(L / "07_redeschis.json").write_text(json.dumps(red, ensure_ascii=False, indent=1), encoding="utf-8")

# ---- U3: a doua cale
u3 = {}
toc = fitz.open(str(R / "Tema_Word_Structurata.pdf")).get_toc()
u3["titluri_docx_vs_semne_pdf_libreoffice"] = {
    "docx": [[int(s.split()[-1]), t] for s, t in red["Tema_Word_Structurata.docx"]["titluri"]],
    "pdf_toc": [[lvl, t] for lvl, t, _ in toc],
}
t_docx = " ".join(p.text for p in docx.Document(str(P / "Tema_Word.docx")).paragraphs)
t_pdf = fitz.open(str(R / "Tema_Word.pdf"))[0].get_text()
u3["cuvinte_Tema_Word_docx_vs_pdf"] = {"docx": len(t_docx.split()), "pdf": len(t_pdf.split())}
t_str = [p.text for p in st.paragraphs if p.style.name == "Normal"][:5]
u3["text_copiat_identic"] = t_str == [p.text for p in docx.Document(str(P / "Tema_Word.docx")).paragraphs]

# ---- U16: diacritice (aceeasi formula ca poarta, scrisa independent)
txt = (L / "innerText.txt").read_text(encoding="utf-8")
txt = re.sub(r"\[/?ASCUNS[^\]]*\]", " ", txt)
lit = len(re.findall(r"[A-Za-zăâîșțşţĂÂÎȘȚŞŢ]", txt))
dia = len(re.findall(r"[ăâîșțşţĂÂÎȘȚŞŢ]", txt))
u3["diacritice_la_1000"] = round(1000 * dia / max(lit, 1), 1)
u3["diacritice_gasite"] = re.findall(r"\w*[ăâîșțşţĂÂÎȘȚŞŢ]\w*", txt)
u3["sedila_s_t"] = len(re.findall(r"[şţŞŢ]", txt))
# cuvinte fara diacritice care in romana le cer (proba mica, numarata)
u3["cuvinte_care_cer_diacritice"] = {w: len(re.findall(r"\b" + w + r"\b", txt, re.I))
                                     for w in ["si", "in", "fereastra", "lectie", "invata", "salveaza", "pasi", "fisier", "fisierul"]}
# cuvinte ale lectiei de citit (compar cu masuri.json)
u3["cuvinte_lectie_numarate_independent"] = len(txt.split())
u3["cuvinte_masuri_json"] = json.loads((L / "masuri.json").read_text(encoding="utf-8"))["cuvinte_total"]
(L / "u3_iesire.json").write_text(json.dumps(u3, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps(red, ensure_ascii=False)[:600])
print({k: v for k, v in u3.items() if k != "diacritice_gasite"})
print("diacritice:", u3["diacritice_gasite"])
