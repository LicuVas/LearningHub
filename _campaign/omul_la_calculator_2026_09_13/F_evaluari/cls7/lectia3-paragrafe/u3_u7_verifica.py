"""U3 + U7 — a doua cale, intr-un proces nou:
(a) PDF-urile randate de LibreOffice citite cu PyMuPDF: pozitia fiecarui paragraf in Ex.3 (Enter-uri goale vs Spacing After 10 pt),
    comparata cu calculul din valorile docx (a doua cale: aritmetica pe setari);
(b) aliniere/indentare observate in PDF (x-ul randurilor) vs setarile din docx redeschis;
(c) cheile chestionarelor din data-quiz vs indiciul (hint) al aceluiasi atom.
Iesire: u3_iesire.json + 07_redeschis.json; tipareste <= 20 randuri."""
import html
import json
import re
from pathlib import Path

import fitz
from docx import Document

L = Path(__file__).resolve().parent
P = L / "produs_elev"
RD = L / "randat_docx"
LESSON = Path(r"C:\00\Projects\LearningHub\content\tic\cls7\m1-word-fundamente\lectia3-paragrafe.html")
res = {}


def linii(pdf, pagina=0):
    out = []
    pg = fitz.open(str(pdf))[pagina]
    for b in pg.get_text("dict")["blocks"]:
        for ln in b.get("lines", []):
            t = "".join(s["text"] for s in ln["spans"]).strip()
            if t:
                out.append({"t": t[:40], "x0": round(ln["bbox"][0], 1), "x1": round(ln["bbox"][2], 1), "y": round(ln["bbox"][1], 1)})
    return out, pg.rect.width


# (a) Ex.3
def starturi(pdf):
    ls, _ = linii(pdf)
    return [x["y"] for x in ls if x["t"].startswith("Paragraful ")]


y_enter = starturi(RD / "Ex3_coleg_enteruri.pdf")
y_sp10 = starturi(RD / "Ex3_corectat_spacing10.pdf")
ls_e, _ = linii(RD / "Ex3_coleg_enteruri.pdf")
rand = sorted({x["y"] for x in ls_e})
pas_rand = round(rand[1] - rand[0], 1)  # inaltimea unui rand in paragraf (interlinie 1,15)
gap_enter = round(y_enter[1] - y_enter[0], 1)
gap_sp10 = round(y_sp10[1] - y_sp10[0], 1)
# a doua cale: aritmetica din setarile docx (2 randuri de text + 8 pt dupa; 2 paragrafe goale x (1 rand + 8 pt))
ls_s, _ = linii(RD / "Ex3_corectat_spacing10.pdf")
n = sum(1 for x in ls_s if y_sp10[0] <= x["y"] < y_sp10[1])  # randuri de text intr-un paragraf
res_n = n
calc_enter = round(n * pas_rand + 8 + 2 * (pas_rand + 8), 1)
calc_sp10 = round(n * pas_rand + 10, 1)
res["ex3"] = {"randuri_pe_paragraf": res_n, "pas_rand_pt": pas_rand, "distanta_intre_inceputuri_enter_pt": gap_enter, "calculat_enter_pt": calc_enter,
              "distanta_intre_inceputuri_spacing10_pt": gap_sp10, "calculat_spacing10_pt": calc_sp10,
              "spatiu_alb_intre_paragrafe_enter_pt": round(gap_enter - res_n * pas_rand, 1),
              "spatiu_alb_intre_paragrafe_spacing10_pt": round(gap_sp10 - res_n * pas_rand, 1),
              "inaltime_ocupata_enter_pt": round(y_enter[-1] - y_enter[0], 1),
              "inaltime_ocupata_spacing10_pt": round(y_sp10[-1] - y_sp10[0], 1)}

# (b) Tema: alinieri observate in PDF
ls1, W = linii(RD / "Tema_Paragrafe.pdf", 0)
res["ex1_linii_pdf"] = [x for x in ls1 if any(k in x["t"] for k in ("Liceul", "Piatra", "16.10", "Doamn", "Popescu", "clasa a VII"))]
ls2, _ = linii(RD / "Tema_Paragrafe.pdf", 1)
res["ex2_x0_randuri"] = [x["x0"] for x in ls2][:8]
ls3, _ = linii(RD / "Tema_Paragrafe.pdf", 2)
res["provocare_x0_randuri"] = [x["x0"] for x in ls3][:6]
res["latime_pagina_pt"] = round(W, 1)

# (c) cheile chestionarelor vs indiciu
src = LESSON.read_text(encoding="utf-8")
q = []
for m in re.finditer(r"data-quiz='(.*?)'", src, re.S):
    for it in json.loads(html.unescape(m.group(1))):
        corect = it["options"]["abc".index(it["correct"])]
        tok = lambda s: set(re.findall(r"[a-z0-9+./]{2,}", s.lower()))
        altele = set().union(*[tok(o) for o in it["options"] if o != corect])
        proprii = tok(corect) - altele
        q.append({"intrebare": it["question"][:60], "raspuns_cheie": corect, "cuvinte_proprii_cheii": sorted(proprii),
                  "cuvant_cheie_in_indiciu": any(w in tok(it["hint"]) for w in proprii)})
res["chestionare"] = q

# U7: redeschid docx intr-un proces nou si citesc setarile
dd = Document(str(P / "Tema_Paragrafe.docx"))
ps = [p for p in dd.paragraphs if p.text.strip()]
red = {
    "fisier": "produs_elev/Tema_Paragrafe.docx",
    "adresa_aliniere": str(ps[0].alignment), "corp_aliniere": str(ps[4].alignment), "corp_interlinie": ps[4].paragraph_format.line_spacing,
    "ultimul_corp_space_after_pt": ps[6].paragraph_format.space_after.pt, "semnatura_aliniere": str(ps[7].alignment),
    "ex2_first_line_cm": round(ps[9].paragraph_format.first_line_indent.cm, 2),
    "bibliografie_left_cm": round(ps[12].paragraph_format.left_indent.cm, 2), "bibliografie_first_line_cm": round(ps[12].paragraph_format.first_line_indent.cm, 2),
    "latime_pagina_mm": round(dd.sections[0].page_width.mm), "inaltime_pagina_mm": round(dd.sections[0].page_height.mm),
    "goale_in_Ex3_coleg": sum(1 for p in Document(str(P / "Ex3_coleg_enteruri.docx")).paragraphs if not p.text.strip()),
    "goale_in_Ex3_corectat": sum(1 for p in Document(str(P / "Ex3_corectat_spacing10.docx")).paragraphs if not p.text.strip()),
}
(L / "07_redeschis.json").write_text(json.dumps(red, ensure_ascii=False, indent=1), encoding="utf-8")
(L / "u3_iesire.json").write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps(res["ex3"], ensure_ascii=False))
print("ex1:", [(x["t"][:14], x["x0"], x["x1"]) for x in res["ex1_linii_pdf"]])
print("ex2 x0:", res["ex2_x0_randuri"], "| provocare x0:", res["provocare_x0_randuri"])
print("chestionare:", [x["cuvant_cheie_in_indiciu"] for x in q])
print("U7:", json.dumps(red, ensure_ascii=False)[:600])
