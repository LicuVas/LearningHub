"""U1 — fac exercitiile lectiei 3 (Paragrafe si aliniere) ca elevul, cu python-docx, pe A4.
Baza = Word pentru Microsoft 365 dupa pagina Microsoft (surse/raw/default_line_spacing_*.txt): interlinie 1,15;
dupa paragraf 8 pt (valoarea pe care o afirma lectia pentru stilul Normal). Font Calibri 11 (Aptos nu exista in LibreOffice).
Ex.1 scrisoarea · Ex.2 First Line Indent 1,27 cm · Ex.3 documentul colegului (Enter-uri goale) si varianta corectata ·
Provocarea: bibliografie cu Hanging Indent 1,27 cm. Iesire: produs_elev/*.docx + u1_iesire.json; tipareste <= 20 randuri."""
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, Mm, Pt

L = Path(__file__).resolve().parent
OUT = L / "produs_elev"
OUT.mkdir(exist_ok=True)


def doc_nou():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = s.right_margin = Cm(2.54)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    pf = st.paragraph_format
    pf.line_spacing, pf.space_after, pf.space_before = 1.15, Pt(8), Pt(0)
    return d


# ---------- Ex.1 + Ex.2 + provocarea: un singur document, ca la ora ----------
d = doc_nou()
adresa = ["Liceul de Arte „Victor Brauner”", "Piatra Neamț", "16.10.2026"]
for t in adresa:
    p = d.add_paragraph(t)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT          # Ctrl+R
d.add_paragraph("Doamnă Director,").alignment = WD_ALIGN_PARAGRAPH.CENTER
corp = [
    "Subsemnatul, elev în clasa a VII-a, vă rog să îmi aprobați organizarea unui concurs de tehnoredactare în laboratorul de informatică, "
    "în ultima săptămână din octombrie, după ore. Concursul ar dura o oră și ar fi deschis tuturor elevilor din gimnaziu.",
    "Participanții ar primi același text și ar trebui să îl formateze după o fișă cu cerințe: aliniere, spațiere și indentare. "
    "Profesorul de informatică ar evalua documentele, iar primii trei ar primi diplome.",
    "Vă mulțumesc pentru înțelegere și aștept răspunsul dumneavoastră.",
]
for i, t in enumerate(corp):
    p = d.add_paragraph(t)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY        # Ctrl+J
    p.paragraph_format.line_spacing = 1.5           # Ctrl+5
    if i == len(corp) - 1:
        p.paragraph_format.space_after = Pt(24)     # Spacing After 24 pt, fara Enter gol
for t in ["Popescu Andrei", "clasa a VII-a"]:
    d.add_paragraph(t).alignment = WD_ALIGN_PARAGRAPH.CENTER   # ce cere Ex.1: semnatura centrata

d.add_page_break()
ex2 = [
    "Paragraful este textul dintre două apăsări ale tastei Enter. Fiecare paragraf are alinierea lui. Poate avea un rând sau zece. Word îl tratează ca pe o unitate.",
    "Indentarea mută textul spre interior față de margine. Prima linie poate fi mutată separat. Asta arată cititorului unde începe ideea nouă. Nu se face cu spații.",
    "Spațierea dintre paragrafe se pune din fereastra Paragraf. Nu se fac rânduri goale. Documentul rămâne ordonat. Se modifică ușor mai târziu.",
]
for t in ex2:
    p = d.add_paragraph(t)
    p.paragraph_format.first_line_indent = Cm(1.27)  # Special > First Line > 1,27 cm

d.add_page_break()
bib = [
    "Popescu, Ion - Tehnoredactarea documentelor pentru elevi și profesori, ediția a doua revăzută - Editura Didactică și Pedagogică - 2019",
    "Ionescu, Maria - Informatică și TIC pentru clasa a VII-a, manual aprobat prin ordin al ministrului - Editura Art Klett - 2019",
    "Georgescu, Dan - Microsoft Word pas cu pas: paragrafe, stiluri, cuprinsuri și bibliografii - Editura Polirom - 2021",
]
for t in bib:
    p = d.add_paragraph(t)
    p.paragraph_format.left_indent = Cm(1.27)        # marcajul din mijloc tras la 1,27 cm = hanging
    p.paragraph_format.first_line_indent = Cm(-1.27)
d.save(OUT / "Tema_Paragrafe.docx")

# ---------- Ex.3: documentul colegului (2 Enter-uri goale intre paragrafe) si corectarea ----------
texte = [f"Paragraful {k}. " + "Acesta este un paragraf de test care ocupă cam două rânduri pe pagină, ca un text obișnuit dintr-un referat. " * 2
         for k in range(1, 6)]
c = doc_nou()
for k, t in enumerate(texte):
    c.add_paragraph(t)
    if k < len(texte) - 1:
        c.add_paragraph("")
        c.add_paragraph("")
c.save(OUT / "Ex3_coleg_enteruri.docx")

r = doc_nou()
for t in texte:
    p = r.add_paragraph(t)
    p.paragraph_format.space_after = Pt(10)          # cerinta Ex.3: Spacing After 10 pt
r.save(OUT / "Ex3_corectat_spacing10.docx")

out = {
    "Tema_Paragrafe.docx": {"paragrafe": len(d.paragraphs), "goale": sum(1 for p in d.paragraphs if not p.text.strip())},
    "Ex3_coleg_enteruri.docx": {"paragrafe": len(c.paragraphs), "goale": sum(1 for p in c.paragraphs if not p.text.strip())},
    "Ex3_corectat_spacing10.docx": {"paragrafe": len(r.paragraphs), "goale": sum(1 for p in r.paragraphs if not p.text.strip())},
}
(L / "u1_iesire.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps(out, ensure_ascii=False))
