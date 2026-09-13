"""U3 + U7 — a doua cale si redeschiderea, intr-un proces Python NOU (nu importa u1_construieste).
1) redeschide .docx-urile si citeste structura (U7);
2) citeste PDF-ul randat de LibreOffice (alta cale decat python-docx) si compara (U3);
3) chestionarele: cheia din data-quiz vs indiciu + lungimea variantelor (R1.1) — din HTML;
4) diacritice si ghilimele din innerText.txt; <img>/<table> din HTML.
Scrie u3_iesire.json, 07_redeschis.json, u9_chestionare_lungimi.json; tipareste <= 20 randuri."""
import html
import json
import re
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from docx.oxml.ns import qn

L = Path(__file__).resolve().parent
LECTIE = Path("C:/00/Projects/LearningHub/content/tic/cls7/m1-word-fundamente/lectia5-tabele.html")
P = L / "produs_elev"
out = {}

# ---- U7: redeschidere ----
red = {}
for f in sorted(P.glob("*.docx")):
    d = Document(f)
    if not d.tables:
        red[f.name] = {"paragrafe_text": [p.text for p in d.paragraphs if p.text.strip()]}
        continue
    t = d.tables[0]
    s = d.sections[0]
    unice = [len({id(c._tc) for c in row.cells}) for row in t.rows]
    look = t._tbl.tblPr.find(qn("w:tblLook"))
    red[f.name] = {"A4_mm": [round(s.page_width.mm), round(s.page_height.mm)], "randuri": len(t.rows),
                   "celule_unice_pe_rand": unice, "celula_1_1": t.cell(0, 0).text[:60],
                   "tblLook": {k.split('}')[1]: v for k, v in look.attrib.items()} if look is not None else None}
(L / "07_redeschis.json").write_text(json.dumps(red, ensure_ascii=False, indent=1), encoding="utf-8")

# ---- U3: PDF LibreOffice vs docx ----
pdf = {}
for f in sorted((L / "randat_docx").glob("*.pdf")):
    doc = fitz.open(f)
    txt = "\n".join(p.get_text() for p in doc)
    pdf[f.stem] = {"pagini": doc.page_count, "text_inceput": txt[:120].replace("\n", " | ")}
prov = pdf.get("Provocare_Conversie", {}).get("text_inceput", "")
out["pdf"] = pdf
out["provocare_titlu_si_primul_elev_in_aceeasi_celula_docx"] = red.get("Provocare_Conversie.docx", {}).get("celule_unice_pe_rand", [None])[0] == 1
out["provocare_pdf_incepe_cu_primul_elev"] = prov.startswith("Popescu Ion")

# ---- chestionare ----
h = LECTIE.read_text(encoding="utf-8")
qs = []
for m in re.finditer(r"data-quiz='([^']+)'", h):
    for q in json.loads(html.unescape(m.group(1))):
        lens = [len(o) for o in q["options"]]
        ci = "abc".index(q["correct"])
        qs.append({"q": q["question"][:70], "corecta": q["options"][ci][:60], "lungimi": lens,
                   "corecta_e_cea_mai_lunga": lens[ci] == max(lens) and lens.count(max(lens)) == 1})
(L / "u9_chestionare_lungimi.json").write_text(json.dumps(qs, ensure_ascii=False, indent=1), encoding="utf-8")
out["chestionare"] = {"total": len(qs), "corecta_cea_mai_lunga": sum(q["corecta_e_cea_mai_lunga"] for q in qs)}

# ---- text ----
it = (L / "innerText.txt").read_text(encoding="utf-8")
litere = sum(c.isalpha() for c in it)
dia = sum(it.count(c) for c in "ăâîșțĂÂÎȘȚşţŞŢ")
out["diacritice_la_1000"] = round(1000 * dia / litere, 2)
out["sedila_s_t"] = sum(it.count(c) for c in "şţŞŢ")
out["ghilimele_drepte"] = it.count('"')
out["ghilimele_romanesti"] = it.count("„")
out["img_in_html"] = len(re.findall(r"<img\b", h))
out["table_in_html"] = len(re.findall(r"<table\b", h))
out["cuvinte_lectie"] = len(it.split())
(L / "u3_iesire.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps({k: v for k, v in out.items() if k != "pdf"}, ensure_ascii=False))
print({k: v["pagini"] for k, v in pdf.items()})
print({k: (v.get("celule_unice_pe_rand"), v.get("A4_mm"), len(v.get("paragrafe_text", []))) for k, v in red.items()})
