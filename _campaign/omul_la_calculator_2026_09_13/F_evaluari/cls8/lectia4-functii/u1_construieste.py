"""U1 - fac sarcinile lectiei lectia4-functii (cls8) ca elevul, in xlsx (openpyxl); LibreOffice recalculeaza apoi (H_randeaza).
Formulele stau in forma de fisier (engleza, virgula), cum cere protocolul.
Fisiere in produs_elev/:
  incearca_note.xlsx  - Incearca: foaia Incearca (pasii 1-4 + bonus E6), foaia Pas6_B4_3 (B4 -> 3, EXACT cum scrie lectia),
                        foaia Pas6_nota10_3 (nota 10, adica B5 -> 3, ce voia probabil lectia), foaia Bonus_sters (B9 golit),
                        foaia Lipire (blocul „Copiaza” lipit din E2 in jos, citit din HTML).
  atomi_note.xlsx     - afirmatiile atomilor: MIN(B2,B5,B8), amplitudinea, AVERAGE cu B9 gol, COUNT/COUNTA A1:A5 si D1:D6,
                        MIN C1:C5, AVERAGE A1:A5 cu goluri, tabelul atomului 7, scenariul atomului 8 EXACT (B7 -> 10) si
                        varianta „nota 4” (B8 -> 10), IF pe 8/5/4/2, proba rapida, Provocarea.
  ex1_note.xlsx       - Ex. 1: foaia Initial, foaia Dan10 (B5 -> 10).
  ex2_meteo.xlsx      - Ex. 2: E2:E4 pe temperaturi, copiate la dreapta in F2:F4 (Translator), E5 amplitudinea.
  ex3_count.xlsx      - Ex. 3: A1:A10 = 5, "N/A", 8, gol, 10, 3, "absent", 7, gol, 9.
Iesire: u1_iesire.json; tipareste <= 20 de randuri."""
import html
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.formula.translate import Translator

L = Path(__file__).resolve().parent
P = L / "produs_elev"
P.mkdir(exist_ok=True)
SRC = Path(r"C:\00\Projects\LearningHub\content\tic\cls8\m1-excel-fundamente\lectia4-functii.html")
s = SRC.read_text(encoding="utf-8")
blocuri = [html.unescape(re.sub(r"<[^>]+>", "", b)) for b in re.findall(r'<div class="code-block">(.*?)</div>', s, re.S)]
out = {"blocuri_copiabile": len(blocuri), "blocuri": [b[:120] for b in blocuri], "blocuri_cu_tab": [("\t" in b) for b in blocuri]}

NOTE = [7, 9, 5, 10, 8, 6, 4, 9]


def note_in(ws, note, col="B", r0=2):
    for i, n in enumerate(note):
        ws[f"{col}{r0 + i}"] = n


def incearca(ws, note):
    ws["A1"], ws["B1"] = "Elev", "Nota"
    note_in(ws, note)
    for r, (et, f) in enumerate([("Nota minima", "=MIN(B2:B9)"), ("Nota maxima", "=MAX(B2:B9)"), ("Media", "=AVERAGE(B2:B9)"),
                                 ("Nr. note", "=COUNT(B2:B9)")], 2):
        ws[f"D{r}"], ws[f"E{r}"] = et, f
    ws["E6"] = "=COUNTA(B2:B9)"
    ws["G1"], ws["G2"] = "adresa_notei_10", "=MATCH(10,B2:B9,0)+1"          # pe ce rand sta nota 10
    ws["G3"], ws["G4"] = "adresa_notei_4", "=MATCH(4,B2:B9,0)+1"


# --- Incearca
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Incearca"
incearca(ws, NOTE)
out["incearca_continut"] = {"B4": ws["B4"].value, "B5": ws["B5"].value, "B7": ws["B7"].value, "B8": ws["B8"].value}
p6 = wb.create_sheet("Pas6_B4_3")                  # „Schimba nota din B4 din 10 in 3”
incearca(p6, NOTE)
p6["B4"] = 3
p6b = wb.create_sheet("Pas6_nota10_3")             # nota 10 sta in B5
incearca(p6b, NOTE)
p6b["B5"] = 3
bs = wb.create_sheet("Bonus_sters")                 # „sterge o nota (lasa celula goala)”
incearca(bs, NOTE)
bs["B9"] = None
lp = wb.create_sheet("Lipire")                      # blocul Copiaza lipit din E2: o linie = o celula
note_in(lp, NOTE)
bl = next((b for b in blocuri if b.strip().startswith("=MIN(B2:B9)")), "")
for i, ln in enumerate([x.strip() for x in bl.strip().split("\n") if x.strip()], 2):
    lp[f"E{i}"] = ln
out["lipire_E2_E6"] = [lp[f"E{r}"].value for r in range(2, 7)]
wb.save(P / "incearca_note.xlsx")

# --- afirmatiile atomilor
wb = openpyxl.Workbook()
a = wb.active
a.title = "Atom3_4"
note_in(a, NOTE)
a["D1"], a["E1"] = "MIN(B2,B5,B8)", "=MIN(B2,B5,B8)"
a["D2"], a["E2"] = "amplitudine", "=MAX(B2:B9)-MIN(B2:B9)"
a["D3"], a["E3"] = "AVERAGE", "=AVERAGE(B2:B9)"
g = wb.create_sheet("Atom4_B9_gol")
note_in(g, NOTE)
g["B9"] = None
g["E1"] = "=AVERAGE(B2:B9)"
c = wb.create_sheet("Atom5_count")
for r, v in enumerate([10, "absent", 8, None, 9], 1):
    c[f"A{r}"] = v
c["C1"], c["C2"] = "=COUNT(A1:A5)", "=COUNTA(A1:A5)"
for r, v in enumerate([10, "absent", 8, None, 9, "scutit"], 1):
    c[f"D{r}"] = v
c["F1"], c["F2"] = "=COUNT(D1:D6)", "=COUNTA(D1:D6)"
q = wb.create_sheet("Intrebari")
for r, v in enumerate([15, 8, 22, 3, 11], 1):
    q[f"C{r}"] = v
q["E1"] = "=MIN(C1:C5)"
for r, v in enumerate([8, 6, 10], 1):
    q[f"A{r}"] = v
q["E2"] = "=AVERAGE(A1:A5)"
t = wb.create_sheet("Atom7_tabel")
note_in(t, NOTE)
for r, (et, f) in enumerate([("Suma notelor", "=SUM(B2:B9)"), ("Nota minima", "=MIN(B2:B9)"), ("Nota maxima", "=MAX(B2:B9)"),
                             ("Media", "=AVERAGE(B2:B9)"), ("Nr. note", "=COUNT(B2:B9)"), ("Celule completate", "=COUNTA(B2:B9)"),
                             ("Amplitudine", "=MAX(B2:B9)-MIN(B2:B9)")], 1):
    t[f"D{r}"], t[f"E{r}"] = et, f
out["atom7_B1_B2"] = [t["B1"].value, t["B2"].value]
for nume, cel in (("Atom8_B7_la_10", "B7"), ("Atom8_nota4_la_10", "B8")):
    sh = wb.create_sheet(nume)
    note_in(sh, NOTE)
    out[f"{nume}_inainte"] = sh[cel].value
    sh[cel] = 10
    sh["E1"], sh["E2"], sh["E3"], sh["E4"] = "=MIN(B2:B9)", "=MAX(B2:B9)", "=AVERAGE(B2:B9)", "=COUNT(B2:B9)"
f = wb.create_sheet("Atom9_IF")
for r, v in enumerate([8, 5, 4, 2, 7], 2):
    f[f"B{r}"] = v
    f[f"C{r}"] = f'=IF(B{r}>=5,"Promovat","Corigent")'
f["D2"] = "=IF(B2>=5,B2*2,0)"
f["A10"] = 4
f["B10"] = '=IF(A10>=5,"Promovat","Corigent")'
f["A11"] = 7
f["B11"] = '=IF(A11>=5,"Promovat","Corigent")'
pv = wb.create_sheet("Provocare")
for r, v in enumerate([4, 6, 5, 3], 2):
    pv[f"B{r}"] = v
pv["C2"] = '=IF(AVERAGE(B2:B5)>=5,"Promovat","Corigent")'
for r, v in enumerate([4, 3, 5, 3], 2):
    pv[f"E{r}"] = v
pv["F2"] = '=IF(AVERAGE(E2:E5)>=5,"Promovat","Corigent")'
wb.save(P / "atomi_note.xlsx")

# --- Ex. 1
wb = openpyxl.Workbook()
for nume in ("Initial", "Dan10"):
    ws = wb.active if nume == "Initial" else wb.create_sheet(nume)
    ws.title = nume
    ws["A1"], ws["B1"] = "Elev", "Nota"
    for r, (el, n) in enumerate(zip(["Ana", "Bogdan", "Carla", "Dan", "Elena"], [8, 6, 9, 5, 7]), 2):
        ws[f"A{r}"], ws[f"B{r}"] = el, n
    ws["D2"], ws["D3"], ws["D4"], ws["D5"], ws["D6"] = "=SUM(B2:B6)", "=MIN(B2:B6)", "=MAX(B2:B6)", "=AVERAGE(B2:B6)", "=COUNT(B2:B6)"
    if nume == "Dan10":
        ws["B5"] = 10
wb.save(P / "ex1_note.xlsx")

# --- Ex. 2
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Meteo"
ws["A1"], ws["B1"], ws["C1"] = "Ziua", "Temp (C)", "Umiditate (%)"
for r, (z, tp, u) in enumerate(zip(["Luni", "Marti", "Miercuri", "Joi", "Vineri", "Sambata", "Duminica"],
                                    [18, 22, 15, 25, 28, 20, 17], [65, 58, 72, 45, 40, 68, 75]), 2):
    ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = z, tp, u
ws["E2"], ws["E3"], ws["E4"] = "=MIN(B2:B8)", "=MAX(B2:B8)", "=AVERAGE(B2:B8)"
for r in (2, 3, 4):                                  # „Repeta pentru umiditate in F2:F4” = copiat la dreapta
    ws[f"F{r}"] = Translator(ws[f"E{r}"].value, origin=f"E{r}").translate_formula(f"F{r}")
ws["E5"] = "=MAX(B2:B8)-MIN(B2:B8)"
ws["H4"] = "=ROUND(E4,2)"
out["ex2_F2_F4"] = [ws["F2"].value, ws["F3"].value, ws["F4"].value]
wb.save(P / "ex2_meteo.xlsx")

# --- Ex. 3
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Ex3"
for r, v in enumerate([5, "N/A", 8, None, 10, 3, "absent", 7, None, 9], 1):
    ws[f"A{r}"] = v
for r, fo in enumerate(["=COUNT(A1:A10)", "=COUNTA(A1:A10)", "=MIN(A1:A10)", "=MAX(A1:A10)", "=AVERAGE(A1:A10)"], 1):
    ws[f"C{r}"] = fo
wb.save(P / "ex3_count.xlsx")

# --- a doua cale: Python pur
py = {"incearca": {"MIN": min(NOTE), "MAX": max(NOTE), "AVERAGE": sum(NOTE) / 8, "COUNT": 8},
      "pas6_B4_3": (lambda n: {"MIN": min(n), "MAX": max(n), "AVERAGE": sum(n) / 8})([7, 9, 3, 10, 8, 6, 4, 9]),
      "pas6_nota10_3": (lambda n: {"MIN": min(n), "MAX": max(n), "AVERAGE": sum(n) / 8})([7, 9, 5, 3, 8, 6, 4, 9]),
      "atom8_B7_la_10": (lambda n: {"MIN": min(n), "MAX": max(n), "AVERAGE": sum(n) / 8})([7, 9, 5, 10, 8, 10, 4, 9]),
      "atom8_nota4_la_10": (lambda n: {"MIN": min(n), "MAX": max(n), "AVERAGE": sum(n) / 8})([7, 9, 5, 10, 8, 6, 10, 9]),
      "ex1": {"SUM": 35, "AVG": 35 / 5, "dupa_SUM": 40, "dupa_MIN": min([8, 6, 9, 10, 7]), "dupa_AVG": 40 / 5},
      "ex2": {"T_avg": round(145 / 7, 4), "U_avg": round(sum([65, 58, 72, 45, 40, 68, 75]) / 7, 4), "T_sum": sum([18, 22, 15, 25, 28, 20, 17])},
      "ex3": {"COUNT": 6, "COUNTA": 8, "AVG": (5 + 8 + 10 + 3 + 7 + 9) / 6}}
out["python_a_doua_cale"] = py
(L / "u1_iesire.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps({k: out[k] for k in out if k != "python_a_doua_cale"}, ensure_ascii=False)[:1500])
print(json.dumps(py, ensure_ascii=False)[:900])
