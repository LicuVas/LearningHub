"""U7 - redeschid catalogul si buletinul intr-un proces Python nou: formule, formatare, grafic, pagina (din produs_elev),
valori calculate (din recalculat/, LibreOffice) si textul din PDF-ul randat. Iesire: 07_redeschis.json; <= 20 randuri."""
import json
import re
import zipfile
from pathlib import Path

import fitz
import openpyxl

L = Path(__file__).resolve().parent
out = {}
for name in ("catalog_incearca", "ex2_buletin"):
    f = L / "produs_elev" / f"{name}.xlsx"
    ws = openpyxl.load_workbook(f).active
    wv = openpyxl.load_workbook(L / "recalculat" / f"{name}.xlsx", data_only=True).active
    z = zipfile.ZipFile(f)
    charts = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
    cx = z.read(charts[0]).decode("utf-8") if charts else ""
    pdf = fitz.open(str(L / "randat_xlsx" / f"{name}.pdf"))
    t1 = pdf[0].get_text()
    out[name] = {
        "A1": ws["A1"].value, "merged": [str(m) for m in ws.merged_cells.ranges],
        "G4_formula": ws["G4"].value, "G13_formula": ws["G13"].value, "G15_formula": ws["G15"].value,
        "G4_valoare": wv["G4"].value, "G13_valoare": wv["G13"].value, "C15_valoare": wv["C15"].value,
        "C18_valoare": wv["C18"].value,
        "antet_bold_fill": [ws["A3"].font.b, ws["A3"].fill.fgColor.rgb, ws["A3"].font.color.rgb if ws["A3"].font.color else None],
        "C4_format": ws["C4"].number_format, "G4_format": ws["G4"].number_format,
        "orientare": ws.page_setup.orientation, "margine_stanga_inch": ws.page_margins.left,
        "antet_pagina": ws.oddHeader.center.text,
        "grafic_referinte": re.findall(r"<(?:c:)?f>([^<]+)</(?:c:)?f>", cx)[:4],
        "pdf_pagini": pdf.page_count,
        "pdf_are_antetul": ws.oddHeader.center.text in t1,
    }
(L / "07_redeschis.json").write_text(json.dumps(out, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
for k, v in out.items():
    print(k, {a: b for a, b in v.items() if a in ("G4_valoare", "G13_valoare", "C18_valoare", "orientare", "antet_pagina", "pdf_pagini", "pdf_are_antetul", "grafic_referinte")})
