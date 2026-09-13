"""U3 + U5 - a doua cale pentru cifrele produsului si reproducerea defectelor.
Cale 1 = LibreOffice (recalculat/), cale 2 = Python pe datele din blocul lectiei / din fisier.
Iesire: 04_a_doua_cale.json + u5_reproducere.txt; tipareste <= 20 de randuri."""
import json
import re
from pathlib import Path

import fitz
import openpyxl

L = Path(__file__).resolve().parent
R = L / "recalculat"
u1 = json.loads((L / "u1_iesire.json").read_text(encoding="utf-8"))
rows = []
rep = []

# 1. media pe elev, randurile 4, 5 si 13 (nu doar primul)
wv = openpyxl.load_workbook(R / "catalog_incearca.xlsx", data_only=True).active
medii_py = list(u1["medii_python"].values())
for r in (4, 5, 13):
    rows.append({"ce": f"media elevului pe randul {r} (G{r})", "cale1": "LibreOffice recalculat, =AVERAGE", "val1": wv[f"G{r}"].value,
                 "cale2": "Python, suma/4 din blocul Copiaza", "val2": medii_py[r - 4]})
# 2. COUNTIF >=7
rows.append({"ce": "BONUS: elevi cu media >= 7", "cale1": "LibreOffice =COUNTIF(G4:G13,\">=7\")", "val1": wv["C18"].value,
             "cale2": "Python, numarare pe mediile calculate", "val2": sum(m >= 7 for m in medii_py)})
# 3. medii cu 3 zecimale (pasul 4 formateaza doar notele)
trei = [round(m, 3) for m in medii_py if round(m, 2) != m]
rows.append({"ce": "medii cu 3 zecimale in coloana G (pasul 4 da 2 zecimale doar notelor)", "cale1": "LibreOffice G4:G13",
             "val1": [wv[f"G{r}"].value for r in range(4, 14) if round(wv[f"G{r}"].value, 2) != wv[f"G{r}"].value],
             "cale2": "Python", "val2": trei})

# 4. Ex. 1: rezolvarea pliata =AVERAGE(B4:F4) vs media corecta pe 5 materii
wc = openpyxl.load_workbook(R / "ex1_catalog15_corect.xlsx", data_only=True).active
wr = openpyxl.load_workbook(R / "ex1_dupa_rezolvare.xlsx", data_only=True).active
dif = [(r, wc[f"H{r}"].value, wr[f"H{r}"].value) for r in range(4, 19) if abs(wc[f"H{r}"].value - wr[f"H{r}"].value) > 1e-9]
rows.append({"ce": "Ex.1: media elevului 1 (H4) - rezolvarea pliata =AVERAGE(B4:F4) vs 5 materii C4:G4",
             "cale1": "LibreOffice, formula din rezolvare", "val1": wr["H4"].value,
             "cale2": "LibreOffice, =AVERAGE(C4:G4) + Python pe 5 note", "val2": [wc["H4"].value, u1["ex1_python_medii_corecte_primele3"][0]]})
rep.append(f"[Ex1] randuri cu media gresita dupa rezolvarea pliata: {len(dif)} din 15; primele: {dif[:3]}")

# 5. Ex. 1: Status =IF(media>=5) vs ROFUIP art. 115 (cel putin 5 la FIECARE disciplina)
contra = []
for r in range(4, 19):
    note = [wc.cell(r, c).value for c in range(3, 8)]
    if wc[f"I{r}"].value == "Promovat" and min(note) < 5:
        contra.append((wc[f"B{r}"].value, note, wc[f"H{r}"].value))
rows.append({"ce": "Ex.1: elevi „Promovat” dupa =IF(media>=5) care au o nota sub 5 la o materie",
             "cale1": "LibreOffice, coloana I (IF din lectie)", "val1": len(contra),
             "cale2": "ROFUIP 2024 art. 115 (1), surse/rofuip_2024_art115.txt: minim 5 la fiecare disciplina, verificat in Python cu MIN pe rand",
             "val2": [c[0] for c in contra]})
rep.append(f"[Ex1] Promovat cu nota < 5 la o materie: {len(contra)}; exemple: {contra[:3]}")

# 6. Ex. 2: axa graficului pe mediile pe materii
txt = fitz.open(str(L / "randat_xlsx" / "ex2_buletin.pdf"))[0].get_text()
axa = [t for t in txt.split("\n") if re.fullmatch(r"7,\d\d?", t.strip())]
wb2 = openpyxl.load_workbook(R / "ex2_buletin.xlsx", data_only=True).active
med_mat = [wb2[f"{c}15"].value for c in "CDEF"]
rows.append({"ce": "Ex.2: diferenta dintre mediile pe materii si unde incepe axa graficului",
             "cale1": "PDF randat LibreOffice, etichetele axei", "val1": axa,
             "cale2": "openpyxl pe recalculat, C15:F15 (max/min)", "val2": [med_mat, round(max(med_mat) / min(med_mat), 4)]})
rep.append(f"[Ex2] axa verticala in PDF: {axa} ; mediile pe materii {med_mat} -> Engleza pare ~jumatate din Matematica desi diferenta e {max(med_mat) - min(med_mat):.2f}")

# 7. Separatorul din HTML: brut vs dupa decodare
rows.append({"ce": "formule din <code> cu ';' (B_context_real §4b: '5 cu punct-virgula')", "cale1": "grep pe HTML brut",
             "val1": u1["formule_cu_punct_virgula_brut"], "cale2": "html.unescape inainte de numarare",
             "val2": [u1["formule_cu_punct_virgula_dupa_unescape"], "entitati: " + ",".join(u1["entitati_in_formulele_cu_punct_virgula"])]})

# 8. lipire ca text (ipoteza ro-RO)
wt = openpyxl.load_workbook(R / "lipire_roRO_text.xlsx", data_only=True).active
rep.append(f"[Lipire text] G4={wt['G4'].value} C15={wt['C15'].value} C18={wt['C18'].value} (note ramase text)")

(L / "04_a_doua_cale.json").write_text(json.dumps(rows, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
(L / "u5_reproducere.txt").write_text("Rulat: python u3_verifica.py (dupa u1_construieste.py si H_randeaza.py xlsx/pdf)\n" + "\n".join(rep), encoding="utf-8")
for r in rows:
    print(r["ce"][:70], "|", r["val1"], "|", str(r["val2"])[:80])
for x in rep:
    print(x[:200])
