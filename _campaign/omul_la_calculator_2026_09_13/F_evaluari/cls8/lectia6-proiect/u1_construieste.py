"""U1 - construiesc catalogul din lectia6-proiect (cls8) exact cum cer pasii 1-6 + BONUS, apoi exercitiile 1 si 2.
Datele se iau din blocul „Copiaza” al lectiei (citit din HTML), nu scrise de mana.
Formulele in fisier = forma de fisier (engleza, virgula). Copierea in jos/dreapta = Translator (ca tragerea in Excel).
Fisiere in produs_elev/:
  catalog_incearca.xlsx     - pasii 1-6 + BONUS COUNTIF (pe C18; lectia nu spune celula, doar „randul 18”)
  lipire_roRO_text.xlsx     - IPOTEZA ro-RO: notele „7.50” raman TEXT dupa lipire (dovada .NET in u4_cultura.txt)
  ex1_catalog15_corect.xlsx - Ex. 1 cu 15 elevi x 5 materii (C:G), Media in H =AVERAGE(C4:G4), Status in I
  ex1_dupa_rezolvare.xlsx   - Ex. 1 cum scrie rezolvarea pliata: Media =AVERAGE(B4:F4) (B = numele, G = materia 5 lipseste)
  ex2_buletin.xlsx          - Ex. 2: rand media clasei/max/min, Column pe mediile pe materii, Landscape, Narrow, antet, 1 pagina
  ex2_rand15_intreg.xlsx    - Ex. 2 capcana: elevul selecteaza tot randul 15 (C15:G15) -> apare si bara „Media”
Note inventate pentru ex. 1: generate determinist, nume „Elev 01..15” (fara date personale).
Iesire: u1_iesire.json; tipareste <= 20 de randuri."""
import html
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

L = Path(__file__).resolve().parent
P = L / "produs_elev"
P.mkdir(exist_ok=True)
SRC = Path("C:/00/Projects/LearningHub/content/tic/cls8/m1-excel-fundamente/lectia6-proiect.html")
s = SRC.read_text(encoding="utf-8")
blocuri = [html.unescape(re.sub(r"<[^>]+>", "", b)) for b in re.findall(r'<div class="code-block"[^>]*>(.*?)</div>', s, re.S)]
out = {"blocuri_copiabile": len(blocuri), "blocuri_cu_tab": [("\t" in b) for b in blocuri]}
linii = [ln.split("\t") for ln in blocuri[0].strip().split("\n")]
out["antet_bloc"] = linii[0]
out["coloane_pe_rand_date"] = sorted({len(r) for r in linii[1:]})

# separatorul in formulele din HTML: brut vs dupa html.unescape (inventarul B §4b numara ';' in brut)
code_raw = re.findall(r"<code[^>]*>(=[^<]*)</code>", s)
out["formule_code_brut"] = len(code_raw)
out["formule_cu_punct_virgula_brut"] = sum(";" in c for c in code_raw)
out["formule_cu_punct_virgula_dupa_unescape"] = sum(";" in html.unescape(c) for c in code_raw)
out["formule_cu_virgula_dupa_unescape"] = sum("," in html.unescape(c) for c in code_raw)
out["entitati_in_formulele_cu_punct_virgula"] = sorted({m for c in code_raw if ";" in c for m in re.findall(r"&\w+;", c)})

thin = Side(style="thin")
ALL = Border(left=thin, right=thin, top=thin, bottom=thin)
RED = PatternFill("solid", fgColor="8B0000")


def drag(ws, src, targets):
    for t in targets:
        ws[t] = Translator(ws[src].value, origin=src).translate_formula(t)


def page(ws, header):
    ws.page_setup.orientation = "landscape"
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.75, bottom=0.75, header=0.3, footer=0.3)  # Narrow (valorile Excel)
    ws.oddHeader.center.text = header


def catalog(conv, name):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catalog"
    ws["A1"] = "CATALOG CLASA a VIII-a B - Anul Scolar 2024-2025"  # pas 1, literal
    ws.merge_cells("A1:G1")
    ws["A1"].alignment = Alignment(horizontal="center")
    for j, v in enumerate(linii[0]):  # pas 2: lipit in A3
        ws.cell(3, 1 + j, v)
    for i, r in enumerate(linii[1:]):
        ws.cell(4 + i, 1, int(r[0]))
        ws.cell(4 + i, 2, r[1])
        for j, v in enumerate(r[2:]):
            ws.cell(4 + i, 3 + j, conv(v))
    ws["G4"] = "=AVERAGE(C4:F4)"  # pas 3
    drag(ws, "G4", [f"G{r}" for r in range(5, 14)])
    ws["A15"], ws["A16"], ws["A17"] = "Media pe materie:", "Nota minima:", "Nota maxima:"
    ws["C15"], ws["C16"], ws["C17"] = "=AVERAGE(C4:C13)", "=MIN(C4:C13)", "=MAX(C4:C13)"
    for r in (15, 16, 17):
        drag(ws, f"C{r}", [f"{c}{r}" for c in "DEFG"])
    for c in ws["A3:G3"][0]:  # pas 4
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = RED
        c.alignment = Alignment(horizontal="center")
    for row in ws["C4:F13"]:
        for c in row:
            c.number_format = "0.00"
    for row in ws["A3:G13"]:
        for c in row:
            c.border = ALL
    ch = BarChart()  # pas 5: B3:B13 + Ctrl G3:G13, Column 2-D Clustered
    ch.type, ch.grouping = "col", "clustered"
    ch.title = "Mediile elevilor"
    ch.add_data(Reference(ws, min_col=7, min_row=3, max_row=13), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=4, max_row=13))
    ch.width, ch.height = 16, 7.5
    ws.add_chart(ch, "I3")
    page(ws, "Catalog Scolar - Clasa a VIII-a")  # pas 6
    ws["A18"], ws["C18"] = "Elevi cu media >= 7:", '=COUNTIF(G4:G13,">=7")'  # BONUS
    ws.column_dimensions["B"].width = 20
    f = P / name
    wb.save(f)
    return f, ws


f1, ws1 = catalog(float, "catalog_incearca.xlsx")
f2, _ = catalog(str, "lipire_roRO_text.xlsx")
note = {r[1]: [float(v) for v in r[2:]] for r in linii[1:]}
out["medii_python"] = {k: sum(v) / len(v) for k, v in note.items()}
out["formule_catalog"] = {c: ws1[c].value for c in ("G4", "G5", "G13", "C15", "D15", "G15", "G16", "G17", "C18")}

# Ex. 1: 15 elevi x 5 materii, note inventate deterministe (4..10, pas 0.5)
mat = ["Romana", "Matematica", "Engleza", "Istorie", "Fizica"]


def nota(i, j):
    return 4 + ((i * 7 + j * 5 + i * j) % 13) * 0.5


def ex1(name, media_formula, media_col):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Catalog - exercitiul 1 (note inventate)"
    hdr = ["Nr.", "Nume elev"] + mat
    for j, v in enumerate(hdr):
        ws.cell(3, 1 + j, v)
    for i in range(15):
        ws.cell(4 + i, 1, i + 1)
        ws.cell(4 + i, 2, f"Elev {i + 1:02d}")
        for j in range(5):
            ws.cell(4 + i, 3 + j, nota(i, j))
    mc = media_col
    ws[f"{mc}3"] = "Media"
    ws[f"{mc}4"] = media_formula
    drag(ws, f"{mc}4", [f"{mc}{r}" for r in range(5, 19)])
    sc = chr(ord(mc) + 1)
    ws[f"{sc}3"] = "Status"
    ws[f"{sc}4"] = f'=IF({mc}4>=5,"Promovat","Nepromovat")'
    drag(ws, f"{sc}4", [f"{sc}{r}" for r in range(5, 19)])
    ws["A20"], ws["C20"] = "Media pe materie:", "=AVERAGE(C4:C18)"
    ws["A21"], ws["C21"] = "Minim:", "=MIN(C4:C18)"
    ws["A22"], ws["C22"] = "Maxim:", "=MAX(C4:C18)"
    for r in (20, 21, 22):
        drag(ws, f"C{r}", [f"{c}{r}" for c in "DEFG"])
    ch = BarChart()
    ch.type = "bar"
    ch.title = "Mediile pe materii"
    ch.add_data(Reference(ws, min_col=3, max_col=7, min_row=20), from_rows=True, titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=3, max_col=7, min_row=3))
    ws.add_chart(ch, "K3")
    f = P / name
    wb.save(f)
    return f


ex1("ex1_catalog15_corect.xlsx", "=AVERAGE(C4:G4)", "H")
ex1("ex1_dupa_rezolvare.xlsx", "=AVERAGE(B4:F4)", "H")
out["ex1_python_medii_corecte_primele3"] = [sum(nota(i, j) for j in range(5)) / 5 for i in range(3)]
out["ex1_python_medii_fara_materia5_primele3"] = [sum(nota(i, j) for j in range(4)) / 4 for i in range(3)]
out["ex1_rezolvare_literal"] = "A1:F1 = Nr., Nume elev, Materia 1..5 (7 titluri in 6 celule) + G=Media => G4 ar suprascrie Materia 5"


def ex2(name, max_col):
    wb = openpyxl.load_workbook(f1)
    ws = wb.active
    ws._charts = []
    ws["A20"], ws["C20"] = "Media generala a clasei:", "=AVERAGE(G4:G13)"
    ws["A21"], ws["C21"] = "Nota maxima din clasa:", "=MAX(G4:G13)"
    ws["A22"], ws["C22"] = "Nota minima:", "=MIN(G4:G13)"
    ws["A23"], ws["C23"] = "Materia cu media cea mai mare:", "(citita de pe randul 15)"
    ch = BarChart()
    ch.type = "col"
    ch.title = "Mediile pe materii"
    ch.add_data(Reference(ws, min_col=3, max_col=max_col, min_row=15), from_rows=True, titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=3, max_col=max_col, min_row=3))
    ch.width, ch.height = 14, 7
    ws.add_chart(ch, "I3")
    page(ws, "Buletinul Clasei a VIII-a")
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    f = P / name
    wb.save(f)


ex2("ex2_buletin.xlsx", 6)
ex2("ex2_rand15_intreg.xlsx", 7)
out["fisiere"] = sorted(p.name for p in P.glob("*.xlsx"))
(L / "u1_iesire.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
for k in ("blocuri_copiabile", "blocuri_cu_tab", "coloane_pe_rand_date", "formule_code_brut", "formule_cu_punct_virgula_brut",
          "formule_cu_punct_virgula_dupa_unescape", "formule_cu_virgula_dupa_unescape", "entitati_in_formulele_cu_punct_virgula",
          "ex1_python_medii_corecte_primele3", "ex1_python_medii_fara_materia5_primele3", "fisiere"):
    print(k, "=", out[k])
