"""U1 - fac sarcinile lectiei lectia2-date (cls8) ca elevul, in fisiere xlsx (openpyxl), apoi LibreOffice recalculeaza (H_randeaza).
Fisiere in produs_elev/:
  incearca_catalog.xlsx  - Incearca: antet A1:D1, 5 elevi, bold+fundal albastru+text alb, All Borders, centrare C2:D6, 2 zecimale pe D.
                           foaia Bonus_Merge: aceeasi foaie DUPA bonusul „imbina A1:D1 si scrie titlul" (antetul B1:D1 dispare).
  ex1_catalog.xlsx       - Ex1 tastat corect + data 15/06/2026 in E2 (ca data) + foaia Lipire: blocul cu | lipit pe Tab (o coloana).
  ex2_factura.xlsx       - Ex2: factura cu numere + formule =C4*D4 trase in jos + =SUM(E4:E8); foaia Ca_in_model: valorile scrise ca in
                           „Structura exemplu" ("5 buc", "3.50 lei") -> text.
  atomi_verificari.xlsx  - afirmatiile atomilor: 46037 = 15/01/2026, 0.85 si 85 cu format procent, telefonul 0741234567.
Iesire: u1_iesire.json; tipareste <= 20 de randuri."""
import datetime as dt
import html
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

L = Path(__file__).resolve().parent
P = L / "produs_elev"
P.mkdir(exist_ok=True)
SRC = Path(r"C:\00\Projects\LearningHub\content\tic\cls8\m1-excel-fundamente\lectia2-date.html")
s = SRC.read_text(encoding="utf-8")
out = {}

subtire = Side(style="thin")
ALL = Border(left=subtire, right=subtire, top=subtire, bottom=subtire)
ALBASTRU = PatternFill("solid", fgColor="0070C0")


def catalog(ws):
    rows = [("Nr. Crt.", "Nume Elev", "Nota", "Media"), (1, "Andrei Pop", 8, 8.5), (2, "Maria Ion", 9, 9.25),
            (3, "Dan Popa", 7, 7), (4, "Ana Lupu", 10, 9.75), (5, "Ioana Rusu", 6, 6.5)]
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    for c in range(1, 5):
        h = ws.cell(1, c)
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = ALBASTRU
    for row in ws["A1:D6"]:
        for c in row:
            c.border = ALL
    for row in ws["C2:D6"]:
        for c in row:
            c.alignment = Alignment(horizontal="center")
    for row in ws["B2:B6"]:
        for c in row:
            c.alignment = Alignment(horizontal="left")
    for row in ws["D2:D6"]:
        for c in row:
            c.number_format = "0.00"
    ws["F1"] = "antete_nevide"
    ws["F2"] = "=COUNTA(A1:D1)"
    ws["G1"] = "note_numerice"
    ws["G2"] = "=COUNT(C2:D6)"


# --- Incearca + bonus
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Catalog"
catalog(ws)
wb2 = wb.create_sheet("Bonus_Merge")
catalog(wb2)
wb2.merge_cells("A1:D1")          # „Selecteaza celulele A1:D1 si imbina-le folosind Home -> Merge & Center"
wb2["A1"] = "CATALOG NOTE - CLASA a VIII-a"
wb2["A1"].alignment = Alignment(horizontal="center")
wb.save(P / "incearca_catalog.xlsx")
out["incearca"] = {"bonus_B1_dupa_merge": wb2["B1"].value, "bonus_C1": wb2["C1"].value, "bonus_D1": wb2["D1"].value}

# --- Ex1
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Catalog"
bloc = html.unescape(re.search(r'<div class="code-example">\s*<code>(Nr\. \|.*?)</code>', s, re.S).group(1))
date = [[x.strip() for x in ln.split("|")] for ln in bloc.split("\n")]
for r, row in enumerate(date, 1):
    for c, v in enumerate(row, 1):
        ws.cell(r, c, int(v) if v.isdigit() else v)
for c in range(1, 5):
    ws.cell(1, c).font = Font(bold=True)
    ws.cell(1, c).fill = ALBASTRU
for row in ws["A1:D5"]:
    for c in row:
        c.border = ALL
for row in ws["C2:D5"]:
    for c in row:
        c.alignment = Alignment(horizontal="center")
ws["E1"] = "Data tezei"
ws["E2"] = dt.date(2026, 6, 15)
ws["E2"].number_format = "DD/MM/YYYY"
ws["G1"] = "note_numerice"
ws["G2"] = "=COUNT(C2:D5)"
ws["H1"] = "serial_data"
ws["H2"] = "=E2*1"
lip = wb.create_sheet("Lipire")
for r, ln in enumerate(bloc.split("\n"), 1):     # textul copiat nu are Tab -> fiecare rand intr-o singura celula
    for c, v in enumerate(ln.split("\t"), 1):
        lip.cell(r, c, v)
lip["G1"] = "numere_dupa_lipire"
lip["G2"] = "=COUNT(A1:F5)"
lip["H1"] = "celule_ocupate_B_la_D"
lip["H2"] = "=COUNTA(B1:D5)"
wb.save(P / "ex1_catalog.xlsx")
out["ex1"] = {"bloc_contine_tab": "\t" in bloc, "linii_bloc": len(bloc.split("\n")), "A1_dupa_lipire": lip["A1"].value}

# --- Ex2
prod = [("Caiet A4", 5, 3.5), ("Pix", 10, 1.2), ("Radiera", 8, 1.0), ("Creion", 12, 0.8), ("Rigla", 6, 2.5)]
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Factura"
ws.merge_cells("A1:E1")
ws["A1"] = "FACTURA NR. 001"
ws["A1"].font = Font(bold=True, size=16)
ws["A1"].alignment = Alignment(horizontal="center")
for c, h in enumerate(["Nr.", "Produs", "Cantitate", "Pret unitar", "Total"], 1):
    ws.cell(3, c, h).font = Font(bold=True)
    ws.cell(3, c).fill = PatternFill("solid", fgColor="BFBFBF")
for i, (p, q, pu) in enumerate(prod):
    r = 4 + i
    ws.cell(r, 1, i + 1)
    ws.cell(r, 2, p)
    ws.cell(r, 3, q)
    ws.cell(r, 4, pu)
ws["E4"] = "=C4*D4"
for r in range(5, 9):                              # tras in jos, ca in Excel
    ws[f"E{r}"] = Translator(ws["E4"].value, origin="E4").translate_formula(f"E{r}")
ws["A9"] = "TOTAL GENERAL"
ws["E9"] = "=SUM(E4:E8)"
for c in range(1, 6):
    ws.cell(9, c).font = Font(bold=True)
    ws.cell(9, c).fill = PatternFill("solid", fgColor="FFFF00")
for row in ws["D4:E9"]:
    for c in row:
        c.number_format = '#,##0.00 "lei"'
for row in ws["A3:E9"]:
    for c in row:
        c.border = ALL
ws["G3"] = "total_exemplu_2_randuri"
ws["G4"] = "=E4+E5"
m = wb.create_sheet("Ca_in_model")
m.append(["Nr.", "Produs", "Cant.", "Pret", "Total"])
m.append([1, "Caiet A4", "5 buc", "3.50 lei", "17.50 lei"])     # „1 Caiet A4, 5 buc, 3.50 lei, Total 17.50 lei" (rezolvarea)
m.append([2, "Pix", "10 buc", "1.20 lei", "12.00 lei"])
m["F1"] = "produs_formula"
m["F2"] = "=C2*D2"
m["G1"] = "suma_totaluri"
m["G2"] = "=SUM(E2:E3)"
m["H1"] = "numere_in_tabel"
m["H2"] = "=COUNT(C2:E3)"
wb.save(P / "ex2_factura.xlsx")
out["ex2"] = {"formula_E5": ws["E5"].value, "formula_E8": ws["E8"].value}

# --- afirmatiile atomilor
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Atomi"
ws["A1"] = "=DATE(2026,1,15)"
ws["B1"] = "serial calculat de LibreOffice pentru 15/01/2026 (lectia: 46037)"
ws["A2"] = 0.85
ws["A2"].number_format = "0%"
ws["A3"] = 85
ws["A3"].number_format = "0%"
ws["A4"] = 741234567                      # 0741234567 tastat intr-o celula General -> numar
ws["A5"] = "0741234567"                   # tastat cu apostrof / celula Text
ws["A6"] = "=LEN(A4)"
ws["A7"] = "=LEN(A5)"
ws["A8"] = 1234.5
ws["A8"].number_format = "#,##0.00"
ws["A9"] = "=A2*100"
ws["A10"] = "=A1*1"
wb.save(P / "atomi_verificari.xlsx")
out["python_serial_15_01_2026"] = (dt.date(2026, 1, 15) - dt.date(1899, 12, 30)).days
out["python_suma_rezolvare_ex2"] = round(sum(q * pu for _, q, pu in prod), 2)
out["python_suma_primele_2_randuri"] = round(5 * 3.5 + 10 * 1.2, 2)
(L / "u1_iesire.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps(out, ensure_ascii=False)[:1500])
