"""U7 - redeschid artefactele intr-un proces Python NOU: formula din fisier + valoarea recalculata de LibreOffice,
plus textul AFISAT din PDF-ul randat de LibreOffice (ce vede ochiul: 85%, 8500%, 17,50 lei). Iesire: 07_redeschis.json."""
import json
from pathlib import Path

import fitz
import openpyxl

L = Path(__file__).resolve().parent
r = {}
for nume, cel in [("incearca_catalog.xlsx", ["Catalog!B1", "Catalog!F2", "Catalog!G2", "Bonus_Merge!A1", "Bonus_Merge!B1", "Bonus_Merge!F2"]),
                  ("ex1_catalog.xlsx", ["Catalog!E2", "Catalog!G2", "Catalog!H2", "Lipire!A2", "Lipire!G2", "Lipire!H2"]),
                  ("ex2_factura.xlsx", ["Factura!E4", "Factura!E5", "Factura!E6", "Factura!E9", "Factura!G4",
                                       "Ca_in_model!F2", "Ca_in_model!G2", "Ca_in_model!H2"]),
                  ("atomi_verificari.xlsx", ["Atomi!A1", "Atomi!A2", "Atomi!A3", "Atomi!A6", "Atomi!A7", "Atomi!A9", "Atomi!A10"])]:
    f = openpyxl.load_workbook(L / "produs_elev" / nume)
    v = openpyxl.load_workbook(L / "recalculat" / nume, data_only=True)
    r[nume] = {"foi": f.sheetnames}
    for c in cel:
        ws, a = c.split("!")
        r[nume][c] = {"in_fisier": f[ws][a].value, "valoare_recalculata_LibreOffice": v[ws][a].value}
pdfs = {}
for p in sorted(L.rglob("*.pdf")):
    d = fitz.open(p)
    pdfs[p.relative_to(L).as_posix()] = " | ".join(pg.get_text().replace("\n", " ; ") for pg in d)[:900]
r["text_afisat_in_pdf_LibreOffice"] = pdfs
(L / "07_redeschis.json").write_text(json.dumps(r, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
for k, x in r.items():
    print(k, json.dumps(x, ensure_ascii=False, default=str)[:420])
