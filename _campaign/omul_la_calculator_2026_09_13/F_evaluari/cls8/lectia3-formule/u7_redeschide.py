"""U7 - redeschid artefactele intr-un proces Python NOU: formula din fisier + valoarea recalculata de LibreOffice,
plus textul AFISAT din PDF-ul randat de LibreOffice (ce vede ochiul pe Windows ro-RO). Iesire: 07_redeschis.json."""
import json
from pathlib import Path

import fitz
import openpyxl

L = Path(__file__).resolve().parent
r = {}
for nume, cel in [("incearca_catalog.xlsx", ["Catalog!E2", "Catalog!E3", "Catalog!E4", "Catalog!F2", "Catalog!F3", "Catalog!F4",
                                            "Catalog!E5", "Catalog!H2", "Bonus!E3", "Bonus!F3", "Bonus!E5"]),
                  ("ex1_medie.xlsx", ["Medie!F2", "Medie!F3", "Medie!F6", "Medie!G2", "Medie!G3", "Medie!G6", "Medie!I2"]),
                  ("ex2_excursie.xlsx", ["Excursie!C3", "Excursie!C4", "Excursie!C5", "Excursie!C8", "Elevi25!C8",
                                         "Fara_dolar!C4", "Fara_dolar!C5", "Fara_dolar!C8"]),
                  ("atomi_verificari.xlsx", ["Atomi!A1", "Atomi!A2", "Atomi!A3", "Atomi!B4", "Atomi!B5", "Atomi!B6",
                                             "Tabel_final!B5", "Tabel_final!C5", "Tabel_final!D5", "Tabel_final!E5", "Tabel_final!F3",
                                             "TVA!C2", "TVA!C5", "TVA21!C2", "TVA21!C5",
                                             "Intrebare_atom7!C3", "Intrebare_atom7!C4", "Provocare!C3", "Provocare!D3"])]:
    f = openpyxl.load_workbook(L / "produs_elev" / nume)
    v = openpyxl.load_workbook(L / "recalculat" / nume, data_only=True)
    r[nume] = {"foi": f.sheetnames}
    for c in cel:
        ws, a = c.split("!")
        r[nume][c] = {"in_fisier": f[ws][a].value, "valoare_recalculata_LibreOffice": v[ws][a].value}
pdfs = {}
for p in sorted((L / "randat_xlsx").glob("*.pdf")):
    d = fitz.open(p)
    pdfs[p.name] = " | ".join(pg.get_text().replace("\n", " ; ") for pg in d)[:700]
r["text_afisat_in_pdf_LibreOffice"] = pdfs
(L / "07_redeschis.json").write_text(json.dumps(r, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
for k, x in r.items():
    if k.endswith(".xlsx"):
        print(k, "; ".join(f"{c}={y['valoare_recalculata_LibreOffice']}" for c, y in x.items() if c != "foi"))
for k, t in pdfs.items():
    print(k, t[:230])
