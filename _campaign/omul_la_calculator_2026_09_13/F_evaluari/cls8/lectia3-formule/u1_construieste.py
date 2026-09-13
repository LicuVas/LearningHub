"""U1 - fac sarcinile lectiei lectia3-formule (cls8) ca elevul, in xlsx (openpyxl); LibreOffice recalculeaza apoi (H_randeaza).
Blocurile „Copiaza” se iau din HTML (html.unescape) si se lipesc pe Tab, ca in Excel.
Fisiere in produs_elev/:
  incearca_catalog.xlsx - Incearca: E2 =B2+C2+D2 tras la E4, F2 =E2/3 tras la F4, E5 =SUM(E2:E4); foaia Bonus: C3 = 10.
  ex1_medie.xlsx        - Ex1: F2 =SUM(B2:E2) tras la F6, G2 =F2/4 tras la G6.
  ex2_excursie.xlsx     - Ex2: C3 =B3*$B$1 tras la C7, C8 =SUM(C3:C7); foaia Elevi25 (B1=25); foaia Fara_dolar (=B3*B1 tras).
  atomi_verificari.xlsx - afirmatiile atomilor: ordinea operatiilor, tabelul final (randul 5), TVA 19% vs 21%, Ex3 (=A1+A2+A3/3),
                          Provocarea (E1 fara $), intrebarea atomului 7 (=B2*B1 tras).
Iesire: u1_iesire.json; tipareste <= 20 de randuri."""
import html
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.formula.translate import Translator

L = Path(__file__).resolve().parent
P = L / "produs_elev"
P.mkdir(exist_ok=True)
SRC = Path(r"C:\00\Projects\LearningHub\content\tic\cls8\m1-excel-fundamente\lectia3-formule.html")
s = SRC.read_text(encoding="utf-8")
blocuri = [html.unescape(b) for b in re.findall(r'<div class="code-block">(.*?)</div>', s, re.S)]
out = {"blocuri_copiabile": len(blocuri), "blocuri_cu_tab": [("\t" in b) for b in blocuri]}


def lipeste(ws, bloc, r0=1):
    for r, ln in enumerate(bloc.split("\n"), r0):
        for c, v in enumerate(ln.split("\t"), 1):
            v = v.strip()
            ws.cell(r, c, int(v) if v.isdigit() else v)


def trage(ws, src, dest_cells):
    for d in dest_cells:
        ws[d] = Translator(ws[src].value, origin=src).translate_formula(d)


# --- Incearca
bl_inc = next(b for b in blocuri if b.startswith("Elev\tNota1\tNota2\tNota3\tTotal"))
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Catalog"
lipeste(ws, bl_inc)
ws["E2"] = "=B2+C2+D2"
trage(ws, "E2", ["E3", "E4"])
ws["F2"] = "=E2/3"
trage(ws, "F2", ["F3", "F4"])
ws["E5"] = "=SUM(E2:E4)"
ws["H1"] = "control_E2+E3+E4"
ws["H2"] = "=E2+E3+E4"
b = wb.create_sheet("Bonus")
lipeste(b, bl_inc)
b["C3"] = 10                                      # „Schimba nota lui Ion de la 6 la 10 (celula C3)”
for c in ("E2", "E3", "E4", "F2", "F3", "F4", "E5"):
    b[c] = ws[c].value
wb.save(P / "incearca_catalog.xlsx")
out["incearca"] = {"E3": ws["E3"].value, "E4": ws["E4"].value, "F4": ws["F4"].value, "valori_python": {
    "E2": 9 + 8 + 10, "E3": 7 + 6 + 8, "E4": 10 + 10 + 9, "E5": 27 + 21 + 29, "lectia_spune_E5": 78}}

# --- Ex1
bl1 = next(b for b in blocuri if b.startswith("Elev\tNota1\tNota2\tNota3\tNota4"))
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Medie"
lipeste(ws, bl1)
ws["F2"] = "=SUM(B2:E2)"
trage(ws, "F2", [f"F{r}" for r in range(3, 7)])
ws["G2"] = "=F2/4"
trage(ws, "G2", [f"G{r}" for r in range(3, 7)])
ws["I1"] = "note_numerice_dupa_lipire"
ws["I2"] = "=COUNT(B2:E6)"
wb.save(P / "ex1_medie.xlsx")
out["ex1"] = {"F3": ws["F3"].value, "G6": ws["G6"].value}

# --- Ex2
bl2 = next(b for b in blocuri if b.startswith("Nr elevi"))
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Excursie"
lipeste(ws, bl2)
ws["C3"] = "=B3*$B$1"
trage(ws, "C3", [f"C{r}" for r in range(4, 8)])
ws["C8"] = "=SUM(C3:C7)"
e25 = wb.create_sheet("Elevi25")
lipeste(e25, bl2)
e25["B1"] = 25
for r in range(3, 9):
    e25[f"C{r}"] = ws[f"C{r}"].value
fd = wb.create_sheet("Fara_dolar")               # rezolvarea spune: fara $ referinta ar merge pe „celule goale”
lipeste(fd, bl2)
fd["C3"] = "=B3*B1"
trage(fd, "C3", [f"C{r}" for r in range(4, 8)])
fd["C8"] = "=SUM(C3:C7)"
wb.save(P / "ex2_excursie.xlsx")
out["ex2"] = {"C5": ws["C5"].value, "Fara_dolar_C4": fd["C4"].value, "Fara_dolar_C5": fd["C5"].value,
              "B2_continut": ws["B2"].value, "python_total_28": 28 * (45 + 120 + 35 + 15 + 20),
              "python_total_25": 25 * (45 + 120 + 35 + 15 + 20)}

# --- afirmatiile atomilor
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Atomi"
ws["A1"] = "=(2+3)*4"
ws["A2"] = "=2+3*4"
ws["A3"] = "=10-3+2"
ws["A4"] = 9
ws["A5"] = 7
ws["A6"] = 8
ws["B4"] = "=A4+A5+A6/3"                          # Ex3: formula colegului
ws["B5"] = "=(A4+A5+A6)/3"
ws["B6"] = "=AVERAGE(A4:A6)"
t = wb.create_sheet("Tabel_final")               # atomul 6: catalogul complet
for r, row in enumerate([("Elev", "Nota1", "Nota2", "Nota3", "Total", "Media"), ("Ana", 9, 8, 10), ("Ion", 7, 6, 8),
                         ("Maria", 10, 10, 9), ("Total",)], 1):
    for c, v in enumerate(row, 1):
        t.cell(r, c, v)
t["E2"] = "=SUM(B2:D2)"
trage(t, "E2", ["E3", "E4"])
t["F2"] = "=E2/3"
trage(t, "F2", ["F3", "F4"])
t["B5"] = "=SUM(B2:B4)"
trage(t, "B5", ["C5", "D5", "E5"])
for r in range(2, 5):
    t[f"F{r}"].number_format = "0.00"
v = wb.create_sheet("TVA")                       # atomul 7: 19% din lectie; foaia TVA21 cu cota in vigoare
for sh, cota in ((v, 0.19), (wb.create_sheet("TVA21"), 0.21)):
    sh["A1"], sh["B1"], sh["C1"] = "TVA", cota, "Formula TVA"
    sh["B1"].number_format = "0%"
    for r, (p, pr) in enumerate([("Caiet", 5), ("Pix", 3), ("Ruler", 8)], 2):
        sh.cell(r, 1, p)
        sh.cell(r, 2, pr)
    sh["C2"] = "=B2*$B$1"
    trage(sh, "C2", ["C3", "C4"])
    sh["C5"] = "=SUM(C2:C4)"
q = wb.create_sheet("Intrebare_atom7")           # =B2*B1 tras in jos, pe tabelul TVA din atom
for c in ("A1", "B1", "A2", "B2", "A3", "B3", "A4", "B4"):
    q[c] = v[c].value
q["C2"] = "=B2*B1"
trage(q, "C2", ["C3", "C4"])
pv = wb.create_sheet("Provocare")                # 5 preturi in B2:B6, reducere 15% in E1
pv["E1"] = 0.15
for r, pr in enumerate([100, 50, 20, 80, 40], 2):
    pv[f"B{r}"] = pr
pv["C2"] = "=B2-B2*E1"
trage(pv, "C2", [f"C{r}" for r in range(3, 7)])
pv["D2"] = "=B2-B2*$E$1"
trage(pv, "D2", [f"D{r}" for r in range(3, 7)])
wb.save(P / "atomi_verificari.xlsx")
out["atomi"] = {"Intrebare_atom7_C4": q["C4"].value, "Provocare_C3": pv["C3"].value, "Tabel_final_D5": t["D5"].value,
                "python_ex3_gresit": round(9 + 7 + 8 / 3, 2), "python_ex3_corect": (9 + 7 + 8) / 3,
                "python_tva19_caiet": round(5 * 0.19, 2), "python_tva21_caiet": round(5 * 0.21, 2)}
(L / "u1_iesire.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps(out, ensure_ascii=False)[:1800])
