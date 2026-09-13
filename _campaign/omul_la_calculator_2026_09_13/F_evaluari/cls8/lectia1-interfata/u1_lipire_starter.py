"""anexa lipirea_tabelului_starter: ce ajunge in foaie daca elevul apasa „Copiaza" la „Structura catalogului tau" si lipeste in A1.
Textul copiat = pre.copyable-code .textContent.trim() (atomic-learning.js, pattern D). Nu am Excel: modelez DOUA lipiri
 (1) text simplu impartit doar pe Tab (nu exista Tab in text -> o coloana)  (2) impartit pe spatii (ca dupa Text in coloane, delimitator Spatiu)
Ambele se compara cu tabelul cerut (A1:D6 = antete + 5 elevi). Iesire: produs_elev/lipire_starter.xlsx + u1_lipire_iesire.json."""
import html
import json
import re
from pathlib import Path

import openpyxl

L = Path(__file__).resolve().parent
s = Path(r"C:\00\Projects\LearningHub\content\tic\cls8\m1-excel-fundamente\lectia1-interfata.html").read_text(encoding="utf-8")
pre = html.unescape(re.search(r'<pre class="copyable-code">(.*?)</pre>', s, re.S).group(1)).strip()
linii = pre.split("\n")
wb = openpyxl.Workbook()
out = {"text_copiat_linii": len(linii), "contine_tab": "\t" in pre}
for nume, split in [("lipire_tab", lambda x: x.split("\t")), ("lipire_spatii", lambda x: x.split())]:
    ws = wb.create_sheet(nume)
    for r, ln in enumerate(linii, 1):
        for c, v in enumerate(split(ln), 1):
            v = v.strip()
            if v == "":
                continue
            ws.cell(r, c, int(v) if v.isdigit() else v)
    ws["K1"] = "=COUNT(A1:H8)"
    nr = sum(1 for row in ws.iter_rows(max_col=8) for c in row if isinstance(c.value, int))
    out[nume] = {"A1": ws["A1"].value, "A3": ws["A3"].value, "B3": ws["B3"].value, "D2": ws["D2"].value,
                 "dimensiune": ws.calculate_dimension(), "numere_in_foaie": nr,
                 "B2_este_9_ca_in_lectie": ws["B2"].value == 9}
del wb["Sheet"]
wb.save(L / "produs_elev" / "lipire_starter.xlsx")
(L / "u1_lipire_iesire.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps(out, ensure_ascii=False, indent=1)[:1200])
