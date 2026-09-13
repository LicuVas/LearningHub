"""U7 - redeschid artefactele intr-un proces Python NOU si citesc valorile (formule + copia recalculata de LibreOffice)."""
import json
from pathlib import Path

import openpyxl

L = Path(__file__).resolve().parent
r = {}
for nume, cel in [("incearca_catalog.xlsx", ["Catalog!A1", "Catalog!D6", "Catalog!E2", "Catalog!E3", "Catalog!B7", "Catalog!G2"]),
                  ("ex1_catalog.xlsx", ["Catalog!B2", "Catalog!C4", "Catalog!B5"]),
                  ("ex2_navigare.xlsx", ["Navigare!M25", "Test!A3", "Test!C2"]),
                  ("rezolvare_pliata_ex1.xlsx", ["Rezolvare_pliata_Ex1!E7"])]:
    f = openpyxl.load_workbook(L / "produs_elev" / nume)
    v = openpyxl.load_workbook(L / "recalculat" / nume, data_only=True)
    r[nume] = {"foi": f.sheetnames}
    for c in cel:
        ws, a = c.split("!")
        r[nume][c] = {"in_fisier": f[ws][a].value, "valoare_recalculata_LibreOffice": v[ws][a].value}
(L / "07_redeschis.json").write_text(json.dumps(r, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
for k, x in r.items():
    print(k, json.dumps(x, ensure_ascii=False, default=str)[:300])
