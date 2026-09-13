"""U1 - fac graficele cerute de lectia5-grafice (cls8) ca elevul, cu openpyxl.chart; apoi H_randeaza.py pdf/xlsx.
Datele se iau din blocurile „Copiaza” ale lectiei (citite din HTML), nu scrise de mana.
Fisiere in produs_elev/:
  incearca_column.xlsx  - Incearca pasii 1-4: A1:B6 + Column 2-D grupat, titlu „Notele clasei a VIII-a”, Data Labels
  incearca_pie.xlsx     - pasul 5: acelasi tabel ca Pie (cu procentele pe felii, cum le arata Excel la etichete procent)
  incearca_line.xlsx    - pasul 5: acelasi tabel ca Line
  incearca_bonus.xlsx   - BONUS: + Chimie/Biologie/Geografie; lectia NU da valorile -> valori alese de „elev” (marcate in foaie)
  lipire_roRO.xlsx      - IPOTEZA ro-RO: „7.2” citit ca data 07.02 (dovada de parsare in u4_cultura.txt) -> serii de ~46000
  ex1_column_bar.xlsx   - Ex. 1: Column + Bar pe aceleasi date
  ex2_buget_pie.xlsx    - Ex. 2 situatia 1 (singura cu cifre): Pie 40/25/20/15
  ex3_coleg.xlsx        - Ex. 3: graficul colegului (Line, galben, fara titlu) + corectat (Column, titlu, etichete)
  serii_provocare.xlsx  - atomul 4 (tabelul cu 3 elevi, 2 serii) Column + Line, cum cere Provocarea
  lipire_text.xlsx      - IPOTEZA 2: „7.2” ramane text -> COUNT si graficul
  stil_3d.xlsx          - atomul 8 „Stil 3D”: acelasi tabel ca Column 3-D
Iesire: u1_iesire.json; tipareste <= 20 de randuri."""
import html
import json
import re
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, BarChart3D, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties

L = Path(__file__).resolve().parent
P = L / "produs_elev"
P.mkdir(exist_ok=True)
SRC = Path(r"C:\00\Projects\LearningHub\content\tic\cls8\m1-excel-fundamente\lectia5-grafice.html")
s = SRC.read_text(encoding="utf-8")
blocuri = [html.unescape(re.sub(r"<[^>]+>", "", b)) for b in re.findall(r'<div class="code-block"[^>]*>(.*?)</div>', s, re.S)]
out = {"blocuri_copiabile": len(blocuri), "blocuri_cu_tab": [("\t" in b) for b in blocuri],
       "blocuri_identice": len(set(blocuri)) == 1}
linii = [ln.split("\t") for ln in blocuri[0].strip().split("\n")]


def tabel(ws, rows, conv=float):
    ws.append(rows[0])
    for r in rows[1:]:
        ws.append([r[0]] + [conv(v) for v in r[1:]])


def pune(ws, ch, anchor, w=15, h=7.5):
    ch.width, ch.height = w, h
    ws.add_chart(ch, anchor)


def refs(ws, n, col=2, cols=1):
    data = Reference(ws, min_col=col, max_col=col + cols - 1, min_row=1, max_row=n)
    cats = Reference(ws, min_col=1, min_row=2, max_row=n)
    return data, cats


def labels(ch, percent=False):
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = not percent
    ch.dataLabels.showPercent = percent
    ch.dataLabels.showCatName = False
    ch.dataLabels.showSerName = False
    ch.dataLabels.showLegendKey = False


def book(name, rows, builders, conv=float, extra=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = name[:30]
    tabel(ws, rows, conv)
    if extra:
        extra(ws)
    n = len(rows)
    for i, b in enumerate(builders):
        ch = b(ws, n)
        pune(ws, ch, f"A{n + 4 + i * 17}")
    ws.page_setup.fitToWidth = 1
    f = P / f"{name}.xlsx"
    wb.save(f)
    return f


def column(title=None, lab=False, style=None):
    def b(ws, n):
        ch = BarChart()
        ch.type, ch.grouping = "col", "clustered"
        d, c = refs(ws, n, cols=len([x for x in ws[1] if x.value]) - 1)
        ch.add_data(d, titles_from_data=True)
        ch.set_categories(c)
        if title:
            ch.title = title
        if lab:
            labels(ch)
        return ch
    return b


def bar(title=None, lab=False):
    def b(ws, n):
        ch = column(title, lab)(ws, n)
        ch.type = "bar"
        return ch
    return b


def line(title=None, color=None):
    def b(ws, n):
        ch = LineChart()
        d, c = refs(ws, n, cols=len([x for x in ws[1] if x.value]) - 1)
        ch.add_data(d, titles_from_data=True)
        ch.set_categories(c)
        if title:
            ch.title = title
        if color:
            for sr in ch.series:
                sr.graphicalProperties.line.solidFill = color
        for sr in ch.series:
            sr.smooth = False
        return ch
    return b


def pie(title=None):
    def b(ws, n):
        ch = PieChart()
        d, c = refs(ws, n)
        ch.add_data(d, titles_from_data=True)
        ch.set_categories(c)
        if title:
            ch.title = title
        labels(ch, percent=True)
        return ch
    return b


def col3d(title=None):
    def b(ws, n):
        ch = BarChart3D()
        ch.type = "col"
        d, c = refs(ws, n)
        ch.add_data(d, titles_from_data=True)
        ch.set_categories(c)
        ch.title = title
        return ch
    return b


T = "Notele clasei a VIII-a"
fis = []
fis.append(book("incearca_column", linii, [column(T, lab=True)]))
def felii(ws):
    # a doua cale pentru procentele de pe felii: ce inseamna „20 %” pe Pie = nota / suma notelor
    from openpyxl.formula.translate import Translator
    ws["D1"] = "Felie % (nota / suma notelor)"
    ws["D2"] = "=B2/SUM($B$2:$B$6)"
    for r in range(3, 7):
        ws[f"D{r}"] = Translator(ws["D2"].value, origin="D2").translate_formula(f"D{r}")
    ws["E1"] = "Suma notelor"
    ws["E2"] = "=SUM(B2:B6)"


fis.append(book("incearca_pie", linii, [pie(T)], extra=felii))
fis.append(book("incearca_line", linii, [line(T)]))
bonus = linii + [["Chimie", "7.9"], ["Biologie", "8.4"], ["Geografie", "7.0"]]
fis.append(book("incearca_bonus", bonus, [column(T, lab=True)],
                extra=lambda ws: ws.__setitem__("A11", "valori Chimie/Biologie/Geografie alese de elev: lectia nu le da")))


# ipoteza ro-RO: „7.2” -> data 07.02.<an curent> (an = 2026, ora se tine in noiembrie 2026)
def ca_data(v):
    zi, luna = v.split(".")
    return date(2026, int(luna), int(zi))


def extra_ro(ws):
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 2).number_format = "dd.mm.yyyy"
    ws["A8"] = "IPOTEZA: Windows ro-RO citeste 7.2 ca data 07.02 (vezi u4_cultura.txt)"


fis.append(book("lipire_roRO", linii, [column(T, lab=True)], conv=ca_data, extra=extra_ro))
# a doua ipoteza (U13): Excel pastreaza „7.2” ca TEXT (aliniat stanga), nu ca data
fis.append(book("lipire_text", linii, [column(T, lab=True)], conv=str,
                extra=lambda ws: (ws.__setitem__("A8", "IPOTEZA 2: valorile raman TEXT; COUNT(B2:B6) ="), ws.__setitem__("C8", "=COUNT(B2:B6)"), ws.__setitem__("D8", "=SUM(B2:B6)"))))
fis[-1] and None
fis.append(book("ex1_column_bar", linii, [column(T, lab=True), bar(T, lab=True)]))
buget = [["Categorie", "Procent"], ["Salarii", "40"], ["Materiale", "25"], ["Intretinere", "20"], ["Alte cheltuieli", "15"]]
fis.append(book("ex2_buget_pie", buget, [pie("Bugetul scolii")],
                extra=lambda ws: (ws.__setitem__("D1", "Total procente"), ws.__setitem__("D2", "=SUM(B2:B5)"))))
clase = [["Clasa", "Elevi"], ["8A", "28"], ["8B", "30"], ["8C", "25"], ["8D", "32"]]
fis.append(book("ex3_coleg", clase, [line(None, color="FFFF00"), column("Numarul de elevi pe clase", lab=True)]))
# atomul 4: tabelul din lectie (3 elevi) - singurul „tabel din lectie cu doua serii”
t4 = [["Elev", "Romana", "Matematica"], ["Ana", "8", "7"], ["Mihai", "6", "9"], ["Maria", "9", "8"]]
fis.append(book("serii_provocare", t4, [column("Romana si Matematica"), line("Romana si Matematica")]))
fis.append(book("stil_3d", linii, [col3d(T)]))

# ce spune lectia despre Provocare vs tabelul din atomul 4
inner = (L / "innerText.txt").read_text(encoding="utf-8")
out["provocare_cere_5_elevi"] = "pentru 5 elevi" in inner
i4 = inner.find("Romana (Seria 1")
out["atom4_elevi_in_tabel"] = [n for n in ["Ana", "Mihai", "Maria", "Ion", "Elena"] if n in inner[i4:i4 + 200]]
# procentele Pie pe medii (ce vede elevul pe felii)
vals = [float(r[1]) for r in linii[1:]]
out["pie_procente_medii"] = {r[0]: round(100 * float(r[1]) / sum(vals), 1) for r in linii[1:]}
out["pie_diferenta_max_min_puncte_procent"] = round(100 * (max(vals) - min(vals)) / sum(vals), 1)
out["anatomie_inaltimi_html"] = re.findall(r'class="chart-bar" style="height: (\d+)%', s)
out["anatomie_inaltimi_daca_axa_0_10"] = [round(v * 10) for v in [7.2, 6.8, 8.1, 6.3]]
out["anatomie_inaltimi_daca_axa_0_max"] = [round(100 * v / 8.1) for v in [7.2, 6.8, 8.1, 6.3]]
out["fisiere"] = [f.name for f in fis]
(L / "u1_iesire.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
for k, v in out.items():
    print(k, str(v)[:150])
