"""U3 - a doua cale: valorile recalculate de LibreOffice (recalculat/) vs calcul independent in Python pe datele initiale;
plus potrivirea rezolvarilor pliate cu enunturile (verificarea semnalarii K_rezolvari_nepotrivite). Scrie u3_iesire.json."""
import csv
import json
from pathlib import Path

import openpyxl

L = Path(__file__).resolve().parent
K = L.parents[2] / "K_rezolvari_nepotrivite" / "K_rezultat.csv"
rez = {}

# 1. Ex3: Valoare = Pret*Stoc dupa sortare, LibreOffice vs Python pe perechile initiale
init = {"Caiet": (5, 40), "Pix": (3, 12), "Rigla": (4, 7), "Creion": (2, 55), "Guma": (1, 7)}
wv = openpyxl.load_workbook(L / "recalculat" / "ex3_stoc_formule.xlsx", data_only=True).active
lo = {wv.cell(r, 1).value: wv.cell(r, 4).value for r in range(2, 7)}
py = {k: p * s for k, (p, s) in init.items()}
rez["ex3_valoare"] = {"libreoffice": lo, "python": py, "egal": all(abs(lo[k] - py[k]) < 1e-9 for k in py),
                      "total_lo": wv["D8"].value, "total_py": sum(py.values())}

# 2. Suma notelor: LibreOffice pe tabelul sortat corect si pe cel rupt vs Python pe datele initiale
note = [8, 6, 9, 5, 7, 10]
s_ok = openpyxl.load_workbook(L / "recalculat" / "ex1_sortat.xlsx", data_only=True).active["B9"].value
s_rupt = openpyxl.load_workbook(L / "recalculat" / "capcana_doar_coloana_B.xlsx", data_only=True).active["B9"].value
rez["suma_note"] = {"lo_sortat_corect": s_ok, "lo_rupt": s_rupt, "python": sum(note)}

# 3. Perechile Elev-Nota: tabel rupt vs initial (a doua cale care CHIAR prinde ruperea)
wr = openpyxl.load_workbook(L / "recalculat" / "capcana_doar_coloana_B.xlsx", data_only=True).active
per = {wr.cell(r, 1).value: wr.cell(r, 2).value for r in range(2, 8)}
ini = dict(zip(["Maria", "Andrei", "Elena", "Bogdan", "Carla", "Dan"], note))
rez["perechi_rupte"] = [k for k in ini if per[k] != ini[k]]

# 4. K_rezultat.csv pentru lectia mea
rez["K"] = [{k: r[k] for k in ("ex_index", "level", "overlap_pct", "shuffled_signal", "verdict", "best_match_other_idx")}
            for r in csv.DictReader(K.open(encoding="utf-8")) if r["file"].endswith("m1-excel-fundamente/lectia7-sortare.html")]
(L / "u3_iesire.json").write_text(json.dumps(rez, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps(rez, ensure_ascii=False)[:1500])
