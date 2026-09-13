"""U1 - fac sarcina elevului din lectia7-sortare, pe aceleasi date, si sortez ca Excel:
rândul întreg se mută (antetul rămâne), formulele cu referințe pe același rând se ajustează la rândul nou.
Scrie produs_elev/*.xlsx si u1_iesire.json; tipareste <= 20 de randuri.
Ipoteza declarata: sortarea Excel pe mai multe niveluri se modeleaza ca sortare dupa cheie compusa
(nivel 1, apoi nivel 2) - rezultatul nu depinde de stabilitate cand cheile nu au egalitati."""
import json
from pathlib import Path

import openpyxl
from openpyxl.formula.translate import Translator

L = Path(__file__).resolve().parent
OUT = L / "produs_elev"
OUT.mkdir(exist_ok=True)
rez = {}

ELEVI = [("Maria", 8), ("Andrei", 6), ("Elena", 9), ("Bogdan", 5), ("Carla", 7), ("Dan", 10)]
CLASE = ["8A", "8B", "8A", "8B", "8A", "8B"]  # "In C2:C7 scrie alternand: 8A, 8B, 8A, 8B, 8A, 8B"


def citeste(ws, r1, r2, c1, c2):
    return [[ws.cell(r, c).value for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]


def sorteaza_randuri(ws, r1, r2, c1, c2, cheie):
    """Ca Excel cu 'Expand the selection': mut randurile intregi; formulele se traduc la randul nou."""
    randuri = []
    for r in range(r1, r2 + 1):
        randuri.append((r, [ws.cell(r, c).value for c in range(c1, c2 + 1)]))
    randuri.sort(key=lambda x: cheie(x[1]))
    for nou, (vechi, vals) in enumerate(randuri, start=r1):
        for j, v in enumerate(vals):
            c = c1 + j
            coord_v = ws.cell(vechi, c).coordinate
            coord_n = ws.cell(nou, c).coordinate
            if isinstance(v, str) and v.startswith("="):
                v = Translator(v, origin=coord_v).translate_formula(coord_n)
            ws.cell(nou, c).value = v


# ---------- Incearca + Ex. 1: Elev / Nota ----------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Ex1"
ws["A1"], ws["B1"] = "Elev", "Nota"
for i, (n, x) in enumerate(ELEVI, start=2):
    ws.cell(i, 1, n)
    ws.cell(i, 2, x)
ws["A9"], ws["B9"] = "Suma notelor", "=SUM(B2:B7)"
ws["A10"], ws["B10"] = "Media", "=AVERAGE(B2:B7)"
initial = citeste(ws, 2, 7, 1, 2)
sorteaza_randuri(ws, 2, 7, 1, 2, lambda v: v[0])            # clic in A3, Data -> A->Z
dupa_az = [r[0] for r in citeste(ws, 2, 7, 1, 2)]
sorteaza_randuri(ws, 2, 7, 1, 2, lambda v: -v[1])           # clic in B3, Data -> Z->A
dupa_za = citeste(ws, 2, 7, 1, 2)
wb.save(OUT / "ex1_sortat.xlsx")
rez["ex1"] = {"initial": initial, "dupa_AZ_pe_Elev": dupa_az, "dupa_ZA_pe_Nota": dupa_za,
              "perechi_pastrate": sorted(map(tuple, dupa_za)) == sorted(map(tuple, initial)),
              "asteptat_lectie_AZ": ["Andrei", "Bogdan", "Carla", "Dan", "Elena", "Maria"],
              "asteptat_lectie_primul_ZA": "Dan (10)"}

# ---------- Capcana: selectez doar coloana B si aleg 'Continue with the current selection' ----------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Capcana"
ws["A1"], ws["B1"] = "Elev", "Nota"
for i, (n, x) in enumerate(ELEVI, start=2):
    ws.cell(i, 1, n)
    ws.cell(i, 2, x)
note = sorted([ws.cell(r, 2).value for r in range(2, 8)], reverse=True)
for r, x in zip(range(2, 8), note):
    ws.cell(r, 2).value = x
ws["A9"], ws["B9"] = "Suma notelor", "=SUM(B2:B7)"
capcana = citeste(ws, 2, 7, 1, 2)
wb.save(OUT / "capcana_doar_coloana_B.xlsx")
rupte = sum(1 for (n, x), (n0, x0) in zip(capcana, ELEVI) if x != x0)
rez["capcana_doar_B"] = {"dupa": capcana, "elevi_cu_nota_altcuiva": rupte,
                          "suma_notelor_neschimbata": sum(r[1] for r in capcana) == sum(x for _, x in ELEVI)}

# ---------- Ex. 2 LITERAL: C1 = Clasa, apoi "coloana A1 = Nr si in A2:A7 numerele 1-6" ----------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Ex2_literal"
ws["A1"], ws["B1"] = "Elev", "Nota"
for i, (n, x) in enumerate(ELEVI, start=2):
    ws.cell(i, 1, n)
    ws.cell(i, 2, x)
ws["C1"] = "Clasa"
for i, c in enumerate(CLASE, start=2):
    ws.cell(i, 3, c)
ws["A1"] = "Nr"                                          # scrie in A1 -> suprascrie „Elev”
for i in range(2, 8):
    ws.cell(i, 1, i - 1)                                 # suprascrie numele
sorteaza_randuri(ws, 2, 7, 1, 3, lambda v: (v[2], -v[1]))
literal = citeste(ws, 1, 7, 1, 3)
ws["B9"] = "=SUM(B2:B7)"
ws["B10"] = "=MAX(B2:B7)"
wb.save(OUT / "ex2_literal_A1_Nr.xlsx")
rez["ex2_literal"] = {"tabel": literal, "nume_ramase": sum(1 for r in literal[1:] if isinstance(r[0], str))}

# ---------- Ex. 2 cum o arata rezolvarea: Nr, Elev, Nota, Clasa (coloana inserata) ----------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Ex2_inserat"
ws["A1"], ws["B1"] = "Elev", "Nota"
for i, (n, x) in enumerate(ELEVI, start=2):
    ws.cell(i, 1, n)
    ws.cell(i, 2, x)
ws["C1"] = "Clasa"
for i, c in enumerate(CLASE, start=2):
    ws.cell(i, 3, c)
ws.insert_cols(1)                                        # Insert column, apoi A1 = Nr
ws["A1"] = "Nr"
for i in range(2, 8):
    ws.cell(i, 1, i - 1)
sorteaza_randuri(ws, 2, 7, 1, 4, lambda v: (v[3], -v[2]))   # Clasa A-Z, apoi Nota Largest->Smallest
doua_niv = citeste(ws, 2, 7, 1, 4)
ws2 = wb.copy_worksheet(ws)
ws2.title = "Ex2_inapoi_dupa_Nr"
sorteaza_randuri(ws2, 2, 7, 1, 4, lambda v: v[0])
inapoi = [r[1] for r in citeste(ws2, 2, 7, 1, 4)]
ws["C9"], ws["C10"] = "=SUM(C2:C7)", "=MAX(C2:C7)"
wb.save(OUT / "ex2_inserat_doua_niveluri.xlsx")
rez["ex2_inserat"] = {"dupa_doua_niveluri": doua_niv, "dupa_sortare_Nr": inapoi,
                      "rezolvare_lectie_8A": ["Elena 9", "Maria 8", "Carla 7"],
                      "rezolvare_lectie_8B": ["Dan 10", "Andrei 6", "Bogdan 5"]}

# ---------- Atomul 4: tabelul-rezultat afisat in lectie (innerText r. 319-345) vs datele lectiei ----------
ATOM4 = [("Dan", 10, "8A"), ("Elena", 9, "8A"), ("Maria", 8, "8A"),
         ("Andrei", 7, "8B"), ("Carla", 6, "8B"), ("Bogdan", 5, "8B")]
corect = [(r[1], r[2], r[3]) for r in doua_niv]
dif = [{"rand": i + 2, "lectie": list(a), "din_date": list(b)} for i, (a, b) in enumerate(zip(ATOM4, corect)) if a != b]
perechi_date = {n: (x, c) for (n, x), c in zip(ELEVI, CLASE)}
perechi_rupte = [n for n, x, c in ATOM4 if perechi_date[n] != (x, c)]
rez["atom4_vs_date"] = {"randuri_diferite": dif, "elevi_cu_nota_sau_clasa_schimbata": perechi_rupte}

# ---------- Ex. 3 + "De ce crezi tu": Produs / Pret / Stoc / Valoare = B*C, sortat dupa Stoc crescator ----------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Ex3_stoc"
ws.append(["Produs", "Pret", "Stoc", "Valoare"])
for p, pr, st in [("Caiet", 5, 40), ("Pix", 3, 12), ("Rigla", 4, 7), ("Creion", 2, 55), ("Guma", 1, 7)]:
    ws.append([p, pr, st, None])
for r in range(2, 7):
    ws.cell(r, 4).value = "=B2*C2" if r == 2 else Translator("=B2*C2", origin="D2").translate_formula(f"D{r}")
sorteaza_randuri(ws, 2, 6, 1, 4, lambda v: v[2])
ws["D8"] = "=SUM(D2:D6)"
ex3 = citeste(ws, 2, 6, 1, 4)
wb.save(OUT / "ex3_stoc_formule.xlsx")
rez["ex3"] = {"dupa_stoc_crescator": ex3,
              "formule_pe_randul_propriu": all(r[3] == f"=B{i}*C{i}" for i, r in enumerate(ex3, start=2)),
              "egalitate_stoc_minim": [r[0] for r in ex3 if r[2] == min(x[2] for x in ex3)]}

(L / "u1_iesire.json").write_text(json.dumps(rez, ensure_ascii=False, indent=1), encoding="utf-8")
print("Ex1 A->Z:", dupa_az)
print("Ex1 Z->A:", dupa_za[:3], "perechi pastrate:", rez["ex1"]["perechi_pastrate"])
print("Capcana doar B:", capcana, "| elevi cu nota altcuiva:", rupte)
print("Ex2 literal (A1=Nr):", literal, "| nume ramase:", rez["ex2_literal"]["nume_ramase"])
print("Ex2 inserat, doua niveluri:", doua_niv)
print("Ex2 dupa Nr:", inapoi)
print("Atom4 diferente:", dif)
print("Atom4 perechi rupte:", perechi_rupte)
print("Ex3:", ex3, "| formule pe rand:", rez["ex3"]["formule_pe_randul_propriu"], "| egalitate:", rez["ex3"]["egalitate_stoc_minim"])
