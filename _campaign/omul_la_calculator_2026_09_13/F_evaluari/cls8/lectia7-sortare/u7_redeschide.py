"""U7 - redeschid artefactele intr-un proces Python nou si citesc valorile care conteaza. Scrie 07_redeschis.json."""
import json
import os
from pathlib import Path

import openpyxl

L = Path(__file__).resolve().parent
P, R = L / "produs_elev", L / "recalculat"
o = {"pid": os.getpid()}
ws = openpyxl.load_workbook(P / "ex2_inserat_doua_niveluri.xlsx")["Ex2_inserat"]
o["ex2_A1_D4"] = [[ws.cell(r, c).value for c in range(1, 5)] for r in range(1, 5)]
o["ex2_foi"] = openpyxl.load_workbook(P / "ex2_inserat_doua_niveluri.xlsx").sheetnames
ws = openpyxl.load_workbook(P / "ex2_literal_A1_Nr.xlsx").active
o["ex2_literal_A1_A3"] = [ws["A1"].value, ws["A2"].value, ws["A3"].value]
w = openpyxl.load_workbook(P / "ex3_stoc_formule.xlsx").active
wv = openpyxl.load_workbook(R / "ex3_stoc_formule.xlsx", data_only=True).active
o["ex3_D2_formula_valoare"] = [w["D2"].value, wv["D2"].value]
o["ex3_D8_formula_valoare"] = [w["D8"].value, wv["D8"].value]
wv = openpyxl.load_workbook(R / "capcana_doar_coloana_B.xlsx", data_only=True).active
o["capcana_A2_B2_B9"] = [wv["A2"].value, wv["B2"].value, wv["B9"].value]
(L / "07_redeschis.json").write_text(json.dumps(o, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps(o, ensure_ascii=False))
