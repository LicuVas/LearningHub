"""U5 - reproducerile semnalarilor grave, din fisierele produse (nu din memorie). Scrie u5_reproducere.txt."""
import json
import re
from pathlib import Path

import openpyxl

L = Path(__file__).resolve().parent
SITE = Path("C:/00/Projects/LearningHub/content/tic/cls8/m1-excel-fundamente")
u1 = json.loads((L / "u1_iesire.json").read_text(encoding="utf-8"))
out = []

out.append("== A. Atomul 4: tabelul-rezultat din lectie vs sortarea pe datele lectiei ==")
src = (SITE / "lectia7-sortare.html").read_text(encoding="utf-8")
m = re.search(r"Rezultat: Clasa \(A-Z\) \+ Nota \(Z-A\).*?</table>", src, re.S)
rows = re.findall(r"<tr><td>(\w+)</td><td>(\d+)</td><td>(8[AB])</td></tr>", m.group(0))
out.append("Lectie (HTML): " + str(rows))
out.append("Datele lectiei (Incearca BONUS + Ex.2): Maria 8 8A, Andrei 6 8B, Elena 9 8A, Bogdan 5 8B, Carla 7 8A, Dan 10 8B")
out.append("Sortat Clasa A-Z, Nota descresc. (produs_elev/ex2_inserat_doua_niveluri.xlsx): " + str(u1["ex2_inserat"]["dupa_doua_niveluri"]))
out.append("Randuri diferite: %d din 6; elevi cu nota/clasa schimbata in tabelul lectiei: %s"
           % (len(u1["atom4_vs_date"]["randuri_diferite"]), u1["atom4_vs_date"]["elevi_cu_nota_sau_clasa_schimbata"]))
wb = openpyxl.load_workbook(L / "produs_elev" / "ex2_inserat_doua_niveluri.xlsx")
ws = wb["Ex2_inserat"]
out.append("Redeschis din xlsx, B2:D4 = " + str([[ws.cell(r, c).value for c in (2, 3, 4)] for r in (2, 3, 4)]))

out.append("")
out.append("== B. Ex. 2 citit literal: 'Adauga coloana A1 = Nr si in A2:A7 scrie numerele 1-6' ==")
wb = openpyxl.load_workbook(L / "produs_elev" / "ex2_literal_A1_Nr.xlsx")
ws = wb.active
out.append("A1:C7 dupa pasii 1-3 facuti literal: " + str([[ws.cell(r, c).value for c in (1, 2, 3)] for r in range(1, 8)]))
out.append("Nume de elevi ramase in tabel: %d (intrebarea 4 'Ce elevi sunt in primele randuri?' nu mai are raspuns)"
           % u1["ex2_literal"]["nume_ramase"])
for f in ["lectia1-interfata", "lectia2-date", "lectia3-formule", "lectia4-functii", "lectia5-grafice", "lectia6-proiect", "lectia7-sortare"]:
    t = (SITE / f"{f}.html").read_text(encoding="utf-8")
    hits = re.findall(r"(?i)insert(?:\s+sheet)?\s+columns?|inser\w*\s+(?:o\s+)?coloan\w*", t)
    out.append(f"  predat 'inserare coloana' in {f}: {len(hits)} aparitii")

out.append("")
out.append("== C. Capcana clasica (selectezi doar coloana B + 'Continue with the current selection') ==")
out.append("Dupa: " + str(u1["capcana_doar_B"]["dupa"]))
out.append("Elevi care au primit nota altcuiva: %d din 6; suma notelor neschimbata: %s -> un total NU prinde ruperea"
           % (u1["capcana_doar_B"]["elevi_cu_nota_altcuiva"], u1["capcana_doar_B"]["suma_notelor_neschimbata"]))
wv = openpyxl.load_workbook(L / "recalculat" / "capcana_doar_coloana_B.xlsx", data_only=True).active
out.append("LibreOffice a recalculat B9 (SUM) = %r" % wv["B9"].value)

out.append("")
out.append("== D. Ex. 3 intrebarea 4 + criteriul (4): 'dupa salvare doar coloana Index mai permite revenirea exacta' ==")
for n, fr in (("sort_en", None), ("undo_en", "even after you have saved"), ("undo_ro", "chiar și după ce ați salvat")):
    if fr:
        t = (L / "surse" / f"{n}.txt").read_text(encoding="utf-8")
        i = t.find(fr)
        out.append(f"{n}: …{t[max(0, i - 120): i + 160]}…")

(L / "u5_reproducere.txt").write_text("\n".join(out), encoding="utf-8")
print("\n".join(out)[:1800])
