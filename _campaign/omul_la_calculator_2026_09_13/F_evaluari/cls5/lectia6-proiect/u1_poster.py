"""U1 - posterul facut exact cum cere lectia (Exercitiul 1 + Provocarea), cu uneltele unui incepator in PowerPoint:
o singura diapozitiva (marimea implicita 16:9 - elevul nu stie sa o schimbe), casete de text + forme (Insert -> Shapes, indiciul 3),
continutul din „Model de continut” (lectia: „foloseste exact aceste informatii”), culorile din „Sfat PRO”.
Numara ce ar face elevul: casete, forme, caractere tastate, clicuri estimate (6/caseta de text, 5/forma, +4 salvare PDF).
Scrie u1_poster.pptx + u1_poster_iesire.json; tipareste <= 20 de randuri."""
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

L = Path(__file__).resolve().parent
VERDE, ALBASTRU, MOV, PORTO, NEGRU = RGBColor(0x2E, 0x9E, 0x44), RGBColor(0x1F, 0x6F, 0xD1), RGBColor(0x7B, 0x3F, 0xB5), RGBColor(0xE8, 0x7A, 0x1E), RGBColor(0x20, 0x20, 0x20)

prs = Presentation()  # 10 x 7.5 in (4:3) in python-pptx; PowerPoint nou foloseste 13.333 x 7.5 -> il setam ca implicitul din PowerPoint
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])  # necompletat
stat = {"casete_text": 0, "forme": 0, "caractere": 0, "fonturi_pt": []}


def caseta(x, y, w, h, linii, pt, culoare=NEGRU, bold=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, ln in enumerate(linii):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(pt)
        r.font.bold = bold
        r.font.color.rgb = culoare
        stat["caractere"] += len(ln) + 1
    stat["casete_text"] += 1
    stat["fonturi_pt"].append(pt)
    return tb


def forma(tip, x, y, w, h, culoare):
    sh = s.shapes.add_shape(tip, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = culoare
    sh.line.fill.background()
    stat["forme"] += 1
    return sh


# Titlu (Faza 1)
caseta(0.3, 0.1, 12.7, 0.9, ["Calculatorul Meu"], 40, PORTO, True)
# Sectiunea 1: 8 componente (Componenta - Tip - Rol), cu forme in loc de imagini (indiciul 3)
comp = [("Monitor", "Output", "afiseaza imagini si text", ALBASTRU, MSO_SHAPE.RECTANGLE),
        ("Tastatura", "Input", "introduc text si comenzi", VERDE, MSO_SHAPE.RECTANGLE),
        ("Mouse", "Input", "controlez cursorul", VERDE, MSO_SHAPE.OVAL),
        ("Procesor (CPU)", "Intern", "proceseaza toate datele", MOV, MSO_SHAPE.RECTANGLE),
        ("Memorie RAM", "Intern", "memorie rapida pentru aplicatiile deschise", MOV, MSO_SHAPE.RECTANGLE),
        ("Hard Disk", "Stocare", "stocare permanenta a fisierelor", MOV, MSO_SHAPE.RECTANGLE),
        ("Boxe", "Output", "reproduc sunete", ALBASTRU, MSO_SHAPE.OVAL),
        ("Imprimanta", "Output", "tipareste documente", ALBASTRU, MSO_SHAPE.RECTANGLE)]
caseta(0.3, 1.0, 4.4, 0.4, ["Componente"], 18, NEGRU, True)
for i, (n, t, r, c, f) in enumerate(comp):
    y = 1.45 + i * 0.72
    forma(f, 0.35, y + 0.08, 0.5, 0.45, c)
    caseta(0.95, y, 3.8, 0.7, [f"{n} - {t}", r], 12, c)
# Sectiunea 2: Input / Output
caseta(4.9, 1.0, 2.7, 0.4, ["Input / Output"], 18, NEGRU, True)
caseta(4.9, 1.45, 2.7, 1.9, ["INPUT (Intrare):", "Tastatura", "Mouse", "Scanner", "Microfon", "Camera web"], 12, VERDE)
caseta(4.9, 3.45, 2.7, 1.9, ["OUTPUT (Iesire):", "Monitor", "Imprimanta", "Boxe", "Casti", "Videoproiector"], 12, ALBASTRU)
# Sectiunea 3: Hardware vs Software
caseta(7.7, 1.0, 2.5, 0.4, ["Hard vs Soft"], 18, NEGRU, True)
caseta(7.7, 1.45, 2.5, 1.9, ["HARDWARE (il atingi):", "Monitor", "Tastatura", "Procesor", "Placa video"], 12, MOV)
caseta(7.7, 3.45, 2.5, 1.9, ["SOFTWARE (il instalezi):", "Windows", "Microsoft Word", "Google Chrome", "Jocuri video"], 12, ALBASTRU)
# Sectiunea 4: Reguli (5 ergonomie + 5 laborator)
caseta(10.3, 1.0, 2.9, 0.4, ["Reguli"], 18, NEGRU, True)
caseta(10.3, 1.45, 2.9, 2.6, ["ERGONOMIE:", "1. Spatele drept", "2. Monitorul la nivelul ochilor", "3. Pauza 2-3 min la 45-60 min",
                               "4. Lumina suficienta", "5. Distanta 50-70 cm"], 11, PORTO)
caseta(10.3, 4.15, 2.9, 2.8, ["LABORATOR:", "1. Nu mananc/beau langa PC", "2. Nu ating ecranul", "3. Inchid corect (Start - Shut Down)",
                               "4. Anunt profesorul daca se strica ceva", "5. Las locul curat"], 11, NEGRU)
# Faza 3: elemente creative (bonusul din Provocare: sageti Tastatura -> Calculator -> Monitor)
forma(MSO_SHAPE.RIGHT_ARROW, 4.95, 5.6, 0.9, 0.35, PORTO)
caseta(5.9, 5.5, 4.2, 0.6, ["Tastatura -> Calculator -> Monitor"], 12, PORTO, True)

out_f = L / "u1_poster.pptx"
prs.save(out_f)
stat["clicuri_estimate"] = stat["casete_text"] * 6 + stat["forme"] * 5 + 4
stat["font_minim_pt"] = min(stat["fonturi_pt"])
stat["fisier"] = out_f.name
stat["latime_in"], stat["inaltime_in"] = 13.333, 7.5
(L / "u1_poster_iesire.json").write_text(json.dumps(stat, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps({k: v for k, v in stat.items() if k != "fonturi_pt"}, ensure_ascii=False))
