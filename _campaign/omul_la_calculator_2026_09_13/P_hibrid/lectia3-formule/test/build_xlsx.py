"""Construieste foile de test pentru lectia3-formule (misiune, ex1, ex2 corect/gresit, ordinea operatiilor, quiz TVA)."""
from openpyxl import Workbook

OUT = r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\P_hibrid\lectia3-formule\test\lectia3_test.xlsx"
wb = Workbook()

# 1. Misiunea "Incearca" exact cum o face elevul
ws = wb.active
ws.title = "misiune"
rows = [["Elev", "Nota1", "Nota2", "Nota3", "Total", "Media"],
        ["Ana", 9, 8, 10], ["Ion", 7, 6, 8], ["Maria", 10, 10, 9]]
for r in rows:
    ws.append(r)
for i in (2, 3, 4):
    ws[f"E{i}"] = f"=B{i}+C{i}+D{i}"
    ws[f"F{i}"] = f"=E{i}/3"
ws["E5"] = "=SUM(E2:E4)"
# tabelul final: B5 tras pana la D5
for col in "BCD":
    ws[f"{col}5"] = f"=SUM({col}2:{col}4)"

# 1b. bonus: Ion 6 -> 10
wsb = wb.create_sheet("misiune_bonus")
for r in rows:
    wsb.append(r)
wsb["C3"] = 10
for i in (2, 3, 4):
    wsb[f"E{i}"] = f"=B{i}+C{i}+D{i}"
    wsb[f"F{i}"] = f"=E{i}/3"
wsb["E5"] = "=SUM(E2:E4)"

# 2. Exercitiul 1
w1 = wb.create_sheet("ex1")
for r in [["Elev", "Nota1", "Nota2", "Nota3", "Nota4", "Suma", "Media"],
          ["Ana", 8, 7, 9, 6], ["Ion", 10, 9, 8, 10], ["Maria", 6, 7, 5, 8],
          ["Radu", 9, 9, 10, 9], ["Elena", 7, 8, 6, 7]]:
    w1.append(r)
for i in range(2, 7):
    w1[f"F{i}"] = f"=SUM(B{i}:E{i})"
    w1[f"G{i}"] = f"=F{i}/4"

# 3. Exercitiul 2 corect, cu 28 si cu 25; si varianta copilului fara $
def ex2(name, n, absolut=True):
    w = wb.create_sheet(name)
    w.append(["Nr elevi", n])
    w.append(["Cheltuiala", "Pret/elev", "Cost total"])
    for c, p in [("Transport", 45), ("Cazare", 120), ("Mancare", 35), ("Bilete muzeu", 15), ("Suveniruri", 20)]:
        w.append([c, p])
    for k, i in enumerate(range(3, 8)):
        if absolut:
            w[f"C{i}"] = f"=B{i}*$B$1"
        else:
            w[f"C{i}"] = f"=B{i}*B{1 + k}"  # ce produce tragerea lui =B3*B1
    w["C8"] = "=SUM(C3:C7)"

ex2("ex2_28", 28)
ex2("ex2_25", 25)
ex2("ex2_fara_dolar", 28, absolut=False)

# 4. Ordinea operatiilor + Ex3
wo = wb.create_sheet("ordine")
wo["A1"], wo["A2"], wo["A3"] = 9, 7, 8
wo["B1"] = "=A1+A2+A3/3"
wo["B2"] = "=(A1+A2+A3)/3"
wo["B3"] = "=(2+3)*4"
wo["B4"] = "=2+3*4"
wo["B5"] = "=10-3+2"
wo["B6"] = "=10-(3+2)"

# 5. Quiz TVA: =B2*B1 tras in jos (varianta gresita din intrebare) + exemplul corect
wt = wb.create_sheet("tva")
wt.append(["TVA", 0.19])
for n, p in [("Caiet", 5), ("Pix", 3), ("Ruler", 8)]:
    wt.append([n, p])
for k, i in enumerate(range(2, 5)):
    wt[f"C{i}"] = f"=B{i}*B{1 + k}"   # gresit
    wt[f"D{i}"] = f"=B{i}*$B$1"       # corect

wb.save(OUT)
print("scris", OUT)
