"""Citeste valorile recalculate de LibreOffice."""
from openpyxl import load_workbook

P = r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\P_hibrid\lectia3-formule\test\recalculat\lectia3_test.xlsx"
wb = load_workbook(P, data_only=True)
for ws in wb.worksheets:
    print("==", ws.title)
    for row in ws.iter_rows():
        vals = [(c.coordinate, c.value) for c in row if c.value is not None]
        if vals:
            print("  ", vals)
