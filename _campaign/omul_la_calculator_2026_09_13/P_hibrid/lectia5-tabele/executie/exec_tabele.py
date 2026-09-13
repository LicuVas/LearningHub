# Executa exercitiile din lectia5-tabele cu python-docx (fara Word): Ex1, Ex2, Ex3, Provocare, conversiile.
import json
from datetime import datetime, timedelta
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\P_hibrid\lectia5-tabele\executie")
OUT.mkdir(parents=True, exist_ok=True)
log = {}


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders')
        tcPr.append(b)
    for edge, (val, sz, color) in kw.items():
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), val)
        e.set(qn('w:sz'), str(sz))
        e.set(qn('w:color'), color)
        b.append(e)


def table_borders(table, outer_sz, inner_val):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(outer_sz)); e.set(qn('w:color'), '000000')
        b.append(e)
    for edge in ('insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), inner_val); e.set(qn('w:sz'), '0'); e.set(qn('w:color'), 'auto')
        b.append(e)
    tblPr.append(b)


# ---------------- Ex1: orar 6 coloane x 8 randuri ----------------
d = Document()
t = d.add_table(rows=8, cols=6)
t.style = 'Table Grid'
hdr = ['Ora', 'Luni', 'Marti', 'Miercuri', 'Joi', 'Vineri']
start = datetime(2026, 1, 1, 8, 0)
ore = []
for i in range(7):
    a = start + timedelta(hours=i)
    ore.append(f"{a:%H:%M}-{a + timedelta(minutes=50):%H:%M}".lstrip('0'))
for j, h in enumerate(hdr):
    c = t.cell(0, j); c.text = h
    r = c.paragraphs[0].runs[0]; r.bold = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor(255, 255, 255)
    shade(c, '1F3864')
for i, o in enumerate(ore):
    t.cell(i + 1, 0).text = o
for row in t.rows:
    for c in row.cells:
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
# indiciul: pauza 10:50-11:10 intre ora 3 (10:00-10:50) si ora 4
new_tr = t.rows[3]._tr  # dupa randul 3 (ora 10:00-10:50)
import copy
pauza = copy.deepcopy(t.rows[4]._tr)
new_tr.addnext(pauza)
t2 = t  # re-read
row = t.rows[4]
m = row.cells[0].merge(row.cells[5])
for p in m.paragraphs[1:]:
    p._element.getparent().remove(p._element)
m.paragraphs[0].text = 'PAUZA 10:50-11:10'
shade(m, 'FFFF00')
p = OUT / 'Ex1_Orar_Scolar.docx'; d.save(p)
log['ex1'] = {
    'ore_generate': ore,
    'randuri_necesare': 1 + len(ore),
    'randuri_cerute': 8,
    'pauza_indiciu': '10:50-11:10',
    'ora_urmatoare_in_orar': ore[3],
    'suprapunere_pauza_cu_ora4_min': int((datetime(2026,1,1,11,10) - datetime(2026,1,1,11,0)).seconds / 60),
    'celule_de_completat_materii': 7 * 5,
    'celule_total_de_tastat': 6 + 7 + 35,
    'apasari_Tab_pentru_col1_cand_completezi_doar_orele': 6 * 7,
    'fisier': str(p),
}

# ---------------- Ex2: browsere 5x5 ----------------
d = Document()
d.sections[0].left_margin = d.sections[0].right_margin = Cm(2.54)
t = d.add_table(rows=5, cols=5); t.style = 'Table Grid'
hdr = ['Caracteristica', 'Google Chrome', 'Mozilla Firefox', 'Microsoft Edge', 'Observatii']
col1 = ['Viteza', 'Consum memorie', 'Extensii disponibile', 'Pret']
vals = [['Rapida', 'Foarte rapida', 'Rapida'], ['Mare', 'Mediu', 'Mic'], ['Foarte multe', 'Multe', 'Multe'], ['Gratuit'] * 3]
for j, h in enumerate(hdr): t.cell(0, j).text = h
for i, c1 in enumerate(col1):
    t.cell(i + 1, 0).text = c1
    for j, v in enumerate(vals[i]): t.cell(i + 1, j + 1).text = v
# pasul 7 literal: "Imbina celulele din randul de sus al coloanei Observatii"
celule_rand_sus_col_obs = [t.cell(0, 4)]
# interpretarea plauzibila: randurile 2-5 din coloana Observatii
mv = t.cell(1, 4).merge(t.cell(4, 4))
page_w = 21.0 - 2 * 2.54
col_rest = (page_w - 3.5) / 4
# semafor: clasificare
reguli = {'verde (pozitiv)': ['Rapida', 'Foarte rapida'], 'galben (mediu)': ['Mediu'], 'rosu deschis (negativ)': ['Mare']}
toate = sorted({v for r in vals for v in r})
clasificate = {v for vs in reguli.values() for v in vs}
p = OUT / 'Ex2_Comparatie_Browsere.docx'; d.save(p)
log['ex2'] = {
    'pas7_numar_celule_in_randul_de_sus_al_coloanei_Observatii': len(celule_rand_sus_col_obs),
    'pas7_se_poate_imbina_o_singura_celula': False,
    'latime_pagina_utila_cm': round(page_w, 2),
    'latime_coloane_2_5_dupa_Distribute_cm': round(col_rest, 3),
    'valori_date': toate,
    'valori_fara_culoare_in_regula_din_lectie': sorted(set(toate) - clasificate),
    'observatie': 'Mic (consum memorie Edge) = pozitiv, dar exemplul verde da doar Rapida; Gratuit/Multe/Foarte multe nu au regula',
    'fisier': str(p),
}

# ---------------- Ex3: formular 4x8 ----------------
d = Document()
t = d.add_table(rows=8, cols=4)
def merge_clean(a, b):
    m = a.merge(b)
    for pp in m.paragraphs[1:]:
        pp._element.getparent().remove(pp._element)
    return m
m = merge_clean(t.cell(0, 0), t.cell(0, 3)); m.paragraphs[0].text = 'FORMULAR DE INSCRIERE'
r = m.paragraphs[0].runs[0]; r.bold = True; r.font.size = Pt(16); r.font.color.rgb = RGBColor(255, 255, 255)
m.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; shade(m, '1F3864')
campuri = []
for ri, et in [(1, 'Nume:'), (2, 'Prenume:'), (3, 'Data nasterii:')]:
    a = merge_clean(t.cell(ri, 0), t.cell(ri, 1)); a.text = et; a.paragraphs[0].runs[0].bold = True
    campuri.append(merge_clean(t.cell(ri, 2), t.cell(ri, 3)))
t.cell(4, 0).text = 'Clasa:'; t.cell(4, 2).text = 'Scoala:'
campuri += [t.cell(4, 1), t.cell(4, 3)]
m6 = merge_clean(t.cell(5, 0), t.cell(5, 3)); m6.text = 'Informatii suplimentare:'; m6.paragraphs[0].runs[0].bold = True
m7 = merge_clean(t.cell(6, 0), t.cell(6, 3)); t.rows[6].height = Cm(3); t.rows[6].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
a = merge_clean(t.cell(7, 0), t.cell(7, 1)); a.text = 'Data:'
b = merge_clean(t.cell(7, 2), t.cell(7, 3)); b.text = 'Semnatura:'
table_borders(t, 16, 'nil')  # 2pt = sz 16 (optimi de punct)
for c in campuri + [m7]:
    cell_border(c, bottom=('single', 4, '808080'))  # 0.5pt
p = OUT / 'Ex3_Formular_Inscriere.docx'; d.save(p)
log['ex3'] = {'celule_distincte_rand8': len({id(c._tc) for c in t.rows[7].cells}),
              'problema_rand8': 'Data: si Semnatura: stau in celulele imbinate; pasul nu lasa celula goala pentru completare (spre deosebire de randurile 2-4)',
              'fisier': str(p)}

# ---------------- Varianta gresita: Merge dupa ce ai scris ----------------
d = Document(); t = d.add_table(rows=1, cols=4); t.style = 'Table Grid'
for j, v in enumerate(['Nume', 'Nota Sem.1', 'Nota Sem.2', 'Media']): t.cell(0, j).text = v
m = t.cell(0, 0).merge(t.cell(0, 3))
log['merge_dupa_text'] = {'paragrafe_in_celula_imbinata': [pp.text for pp in m.paragraphs]}
d.save(OUT / 'Gresit_merge_dupa_text.docx')

# ---------------- Conversii text -> tabel (simularea separatorului) ----------------
ex_tab = "Nume\tVarsta\tOras\nAna\t14\tBucuresti\nIon\t13\tCluj"
rows = [r.split('\t') for r in ex_tab.split('\n')]
log['conv_tab'] = {'randuri': len(rows), 'coloane': max(map(len, rows))}
email = "Popescu Ion, 8, 9, 7\nIonescu Ana, 10, 9, 10"
rows = [r.split(',') for r in email.split('\n')]
log['conv_email_commas'] = {'randuri': len(rows), 'coloane': max(map(len, rows)), 'celule': rows,
                            'exista_rand_antet_Nume_Nota1': any('Nume' in c for r in rows for c in r)}
# Provocarea: 5 randuri x 3 valori separate prin virgula
prov = ["Ana, 9, 2", "Ion, 8, 0", "Maria, 10, 1", "Dan, 7, 4", "Ioana, 9, 3"]
virgule = [s.count(',') for s in prov]
d = Document(); t = d.add_table(rows=5, cols=3); t.style = 'Table Grid'
for i, s in enumerate(prov):
    for j, v in enumerate(s.split(',')): t.cell(i, j).text = v.strip()
m = t.cell(0, 0).merge(t.cell(0, 2))
log['provocare'] = {'virgule_pe_rand_corect': virgule,
                    'intrebarea_spune': 'doua virgule in loc de una',
                    'dupa_merge_rand1_continut': [pp.text for pp in m.paragraphs],
                    'randuri_de_date_ramase_intacte': 4}
# rand cu o virgula in plus -> coloane
gresit = "Ana, 9,, 2"
log['provocare']['rand_cu_virgula_dubla_coloane'] = len(gresit.split(','))
d.save(OUT / 'Provocare_merge_rand1.docx')

(OUT / 'executie_log.json').write_text(json.dumps(log, ensure_ascii=False, indent=1), encoding='utf-8')
print(json.dumps(log, ensure_ascii=False, indent=1))
