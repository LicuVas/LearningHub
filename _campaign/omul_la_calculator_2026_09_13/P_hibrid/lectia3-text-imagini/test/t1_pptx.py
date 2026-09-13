"""Trecerea 2: fac Exercitiul 1 exact dupa 'Vezi rezolvarea' + varianta gresita a copilului.
Masor: regula 'Minim 1 imagine per slide' vs rezolvare, 6x6, marimi font, incapere pe slide 16:9."""
import json
from pathlib import Path
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Pt, Emu, Inches
from pptx.dml.color import RGBColor

D = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\P_hibrid\lectia3-text-imagini\test")
img = D / "poza_800x600.png"
im = Image.new("RGB", (800, 600), (70, 130, 180))
ImageDraw.Draw(im).ellipse((250, 150, 550, 450), fill=(255, 200, 0))
im.save(img)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
rez = {}

def titlu(s, text, size=40):
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(1.2))
    r = tb.text_frame.paragraphs[0].add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = True; r.font.name = "Calibri"; r.font.color.rgb = RGBColor(0, 51, 102)
    return tb

def lista(s, items, numerotata=False, size=24, top=1.7, width=7.0):
    tb = s.shapes.add_textbox(Inches(0.6), Inches(top), Inches(width), Inches(5))
    tf = tb.text_frame; tf.word_wrap = True
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = (f"{i+1}. " if numerotata else "\u2022 ") + t
        r.font.size = Pt(size); r.font.name = "Calibri"
    return tb

blank = prs.slide_layouts[6]
# Slide 1 - coperta (rezolvare: titlu, Ctrl+B, 40pt, culoare, O imagine, fundal colorat)
s = prs.slides.add_slide(blank); titlu(s, "Despre mine - Ana")
s.shapes.add_picture(str(img), Inches(4.5), Inches(2), width=Inches(4.3))
s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(230, 240, 255)
# Slide 2 - hobby-uri cu imagine mica langa fiecare
s = prs.slides.add_slide(blank); titlu(s, "Hobby-urile mele")
hob = ["Desenez peisaje in weekend", "Joc sah cu bunicul", "Citesc romane de aventura", "Merg cu bicicleta", "Cant la chitara"]
lista(s, hob)
for i in range(5):
    s.shapes.add_picture(str(img), Inches(8.0), Inches(1.75 + i * 0.52), height=Inches(0.45))
# Slide 3 - familie: rezolvarea NU pune imagine
s = prs.slides.add_slide(blank); titlu(s, "Familia mea")
lista(s, ["Mama - Maria", "Tata - Ion", "Sora - Ioana"])
# Slide 4 - obiective numerotate: rezolvarea NU pune imagine
s = prs.slides.add_slide(blank); titlu(s, "Obiectivele mele")
lista(s, ["Invat sa programez jocuri", "Citesc douazeci de carti", "Castig concursul de sah"], numerotata=True)
# Slide 5 - contact fictiv: forme + gradient; rezolvarea NU pune imagine
s = prs.slides.add_slide(blank); titlu(s, "Contact (fictiv)")
from pptx.enum.shapes import MSO_SHAPE
for i, t in enumerate(["elev_test@scoala.ro", "0700-000-000"]):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1 + i * 6), Inches(3), Inches(5), Inches(1.5))
    sh.text_frame.text = t; sh.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
s.background.fill.gradient()
prs.save(D / "ex1_dupa_rezolvare.pptx")

# verificari pe fisier (citit inapoi de pe disc)
p2 = Presentation(D / "ex1_dupa_rezolvare.pptx")
per = []
for n, sl in enumerate(p2.slides, 1):
    pics = sum(1 for sh in sl.shapes if sh.shape_type == 13)
    randuri, max_cuv, min_pt = 0, 0, 999
    for sh in sl.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                txt = "".join(r.text for r in p.runs).replace("\u2022", "").strip()
                if txt:
                    randuri += 1; max_cuv = max(max_cuv, len(txt.split()))
                for r in p.runs:
                    if r.font.size: min_pt = min(min_pt, r.font.size.pt)
    per.append({"slide": n, "imagini": pics, "randuri_text": randuri, "max_cuvinte_rand": max_cuv, "font_min_pt": min_pt})
rez["ex1_rezolvare"] = per
rez["slide_uri_fara_imagine"] = [x["slide"] for x in per if x["imagini"] == 0]

# 'Slide BUN' din atomul 10: titlu 40pt + 4 bullets 24pt + imagine 50% din slide
s = prs.slides.add_slide(blank); titlu(s, "Ciclul apei in natura")
lista(s, ["Apa se evapora la soare", "Vaporii formeaza norii", "Norii aduc ploaia", "Apa ajunge in rauri"], width=6.2)
s.shapes.add_picture(str(img), Inches(6.9), Inches(1.6), width=W // 2 - Inches(0.3))
# varianta gresita a copilului: trage de marginea laterala -> imagine turtita
s = prs.slides.add_slide(blank); titlu(s, "Gresit: tras de latura")
s.shapes.add_picture(str(img), Inches(1), Inches(2), width=Inches(10), height=Inches(3))
prs.save(D / "ex1_plus_bun_si_gresit.pptx")
pb = Presentation(D / "ex1_plus_bun_si_gresit.pptx")
sb = list(pb.slides)[5]
pic = [sh for sh in sb.shapes if sh.shape_type == 13][0]
rez["slide_bun_imagine_procent_latime"] = round(pic.width / W * 100, 1)
rez["slide_bun_imagine_iese_jos"] = (pic.top + pic.height) > H
g = [sh for sh in list(pb.slides)[6].shapes if sh.shape_type == 13][0]
rez["gresit_raport_aspect"] = round(g.width / g.height, 2)
rez["original_raport_aspect"] = round(800 / 600, 2)
print(json.dumps(rez, indent=1, ensure_ascii=False))
(D / "t1_rezultat.json").write_text(json.dumps(rez, indent=1, ensure_ascii=False), encoding="utf-8")
