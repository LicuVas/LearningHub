"""Ex.3 (Regulamentul clasei) construit in cele doua feluri, cu liste REALE Word (numbering.xml), ca sa aliniez rezolvarea.

A  Ex3_A_doar_Tab.docx          Numerotare + Define New Number Format (nivel 1 = I, II, III) + Tab pe sub-puncte.
                                Tab muta sub-punctul pe nivelul 2 al ACELEIASI liste; nivelul 2 ramane in formatul listei
                                numerotate (a, b, c) - exact ce spune lectia la atomul 4 si la Ex.2.
B  Ex3_B_Tab_apoi_Marcatori.docx  Ca A, apoi clic pe o litera a. (selecteaza toate elementele de pe nivelul 2 -
                                Microsoft: "you select all of the list items that are at that particular level")
                                si Marcatori (Bullets): doar nivelul 2 devine marcator, nivelul 1 ramane I, II, III.
Iesire: docx/*.docx; verificare: reciteste numbering.xml si tipareste formatul fiecarui paragraf.
"""
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Mm, Pt

S = Path(__file__).resolve().parent
OUT = S / "docx"
OUT.mkdir(exist_ok=True)

REGULI = [("Respecta-ti colegii", ["Asculta-l pe cel care vorbeste", "Cere cuvantul ridicand mana"]),
          ("Fii punctual", ["Intra in clasa inainte de sunet", "Anunta daca intarzii"]),
          ("Pastreaza curatenia", ["Nu lasa gunoaie pe banca", "Sterge tabla la sfarsitul orei", "Pastreaza-ti locul ordonat"]),
          ("Participa activ la ore", ["Pune intrebari cand nu intelegi", "Ajuta-ti colegii la exercitii"]),
          ("Ai grija de materialele scolii", ["Inchide calculatorul corect", "Pune scaunul la loc"])]


def construieste(nume, nivel2):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = s.right_margin = Cm(2.54)
    numbering = d.part.numbering_part.element
    fmt2, txt2, font2 = nivel2
    rpr2 = f'<w:rPr><w:rFonts w:ascii="{font2}" w:hAnsi="{font2}"/></w:rPr>' if font2 else ""
    an = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="901"><w:multiLevelType w:val="hybridMultilevel"/>'
        f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="upperRoman"/><w:lvlText w:val="%1."/><w:lvlJc w:val="left"/>'
        f'<w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl>'
        f'<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="{fmt2}"/><w:lvlText w:val="{txt2}"/><w:lvlJc w:val="left"/>'
        f'<w:pPr><w:ind w:left="1440" w:hanging="360"/></w:pPr>{rpr2}</w:lvl></w:abstractNum>')
    numbering.append(an)
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="901"><w:abstractNumId w:val="901"/></w:num>'))
    p = d.add_paragraph()
    r = p.add_run("Regulamentul clasei")
    r.bold, r.font.size = True, Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for regula, subs in REGULI:
        for text, ilvl in [(regula, 0)] + [(x, 1) for x in subs]:
            q = d.add_paragraph(text)
            q._p.get_or_add_pPr().insert(0, parse_xml(
                f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{ilvl}"/><w:numId w:val="901"/></w:numPr>'))
    path = OUT / nume
    d.save(str(path))
    return path


def reciteste(path):
    """Simuleaza ce afiseaza Word: contor pe nivel, format din numbering.xml."""
    d = Document(str(path))
    lvls = {}
    for lvl in d.part.numbering_part.element.iter(qn("w:lvl")):
        lvls[int(lvl.get(qn("w:ilvl")))] = (lvl.find(qn("w:numFmt")).get(qn("w:val")), lvl.find(qn("w:lvlText")).get(qn("w:val")))
    roman = ["I", "II", "III", "IV", "V", "VI"]
    cnt = [0, 0]
    rows = []
    for p in d.paragraphs:
        np_ = p._p.pPr.find(qn("w:numPr")) if p._p.pPr is not None else None
        if np_ is None:
            continue
        il = int(np_.find(qn("w:ilvl")).get(qn("w:val")))
        cnt[il] += 1
        if il == 0:
            cnt[1] = 0
        fmt, txt = lvls[il]
        if fmt == "upperRoman":
            m = roman[cnt[il] - 1] + "."
        elif fmt == "lowerLetter":
            m = chr(96 + cnt[il]) + "."
        else:
            m = txt
        rows.append(("    " * il) + f"{m} {p.text}")
    return rows


rez = {}
for nume, niv2 in [("Ex3_A_doar_Tab.docx", ("lowerLetter", "%2.", None)),
                   ("Ex3_B_Tab_apoi_Marcatori.docx", ("bullet", "•", None))]:
    pth = construieste(nume, niv2)
    rez[nume] = reciteste(pth)
    print(nume)
    for row in rez[nume][:6]:
        print("  " + row)
(S / "s_ex3_iesire.json").write_text(json.dumps(rez, ensure_ascii=False, indent=1), encoding="utf-8")
