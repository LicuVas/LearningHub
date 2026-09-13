"""Reconstruieste provocarea din lectia5-tabele in docx: varianta veche (imbini primul rand) vs noua (rand nou deasupra).
Iesire: docx/Provocare_veche.docx, docx/Provocare_noua.docx + construieste_iesire.json"""
import copy
import json
from pathlib import Path
from docx import Document

S = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls7\lectia5-tabele")
(S / "docx").mkdir(exist_ok=True)
TEXT = ["Andrei Pop, 9, 2", "Maria Stan, 10, 0", "Ioana Dinu, 8, 5", "Radu Olt, 7, 3", "Elena Voda, 9, 1"]  # elevi imaginari


def converteste(doc):
    rows = [r.split(",") for r in TEXT]  # separator Commas
    t = doc.add_table(rows=len(rows), cols=3)
    t.style = "Table Grid"
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            t.cell(i, j).text = v.strip()
    return t


def elevi(t):
    return [r.cells[0].text for r in t.rows if r.cells[0].text in [x.split(",")[0] for x in TEXT]]


rez = {}
# varianta veche: imbina primul rand si scrie titlul
d = Document(); t = converteste(d)
m = t.cell(0, 0).merge(t.cell(0, 2))
m.add_paragraph("Situatia clasei")
d.save(S / "docx" / "Provocare_veche.docx")
rez["veche"] = {"randuri": len(t.rows), "rand1": t.rows[0].cells[0].text, "elevi_curati": len(elevi(t))}

# varianta noua: insereaza un rand deasupra, imbina randul NOU, scrie titlul
d = Document(); t = converteste(d)
nou = copy.deepcopy(t.rows[0]._tr)
for tc in nou.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
    tc.text = ""
t.rows[0]._tr.addprevious(nou)  # = Insert Above
m = t.cell(0, 0).merge(t.cell(0, 2))
m.text = "Situatia clasei"
d.save(S / "docx" / "Provocare_noua.docx")
rez["noua"] = {"randuri": len(t.rows), "rand1": t.rows[0].cells[0].text, "elevi_curati": len(elevi(t))}
print(json.dumps(rez, ensure_ascii=False, indent=1))
(S / "construieste_iesire.json").write_text(json.dumps(rez, ensure_ascii=False, indent=1), encoding="utf-8")
