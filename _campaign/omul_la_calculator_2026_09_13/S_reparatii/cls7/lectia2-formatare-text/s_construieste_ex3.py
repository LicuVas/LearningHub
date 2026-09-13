"""Reconstruieste Ex.3 in varianta noua din lectie: 4 subtitluri cu stilul Titlu 1 (Heading 1),
fiecare cu un rand de text sub el (fara randuri goale), stilul modificat O DATA (Arial 16, aldin, albastru).
Iese: S\\ex3_stil_Heading1.docx; apoi H_randeaza.py pdf -> semne de carte in PDF = titluri recunoscute."""
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

S = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls7\lectia2-formatare-text")
d = Document()
h = d.styles["Heading 1"]
h.font.name = "Arial"
h.font.size = Pt(16)
h.font.bold = True
h.font.color.rgb = RGBColor(0x1F, 0x4E, 0xC8)
for s in ["Introducere", "Capitolul 1", "Capitolul 2", "Concluzii"]:
    d.add_paragraph(s, style="Heading 1")
    d.add_paragraph("Un rand de text sub subtitlu.")
out = S / "ex3_stil_Heading1.docx"
d.save(out)

d2 = Document(out)
heads = [(p.text, p.style.name) for p in d2.paragraphs if p.style.name.startswith("Heading")]
goale = sum(1 for p in d2.paragraphs if not p.text.strip())
print(json.dumps({"titluri": heads, "paragrafe_goale": goale,
                  "stil": {"font": d2.styles["Heading 1"].font.name, "pt": d2.styles["Heading 1"].font.size.pt,
                           "bold": d2.styles["Heading 1"].font.bold}}, ensure_ascii=False))
