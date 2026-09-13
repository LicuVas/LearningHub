"""Reconstructie independenta Ex.3: varianta manuala (pensula) vs stil Heading 1, fara randuri goale."""
from docx import Document
from docx.shared import Pt, RGBColor
from pathlib import Path

D = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls7\lectia2-formatare-text\verif")
T = ["Introducere", "Capitolul 1", "Capitolul 2", "Concluzii"]

# manual
d = Document()
for t in T:
    p = d.add_paragraph()
    r = p.add_run(t); r.bold = True; r.font.name = "Arial"; r.font.size = Pt(16); r.font.color.rgb = RGBColor(0, 0x70, 0xC0)
    d.add_paragraph("O propozitie de text sub subtitlu.")
d.save(D / "v_manual.docx")

# stil
d = Document()
h = d.styles["Heading 1"]
h.font.name = "Arial"; h.font.size = Pt(16); h.font.bold = True; h.font.color.rgb = RGBColor(0, 0x70, 0xC0)
for t in T:
    d.add_paragraph(t, style="Heading 1")
    d.add_paragraph("O propozitie de text sub subtitlu.")
d.save(D / "v_stil.docx")
d2 = Document(D / "v_stil.docx")
print("stiluri:", [(p.style.name, p.text) for p in d2.paragraphs])
print("goale:", sum(1 for p in d2.paragraphs if not p.text.strip()))
