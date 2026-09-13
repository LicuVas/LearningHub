"""Reconstructie independenta Ex.3: alt text, masor distanta intre liniile de baza (nu spatiul alb)."""
import subprocess, sys, json
from pathlib import Path
import pdfplumber
from docx import Document
from docx.shared import Pt

D = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls7\lectia3-paragrafe\verif\docs")
D.mkdir(parents=True, exist_ok=True)
H = r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\H_randeaza.py"
TXT = "Text scurt numarul {n} despre ora de informatica."


def build(name, goale, after):
    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    st.paragraph_format.line_spacing = 1.15
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.space_before = Pt(0)
    for n in range(1, 6):
        p = d.add_paragraph(TXT.format(n=n))
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if n < 5:
            for _ in range(goale):
                d.add_paragraph("")
    f = D / (name + ".docx")
    d.save(f)
    r = subprocess.run([sys.executable, H, "pdf", str(f), str(D)], capture_output=True, text=True)
    pdf = next(D.rglob(name + ".pdf"))
    with pdfplumber.open(pdf) as P:
        pg = P.pages[0]
        lines = sorted({round(c["bottom"], 1) for c in pg.chars})
        words = pg.extract_words()
        tops = sorted({round(w["top"], 1) for w in words})
        bots = sorted({round(w["bottom"], 1) for w in words})
    gaps = [round(b - a, 1) for a, b in zip(bots, bots[1:])]
    white = [round(t - b, 1) for b, t in zip(bots, tops[1:])]
    return {"doc": name, "baseline_gaps": gaps, "white": white, "total": round(bots[-1] - tops[0], 1)}


out = [build("v_2goale", 2, None), build("v_3goale", 3, None), build("v_after10", 0, 10),
       build("v_after55", 0, 55), build("v_after78", 0, 78)]
print(json.dumps(out, indent=1))
