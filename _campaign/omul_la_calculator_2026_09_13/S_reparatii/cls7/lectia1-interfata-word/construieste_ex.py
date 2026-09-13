"""Reconstruieste produsele Ex.1 si Ex.3 dupa enuntul reparat si verifica numele + stilurile de titlu."""
from pathlib import Path
from docx import Document

OUT = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls7\lectia1-interfata-word\produs")
OUT.mkdir(parents=True, exist_ok=True)

# Ex.1
d = Document()
d.add_paragraph("Fisa mea de observatie - Interfata Word")
for t in ["Bold (Ctrl+B) - ingroasa textul.", "Italic (Ctrl+I) - inclina textul.", "Underline (Ctrl+U) - subliniaza textul."]:
    d.add_paragraph(t)
d.save(OUT / "7A_Popescu_Interfata.docx")

# Ex.3 - primul document
text = ["Word este un procesor de texte.", "Panglica are file si grupuri.", "Bara de acces rapid se poate personaliza.",
        "Salvez cu Ctrl+S.", "Pot exporta documentul ca PDF.", "Stilurile de titlu fac documentul ordonat."]
d = Document()
for t in text:
    d.add_paragraph(t)
d.add_paragraph("Notita QAT: Print Preview - verific pagina; Save As - copie PDF; Spelling & Grammar - corectez.")
d.save(OUT / "7A_Popescu_Tema_Word.docx")

# Ex.3 - documentul structurat: Titlu 1 deasupra, subtitlu Heading 2 sub el, al doilea Titlu 1 la mijloc
d = Document()
d.add_paragraph("Ce am invatat despre Word", style="Heading 1")
d.add_paragraph("Fereastra si panglica", style="Heading 2")
for t in text[:3]:
    d.add_paragraph(t)
d.add_paragraph("Ce voi exersa", style="Heading 1")
for t in text[3:]:
    d.add_paragraph(t)
d.save(OUT / "7A_Popescu_Tema_Word_Structurata.docx")

r = Document(OUT / "7A_Popescu_Tema_Word_Structurata.docx")
print([(p.style.name, p.text) for p in r.paragraphs if p.style.name.startswith("Heading")])
