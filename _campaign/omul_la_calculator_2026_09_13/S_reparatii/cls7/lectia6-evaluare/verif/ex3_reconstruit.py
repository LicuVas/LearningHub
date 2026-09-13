import os, subprocess, sys
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image
V = os.path.dirname(os.path.abspath(__file__))
img = os.path.join(V, "poza.png")
Image.new("RGB", (400, 300), (200, 60, 60)).save(img)
d = Document()
d.add_paragraph("Titlu referat").alignment = WD_ALIGN_PARAGRAPH.CENTER
d.add_paragraph("Paragraf de continut inainte de imagine. " * 4).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(img, width=Cm(6))  # python-docx insereaza inline (In Line with Text)
cap = d.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cap.add_run("LEGENDA: imaginea arata un dreptunghi"); r.italic = True
d.add_paragraph("Text dupa legenda. " * 10)
out = os.path.join(V, "ex3_verif.docx"); d.save(out)
# tabel Ex.2: 3 coloane x 6 randuri
d2 = Document(); t = d2.add_table(rows=6, cols=3)
lectii = ["Interfata", "Formatare text", "Paragrafe", "Liste", "Tabele"]
t.cell(0, 0).text = "Lectia"
for i, l in enumerate(lectii):
    t.cell(i + 1, 0).text = l
d2.save(os.path.join(V, "ex2_verif.docx"))
d2b = Document(r"" + os.path.join(V, "ex2_verif.docx"))
tt = d2b.tables[0]
print("Ex2 randuri", len(tt.rows), "coloane", len(tt.columns), "lectii", sum(1 for rr in tt.rows[1:] if rr.cells[0].text))
H = r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\H_randeaza.py"
print(subprocess.run([sys.executable, H, "pdf", out, V], capture_output=True, text=True).stdout)
import fitz
doc = fitz.open(os.path.join(V, "randat_docx", "ex3_verif.pdf"))
pg = doc[0]
imgs = pg.get_image_info()
print("imagine bbox", [i["bbox"] for i in imgs])
for b in pg.get_text("blocks"):
    if "LEGENDA" in b[4]:
        print("legenda bbox", b[:4])
