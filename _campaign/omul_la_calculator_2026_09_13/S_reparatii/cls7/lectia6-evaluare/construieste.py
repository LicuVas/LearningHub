"""Reconstruieste documentele-model pentru reparatiile Ex.2 (tabel) si Ex.3 (imagine + legenda)."""
import json
import subprocess
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

CAMP = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13")
S = CAMP / "S_reparatii" / "cls7" / "lectia6-evaluare"
OUT = S / "construit"
IMG = CAMP / "F_evaluari" / "cls7" / "lectia6-evaluare" / "produs_elev" / "imagine_scoala_400x300.png"
rez = {}

# Ex.2: 3 coloane x 6 randuri = antet + 5 lectii
d = docx.Document()
lectii = ["Lectia 1 - Interfata Word", "Lectia 2 - Formatare text", "Lectia 3 - Paragrafe", "Lectia 4 - Liste", "Lectia 5 - Tabele"]
t = d.add_table(rows=6, cols=3)
t.style = "Table Grid"
for j, h in enumerate(["Lectia", "Ce am invatat", "Nota mea"]):
    t.cell(0, j).text = h
for i, l in enumerate(lectii, start=1):
    t.cell(i, 0).text = l
    t.cell(i, 1).text = "..."
    t.cell(i, 2).text = "4"
p = OUT / "Ex2_tabel_3x6.docx"
d.save(p)
r = docx.Document(p).tables[0]
rez["ex2"] = {"randuri": len(r.rows), "coloane": len(r.columns), "randuri_lectii": sum(1 for row in r.rows[1:] if row.cells[0].text.startswith("Lectia"))}
# cerinta veche: 3x5 -> cate lectii incap
rez["ex2_vechi_3x5_incap_lectii"] = 5 - 1

# Ex.3: imagine In Line with Text (implicita la inserare) + legenda centrata, italic, dedesubt
d = docx.Document()
d.add_paragraph("Scoala mea").alignment = WD_ALIGN_PARAGRAPH.CENTER
pi = d.add_paragraph()
pi.alignment = WD_ALIGN_PARAGRAPH.CENTER
pi.add_run().add_picture(str(IMG), width=Cm(8))
leg = d.add_paragraph()
leg.alignment = WD_ALIGN_PARAGRAPH.CENTER
rl = leg.add_run("Figura 1. Cladirea scolii")
rl.italic = True
d.add_paragraph("Text dupa imagine. " * 20)
p3 = OUT / "Ex3_imagine_inline_legenda.docx"
d.save(p3)

rr = subprocess.run([sys.executable, str(CAMP / "H_randeaza.py"), "pdf", str(p3), str(OUT)], capture_output=True, text=True)
rez["randare_exit"] = rr.returncode
import fitz
pdf = OUT / "randat_docx" / "Ex3_imagine_inline_legenda.pdf"
doc = fitz.open(pdf)
pg = doc[0]
imgs = pg.get_image_info()
img_box = imgs[0]["bbox"] if imgs else None
leg_box = None
for b in pg.get_text("blocks"):
    if "Figura 1" in b[4]:
        leg_box = b[:4]
rez["ex3"] = {"imagine_bbox": img_box, "legenda_bbox": leg_box,
              "legenda_sub_imagine": bool(img_box and leg_box and leg_box[1] >= img_box[3] - 1)}
print(json.dumps(rez, indent=1))
(S / "construit" / "rezultat.json").write_text(json.dumps(rez, indent=1), encoding="utf-8")
