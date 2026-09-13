"""Construieste catalogul din enuntul Ex. 1 (15 elevi, Nr., Nume, 5 materii) cu formulele din rezolvarea NOUA
si, pe o foaie separata, cu formula VECHE (=AVERAGE(B4:F4)), ca sa se vada diferenta dupa recalcularea LibreOffice.
Structura noua: A3:I3 antet; A=Nr., B=Nume, C:G = 5 materii, H = Media, I = Status; randuri 4-18; sumar 20-22.
"""
from pathlib import Path
import random
from openpyxl import Workbook

OUT = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls8\lectia6-proiect\catalog_ex1.xlsx")
random.seed(8)
materii = ["Romana", "Matematica", "Engleza", "Istorie", "Biologie"]
elevi = [("Elev %02d" % i, [random.randint(3, 10) for _ in materii]) for i in range(1, 16)]

wb = Workbook()
for nume_foaie, veche in (("nou", False), ("vechi", True)):
    ws = wb.active if nume_foaie == "nou" else wb.create_sheet()
    ws.title = nume_foaie
    ws["A1"] = "Catalogul unei clase imaginare"
    ws.append([])
    ws.append(["Nr.", "Nume elev"] + materii + ["Media", "Status"])  # randul 3
    for k, (n, note) in enumerate(elevi):
        r = 4 + k
        ws.append([k + 1, n] + note)
        if veche:
            ws[f"H{r}"] = f"=AVERAGE(B{r}:F{r})"
        else:
            ws[f"H{r}"] = f"=AVERAGE(C{r}:G{r})"
        ws[f"I{r}"] = f'=IF(H{r}>=5,"Promovat","Nepromovat")'
    for r, et, fn in ((20, "Media pe materie:", "AVERAGE"), (21, "Nota minima:", "MIN"), (22, "Nota maxima:", "MAX")):
        ws[f"A{r}"] = et
        for col in "CDEFG":
            ws[f"{col}{r}"] = f"={fn}({col}4:{col}18)"
wb.save(OUT)
import json
Path(OUT.with_suffix(".json")).write_text(json.dumps(elevi), encoding="utf-8")
print("scris", OUT)
