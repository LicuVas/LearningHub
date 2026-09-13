"""Citeste valorile recalculate de LibreOffice si le compara cu calculul Python."""
from pathlib import Path
import json
from statistics import mean
from openpyxl import load_workbook

D = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls8\lectia6-proiect")
elevi = json.loads((D / "catalog_ex1.json").read_text(encoding="utf-8"))
wb = load_workbook(D / "recalculat" / "catalog_ex1.xlsx", data_only=True)
for foaie in ("nou", "vechi"):
    ws = wb[foaie]
    gresite, erori = 0, []
    for k, (n, note) in enumerate(elevi):
        r = 4 + k
        h, i = ws[f"H{r}"].value, ws[f"I{r}"].value
        ok = isinstance(h, (int, float)) and abs(h - mean(note)) < 1e-9
        if not ok:
            gresite += 1
            if len(erori) < 3:
                erori.append((r, round(mean(note), 3), h))
        exp_status = "Promovat" if isinstance(h, (int, float)) and h >= 5 else "Nepromovat"
        if i != exp_status:
            erori.append(("status", r, i))
    sumar_ok = all(
        abs(ws[f"{c}{row}"].value - fn([e[1][j] for e in elevi])) < 1e-9
        for row, fn in ((20, mean), (21, min), (22, max))
        for j, c in enumerate("CDEFG"))
    print(f"[{foaie}] medii gresite: {gresite} din 15; exemple {erori}; sumar C20:G22 corect: {sumar_ok}")
