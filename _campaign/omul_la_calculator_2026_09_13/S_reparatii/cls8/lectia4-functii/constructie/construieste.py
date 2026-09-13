"""Construieste foile lectiei 4 (cls8) in starile din pasii reparati, pentru recalcul in LibreOffice (H_randeaza.py xlsx)."""
import openpyxl
from pathlib import Path

OUT = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls8\lectia4-functii\constructie")
NOTE = [7, 9, 5, 10, 8, 6, 4, 9]

wb = openpyxl.Workbook()


def foaie(nume, note):
    ws = wb.create_sheet(nume)
    ws["A1"], ws["B1"] = "Elev", "Nota"
    for i, v in enumerate(note):
        ws.cell(row=2 + i, column=2, value=v)
    for r, (et, f) in enumerate([("Nota minima", "=MIN(B2:B9)"), ("Nota maxima", "=MAX(B2:B9)"),
                                 ("Media", "=AVERAGE(B2:B9)"), ("Nr. note", "=COUNT(B2:B9)"),
                                 ("Completate", "=COUNTA(B2:B9)"), ("Suma", "=SUM(B2:B9)")], start=2):
        ws.cell(row=r, column=4, value=et)
        ws.cell(row=r, column=5, value=f)
    return ws


wb.remove(wb.active)
foaie("initial", NOTE)
# Incearca pas 6 reparat: B5 (nota 10) -> 3
n6 = NOTE[:]; n6[3] = 3
foaie("pas6_B5_3", n6)
# Bonus reparat: dupa pas 6, B9 -> "absent"
nb = n6[:]; nb[7] = "absent"
foaie("bonus_B9_absent", nb)
# Bonus vechi: dupa pas 6, B9 goala (dovada ca nu arata diferenta)
ng = n6[:]; ng[7] = None
foaie("bonus_vechi_B9_gol", ng)
# Atom 8 reparat: B8 (nota 4) -> 10
n8 = NOTE[:]; n8[6] = 10
foaie("atom8_B8_10", n8)
# Adresele: pe ce rand sta fiecare nota
ws = wb.create_sheet("adrese")
for i, v in enumerate(NOTE):
    ws.cell(row=1 + i, column=1, value=f"B{2 + i}")
    ws.cell(row=1 + i, column=2, value=v)
# IF: proba din atomul 9 (fisierul xlsx pastreaza forma engleza cu virgula, v. H_randeaza.py)
ws = wb.create_sheet("if")
ws["A1"], ws["B1"] = 4, '=IF(A1>=5,"Promovat","Corigent")'
ws["A2"], ws["B2"] = 7, '=IF(A2>=5,"Promovat","Corigent")'
ws["A3"], ws["B3"] = 8, "=IF(A3>=5,A3*2,0)"
ws["A4"], ws["B4"] = 3, "=MIN(A1,A2,A4)"
# Ex. 2 meteo
ws = wb.create_sheet("meteo")
for i, (z, t, u) in enumerate(zip(["Luni", "Marti", "Miercuri", "Joi", "Vineri", "Sambata", "Duminica"],
                                  [18, 22, 15, 25, 28, 20, 17], [65, 58, 72, 45, 40, 68, 75]), start=2):
    ws.cell(row=i, column=1, value=z); ws.cell(row=i, column=2, value=t); ws.cell(row=i, column=3, value=u)
for r, (a, b) in enumerate([("=MIN(B2:B8)", "=MIN(C2:C8)"), ("=MAX(B2:B8)", "=MAX(C2:C8)"),
                            ("=AVERAGE(B2:B8)", "=AVERAGE(C2:C8)")], start=2):
    ws.cell(row=r, column=5, value=a); ws.cell(row=r, column=6, value=b)
ws["E5"] = "=MAX(B2:B8)-MIN(B2:B8)"
ws["G4"] = "=ROUND(E4,2)"
# Atom 6 quiz: 8, 6, 10 in A1:A3, AVERAGE(A1:A5)
ws = wb.create_sheet("quiz6")
ws["A1"], ws["A2"], ws["A3"], ws["B1"] = 8, 6, 10, "=AVERAGE(A1:A5)"
wb.save(OUT / "l4_verificare.xlsx")
print("scris", OUT / "l4_verificare.xlsx")
