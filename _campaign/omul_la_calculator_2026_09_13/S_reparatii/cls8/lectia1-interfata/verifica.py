"""Verific lectia reparata contra fisierelor construite: fiecare cifra/adresa din rezolvari exista in xlsx recalculat."""
import html
import json
import re
from pathlib import Path

import openpyxl

S = Path(__file__).resolve().parent
LES = Path(r"C:\00\Projects\LearningHub\content\tic\cls8\m1-excel-fundamente\lectia1-interfata.html")
t = LES.read_text(encoding="utf-8")
R = S / "recalculat"
ok = []


def txt(s):
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", s)))


sols = [txt(m) for m in re.findall(r'<div class="practice-solution-body">(.*?)</div>\s*</details>', t, re.S)]
assert len(sols) == 3, len(sols)

# Ex1
wb = openpyxl.load_workbook(R / "8A_Popescu_Catalog.xlsx", data_only=True)
ws = wb["Catalog"]
tab = [[c.value for c in r] for r in ws.iter_rows()]
for row in tab:
    assert all(str(v) in sols[0] for v in row), row
assert "B2" in sols[0] and "C4" in sols[0] and "A1:C4" in sols[0]
assert ws.max_row == 4 and ws.max_column == 3  # Ctrl+End = C4
ok.append("ex1: tabelul A1:C4, foaia Catalog, B2/C4 si Ctrl+End C4 = xlsx recalculat")

# Ex2
wb = openpyxl.load_workbook(R / "8A_Popescu_Navigare.xlsx", data_only=True)
v = wb["Verificare"]
assert v["B1"].value == 50 and "50 de celule" in sols[1]
assert v["B2"].value == "Am ajuns aici!" and "M25" in sols[1]
assert v["B3"].value == 3
nav, test = wb["Navigare"], wb["Test"]
assert (nav.max_row, nav.max_column) == (25, 13) and (test.max_row, test.max_column) == (3, 1)
assert "A3" in sols[1] and "2 foi" in sols[1] and "Navigare" in sols[1] and "Test" in sols[1]
ok.append("ex2: 50 celule (LibreOffice ROWS*COLUMNS), Ctrl+End Navigare=M25 (rand 25, col 13), Test=A3, 2 foi")

# Ex3
wb = openpyxl.load_workbook(R / "8A_Popescu_Ex3_exemple.xlsx", data_only=True)
m = wb["Matematica"]
assert round(m["E2"].value, 2) == 8.67 and "8,67" in sols[2] and "=AVERAGE(B2:D2)" in sols[2]
assert m["H1"].value == 28 and m["H2"].value == 12 and "a 28-a coloana" in sols[2]
assert "PEMDAS" not in sols[2] and "discount" not in sols[0] and "#DIV/0!" not in sols[1]
ok.append("ex3: AVERAGE(9,7,10)=8,67 si COLUMN(AB12)=28, ROW=12 recalculate; nicio urma din rezolvarile straine")

# Blocul Copiaza
pre = html.unescape(re.search(r'<pre class="copyable-code">(.*?)</pre>', t, re.S).group(1)).strip()
assert pre == (S / "bloc_copiaza.txt").read_text(encoding="utf-8"), repr(pre[:80])
wb = openpyxl.load_workbook(R / "lipire_bloc_nou.xlsx", data_only=True)
assert wb.active["F1"].value == 15 and wb.active["B2"].value == 9
ok.append("copiaza: textul din <pre> = tabel cu Tab; lipit -> A1:D6, COUNT(B2:D6)=15")

# Chestionare
for aid, raw in re.findall(r'<div class="atom" id="(atom-\d+)" data-quiz=\'(.*?)\'>', t):
    qs = json.loads(html.unescape(raw))
    print(aid, " | ".join(q["question"][:55] for q in qs))
print("\n".join(ok))
