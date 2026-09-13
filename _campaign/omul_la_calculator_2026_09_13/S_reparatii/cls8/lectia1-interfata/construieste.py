"""S - construiesc rezolvarile corecte ale celor 3 exercitii din cls8/lectia1-interfata in .xlsx
si modelez lipirea noului bloc „Copiaza". Iesire: xlsx/*.xlsx + construieste_iesire.json.
Recalcularea (valori reale) se face apoi cu H_randeaza.py xlsx, iar verifica.py citeste recalculat/."""
import json
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

S = Path(__file__).resolve().parent
X = S / "xlsx"
X.mkdir(exist_ok=True)
out = {}


def ultima_celula(ws):
    """Ctrl+End = coltul dreapta-jos al zonei folosite (randul maxim, coloana maxima cu continut)."""
    rmax = cmax = 0
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                rmax, cmax = max(rmax, c.row), max(cmax, c.column)
    return f"{get_column_letter(cmax)}{rmax}" if rmax else "A1"


# ---------- Exercitiul 1: catalogul ghidat, exact pasii din cerinta ----------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Catalog"  # Pas 3
for r, row in enumerate([("Nume", "Nota Matematica", "Nota Romana"), ("Ion", 8, 9), ("Ana", 10, 7), ("Mihai", 6, 8)], 1):
    for c, v in enumerate(row, 1):
        ws.cell(r, c, v)
note = [c.coordinate for row in ws.iter_rows(min_row=2) for c in row if isinstance(c.value, (int, float))]
out["ex1"] = {
    "foi": wb.sheetnames,
    "antete": [ws["A1"].value, ws["B1"].value, ws["C1"].value],
    "prima_nota": note[0], "ultima_nota": note[-1],
    "ctrl_end": ultima_celula(ws),
    "numar_note": len(note),
    "tabel": [[c.value for c in row] for row in ws.iter_rows()],
}
wb.save(X / "8A_Popescu_Catalog.xlsx")

# ---------- Exercitiul 2: navigare si selectare ----------
wb = openpyxl.Workbook()
ws = wb.active
ws["M25"] = "Am ajuns aici!"          # pas 1
sel = ws["A1:E10"]                      # pas 3 (selectia nu scrie nimic in foaie)
ws.title = "Navigare"                   # pas 4a
t = wb.create_sheet()                   # pas 4b: butonul + da o foaie cu nume implicit
out_nume_implicit = t.title
t.title = "Test"                        # pas 4c: redenumire
for i, n in enumerate(["Maria", "Andrei", "Ioana"], 1):  # pas 5: nume inventate
    t[f"A{i}"] = n
out["ex2"] = {
    "foi": wb.sheetnames,
    "numar_foi": len(wb.sheetnames),
    "celule_in_A1:E10": sum(len(r) for r in sel),
    "ctrl_end_Navigare": ultima_celula(ws),
    "ctrl_end_Test": ultima_celula(t),
    "M25": ws["M25"].value,
    "nume_implicit_foaie_noua_openpyxl": out_nume_implicit,
}
# verificare in foaie separata (formule recalculate de LibreOffice)
v = wb.create_sheet("Verificare")
v["A1"] = "celule A1:E10"
v["B1"] = "=ROWS(Navigare!A1:E10)*COLUMNS(Navigare!A1:E10)"
v["A2"] = "text in M25"
v["B2"] = "=Navigare!M25"
v["A3"] = "nume in Test"
v["B3"] = "=COUNTA(Test!A1:A3)"
wb.save(X / "8A_Popescu_Navigare.xlsx")
wb.remove(v)

# ---------- Exercitiul 3: exemplele din raspunsurile-model, construite ----------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Matematica"
for r, row in enumerate([("Nume", "Nota1", "Nota2", "Nota3", "Media"), ("Ion", 9, 7, 10)], 1):
    for c, v2 in enumerate(row, 1):
        ws.cell(r, c, v2)
ws["E2"] = "=AVERAGE(B2:D2)"   # Q1: celula arata rezultatul, bara de formule arata formula
wb.create_sheet("Romana")        # Q2: un registru (fisier) cu mai multe foi
wb.create_sheet("Note finale")
ws["G1"] = "coloana AB ="
ws["H1"] = "=COLUMN(AB12)"       # Q3: AB12 = coloana 28, randul 12
ws["G2"] = "rand AB12 ="
ws["H2"] = "=ROW(AB12)"
out["ex3"] = {
    "Q1_formula_E2": ws["E2"].value,
    "Q1_media_python": round((9 + 7 + 10) / 3, 2),
    "Q2_foi_in_registru": wb.sheetnames,
    "Q3_AB_index_python": column_index_from_string("AB"),
}
wb.save(X / "8A_Popescu_Ex3_exemple.xlsx")

# ---------- Blocul „Copiaza" nou: modelez lipirea (Excel imparte pe Tab si pe rand nou) ----------
bloc = "Nume\tNota1\tNota2\tNota3\nIon\t9\t7\t10\nAna\t8\t10\t9\nMihai\t7\t6\t8\nElena\t10\t9\t10\nAlex\t6\t8\t7"
wb = openpyxl.Workbook()
ws = wb.active
for r, linie in enumerate(bloc.split("\n"), 1):
    for c, cel in enumerate(linie.split("\t"), 1):
        ws.cell(r, c, int(cel) if cel.isdigit() else cel)
ws["F1"] = "=COUNT(B2:D6)"
ws["F2"] = "=AVERAGE(B2:D2)"
out["lipire"] = {"zona": ws.calculate_dimension(), "B2": ws["B2"].value, "A6": ws["A6"].value, "D6": ws["D6"].value}
wb.save(X / "lipire_bloc_nou.xlsx")
(S / "bloc_copiaza.txt").write_text(bloc, encoding="utf-8")

(S / "construieste_iesire.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps(out, ensure_ascii=False, indent=1))
