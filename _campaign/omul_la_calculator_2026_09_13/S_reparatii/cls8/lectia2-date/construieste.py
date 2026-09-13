"""Reconstruieste modelele reparate (factura, bonus, lipire cu Tab) si le recalculeaza prin H_randeaza.py."""
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook

S = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls8\lectia2-date")
H = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\H_randeaza.py")

wb = Workbook()

# 1) Factura: numere in celule, unitatea in antet
ws = wb.active
ws.title = "Factura"
ws.merge_cells("A1:E1")
ws["A1"] = "FACTURA NR. 001"
ws.append([])
ws.append(["Nr.", "Produs", "Cantitate (buc)", "Pret unitar (lei)", "Total (lei)"])
produse = [("Caiet A4", 5, 3.5), ("Pix", 10, 1.2), ("Radiera", 8, 1.0), ("Creion", 12, 0.8), ("Rigla", 6, 2.5)]
for i, (p, c, pr) in enumerate(produse, start=1):
    r = 3 + i
    ws.append([i, p, c, pr, f"=C{r}*D{r}"])
ws["D9"] = "TOTAL GENERAL"
ws["E9"] = "=SUM(E4:E8)"
for r in range(4, 10):
    ws[f"D{r}"].number_format = '#,##0.00 "lei"'
    ws[f"E{r}"].number_format = '#,##0.00 "lei"'
ws["D9"].number_format = "General"
ws["G4"] = "=COUNT(C4:D8)"   # 10 = toate cantitatile si preturile sunt numere

# 2) Bonus: inserezi rand deasupra, apoi imbini A1:D1
wb2 = wb.create_sheet("Bonus")
wb2.append(["Nr. Crt.", "Nume Elev", "Nota", "Media"])
for i, (n, x, m) in enumerate([("Elev A", 9, 8.5), ("Elev B", 8, 9.25), ("Elev C", 10, 9.5), ("Elev D", 7, 7.75), ("Elev E", 9, 8)], 1):
    wb2.append([i, n, x, m])
wb2.insert_rows(1)
wb2.merge_cells("A1:D1")
wb2["A1"] = "CATALOG NOTE - CLASA a VIII-a"
wb2["F1"] = "=COUNTA(A2:D2)"  # 4 = antetul a ramas intreg in randul 2

# 3) Lipire: blocul din lectie, impartit pe Tab cum face Excel la lipire
bloc = Path(S / "bloc_ex1.txt").read_text(encoding="utf-8")
wl = wb.create_sheet("Lipire")
for linie in bloc.splitlines():
    celule = []
    for v in linie.split("\t"):
        try:
            celule.append(int(v))
        except ValueError:
            celule.append(v)
    wl.append(celule)
wl["G1"] = "=COUNT(C2:D5)"  # 8 = toate notele au intrat ca numere
wl["G2"] = "=COUNTA(A1:D5)"  # 20 = 5 randuri x 4 coloane

out = S / "model_reparat.xlsx"
wb.save(out)
rc = subprocess.run([sys.executable, str(H), "xlsx", str(out), str(S)]).returncode
print("H_randeaza exit", rc)

from openpyxl import load_workbook
rw = load_workbook(S / "recalculat" / "model_reparat.xlsx", data_only=True)
f = rw["Factura"]
print("Factura E4..E8:", [f[f"E{r}"].value for r in range(4, 9)], "E9:", f["E9"].value, "COUNT numere:", f["G4"].value)
b = rw["Bonus"]
print("Bonus A1:", b["A1"].value, "| rand 2:", [b.cell(2, c).value for c in range(1, 5)], "| COUNTA antet:", b["F1"].value)
l = rw["Lipire"]
print("Lipire COUNT note:", l["G1"].value, "| COUNTA:", l["G2"].value, "| B2:", l["B2"].value)
