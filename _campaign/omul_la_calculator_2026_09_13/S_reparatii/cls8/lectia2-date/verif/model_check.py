import re
import openpyxl

S = r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls8\lectia2-date"
L = r"C:\00\Projects\LearningHub\content\tic\cls8\m1-excel-fundamente\lectia2-date.html"

# 1) factura: independent din textul rezolvarii de pe pagina
html = open(L, encoding="utf-8").read()
m = re.search(r"Randurile 4-8 \(in celule doar numere; unitatile sunt in antet\): (.*?)</li>", html)
rows = [r.strip() for r in m.group(1).rstrip(".").split(";")]
tot = 0
for r in rows:
    nr, prod, cant, pret, total = [x.strip() for x in r.split(", ")]
    c, p, t = int(cant), float(pret.replace(",", ".")), float(total.replace(",", "."))
    ok = abs(c * p - t) < 1e-9
    tot += t
    print(nr, prod, c, p, t, "OK" if ok else "GRESIT")
print("suma", round(tot, 2))

# 2) blocul copiabil din Ex.1 split pe Tab (cum il lipeste Excel)
b = re.search(r'<div class="copyable-code">\s*<button class="copy-btn">Copiaza</button>\s*<div class="code-block">(.*?)</div>', html, re.S).group(1)
b = b.replace("&#9;", "\t")
grid = [line.split("\t") for line in b.strip().split("\n")]
print("grid", len(grid), "x", {len(g) for g in grid}, grid[1])
nums = sum(1 for g in grid[1:] for v in g[2:] if v.strip().isdigit())
print("note numerice C2:D5 =", nums)

# 3) bonus: antet in randul 1, inserare rand, merge A1:D1
wb = openpyxl.Workbook()
ws = wb.active
ws.append(["Nr. Crt.", "Nume Elev", "Nota", "Media"])
ws.append([1, "X", 8, 8.5])
ws.insert_rows(1)
ws.merge_cells("A1:D1")
ws["A1"] = "CATALOG NOTE - CLASA a VIII-a"
print("rand2", [c.value for c in ws[2]])

# 4) modelul reparatorului
wm = openpyxl.load_workbook(S + r"\recalculat\model_reparat.xlsx", data_only=True)
for sh in wm.sheetnames:
    print("foaie", sh, [[c.value for c in r] for r in wm[sh].iter_rows(max_row=9)][:9])
