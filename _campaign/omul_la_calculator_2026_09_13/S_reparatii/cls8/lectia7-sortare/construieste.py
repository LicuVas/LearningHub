"""Construieste tabelele-model ale lectiei 7 (sortare) din datele lectiei si le scrie in xlsx.
Sortarea Excel e stabila; aici folosim sorted() (stabil) cu aceleasi chei.
"""
import json
from pathlib import Path
import openpyxl

S = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls8\lectia7-sortare")
out = {}

# Datele lectiei (Incearca pas 1 + Ex.1 + Ex.2)
elevi = [("Maria", 8), ("Andrei", 6), ("Elena", 9), ("Bogdan", 5), ("Carla", 7), ("Dan", 10)]
clase = ["8A", "8B", "8A", "8B", "8A", "8B"]

# Ex.2 corect: Ex.1 se termina cu Ctrl+Z de doua ori -> ordinea initiala; se adauga Clasa, apoi se INSEREAZA coloana Nr
tabel = [(i + 1, n, v, clase[i]) for i, (n, v) in enumerate(elevi)]
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Ex2_initial"
ws.append(["Nr", "Elev", "Nota", "Clasa"])
for r in tabel:
    ws.append(list(r))
sortat = sorted(sorted(tabel, key=lambda r: -r[2]), key=lambda r: r[3])  # nivel1 Clasa A-Z, nivel2 Nota desc
ws2 = wb.create_sheet("Ex2_sortat_Clasa_Nota")
ws2.append(["Nr", "Elev", "Nota", "Clasa"])
for r in sortat:
    ws2.append(list(r))
ws2["F1"] = "Suma note"; ws2["F2"] = "=SUM(C2:C7)"
out["ex2_sortat"] = sortat
out["ex2_revenire_Nr"] = sorted(sortat, key=lambda r: r[0])

# Incearca: dupa pasul 3 (Z-A pe Nota) tabelul e in ordinea notelor; BONUS scris direct peste acea ordine
dupa_pas3 = sorted(elevi, key=lambda r: -r[1])
bonus_fara_undo = [(n, v, clase[i]) for i, (n, v) in enumerate(dupa_pas3)]
out["bonus_fara_undo_asociere"] = bonus_fara_undo
out["bonus_fara_undo_sortat"] = sorted(sorted(bonus_fara_undo, key=lambda r: -r[1]), key=lambda r: r[2])
bonus_cu_undo = [(n, v, clase[i]) for i, (n, v) in enumerate(elevi)]
out["bonus_cu_undo_sortat"] = sorted(sorted(bonus_cu_undo, key=lambda r: -r[1]), key=lambda r: r[2])

# Capcana: doar coloana Nota sortata (Continue with the current selection)
note_desc = sorted([v for _, v in elevi], reverse=True)
capcana = [(n, note_desc[i]) for i, (n, _) in enumerate(elevi)]
ws3 = wb.create_sheet("Capcana_doar_B")
ws3.append(["Elev", "Nota"])
for r in capcana:
    ws3.append(list(r))
ws3["D1"] = "Suma"; ws3["D2"] = "=SUM(B2:B7)"
out["capcana"] = capcana
out["capcana_suma"] = sum(v for _, v in capcana)
out["capcana_gresiti"] = sum(1 for (n, v), (n0, v0) in zip(capcana, elevi) if v != v0)

# Capcana si in sens CRESCATOR (doar coloana B): ce rand de control prinde greseala in AMBELE sensuri?
note_asc = sorted([v for _, v in elevi])
capcana_asc = [(n, note_asc[i]) for i, (n, _) in enumerate(elevi)]
ws4 = wb.create_sheet("Capcana_doar_B_crescator")
ws4.append(["Elev", "Nota"])
for r in capcana_asc:
    ws4.append(list(r))
ws4["D1"] = "Suma"; ws4["D2"] = "=SUM(B2:B7)"
out["capcana_asc"] = capcana_asc
out["neschimbati_desc"] = [n for (n, v), (_, v0) in zip(capcana, elevi) if v == v0]
out["neschimbati_asc"] = [n for (n, v), (_, v0) in zip(capcana_asc, elevi) if v == v0]
out["control_bun_ambele_sensuri"] = [n for n, _ in elevi if n not in out["neschimbati_desc"] + out["neschimbati_asc"]]

f = S / "model_sortare.xlsx"
wb.save(f)
print(json.dumps(out, ensure_ascii=False, indent=1))
