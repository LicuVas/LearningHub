"""Construieste graficele reparatiei lectiei 5 (cls8) si recalculeaza cifrele scrise in lectie."""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference

S = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls8\lectia5-grafice")
date = [("Materie", "Nota medie"), ("Romana", 7.2), ("Matematica", 6.8), ("Engleza", 8.1), ("Istorie", 7.5), ("Fizica", 6.3)]
out = {}


def foaie(wb, nume):
    ws = wb.create_sheet(nume)
    for r in date:
        ws.append(r)
    return ws


wb = Workbook()
wb.remove(wb.active)
# Incearca, varianta noua: Column -> Bar -> Line (fara Pie)
for nume, ch in (("Column", BarChart()), ("Bar", BarChart()), ("Line", LineChart())):
    ws = foaie(wb, nume)
    if nume == "Column":
        ch.type = "col"
    if nume == "Bar":
        ch.type = "bar"
    ch.title = "Notele clasei a VIII-a (" + nume + ")"
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=6), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=6))
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 10
    ws.add_chart(ch, "D2")

# Axa pornita de la 6 vs de la 0
ws = foaie(wb, "Axa_de_la_6")
ch = BarChart()
ch.type = "col"
ch.title = "Aceleasi note, axa pornita de la 6"
ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=6), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=6))
ch.y_axis.scaling.min = 6
ch.y_axis.scaling.max = 8.5
ws.add_chart(ch, "D2")

# Pie pe medii: procentele fara sens (de ce l-am scos)
ws = foaie(wb, "Pie_pe_medii")
ch = PieChart()
ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=6), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=6))
ws.add_chart(ch, "D2")
wb.save(S / "s_incearca_nou.xlsx")

note = dict(date[1:])
tot = sum(note.values())
out["pie_procente_medii"] = {k: round(100 * v / tot, 1) for k, v in note.items()}
f, e = note["Fizica"], note["Engleza"]
out["fizica_sub_engleza_pct"] = round(100 * (e - f) / e, 1)
out["axa_de_la_6_raport_bare"] = round((e - 6) / (f - 6), 2)
out["axa_de_la_0_raport_bare"] = round(e / f, 2)
# schema HTML: inaltime bara = nota * 9 px (axa 0-10 = 90 px)
out["schema_px"] = {k: round(v * 9) for k, v in note.items()}
(S / "s_cifre.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps(out, ensure_ascii=False, indent=1))
