"""Reconstruieste cifrele reparate din lectia3-formule (cls8) intr-un xlsx; se recalculeaza cu H_randeaza.py xlsx."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.formula.translate import Translator

S = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls8\lectia3-formule")
wb = Workbook()

# Misiunea (try section)
ws = wb.active
ws.title = "misiune"
for r in [["Elev", "Nota1", "Nota2", "Nota3", "Total", "Media"], ["Ana", 9, 8, 10], ["Ion", 7, 6, 8], ["Maria", 10, 10, 9]]:
    ws.append(r)
for i in (2, 3, 4):
    ws[f"E{i}"] = f"=B{i}+C{i}+D{i}"
    ws[f"F{i}"] = f"=E{i}/3"
ws["E5"] = "=SUM(E2:E4)"

# Ex. 1
ws = wb.create_sheet("ex1")
ws.append(["Elev", "Nota1", "Nota2", "Nota3", "Nota4", "Suma", "Media"])
for r in [["Ana", 8, 7, 9, 6], ["Ion", 10, 9, 8, 10], ["Maria", 6, 7, 5, 8], ["Radu", 9, 9, 10, 9], ["Elena", 7, 8, 6, 7]]:
    ws.append(r)
for i in range(2, 7):
    ws[f"F{i}"] = f"=SUM(B{i}:E{i})"
    ws[f"G{i}"] = f"=F{i}/4"

# Ex. 2 cu $ si fara $ (varianta gresita, trasa in jos din C3)
for name, f0 in (("ex2_cu_dolar", "=B3*$B$1"), ("ex2_fara_dolar", "=B3*B1")):
    ws = wb.create_sheet(name)
    for r in [["Nr elevi", 28], ["Cheltuiala", "Pret/elev", "Cost total"], ["Transport", 45], ["Cazare", 120], ["Mancare", 35], ["Bilete muzeu", 15], ["Suveniruri", 20]]:
        ws.append(r)
    for i in range(3, 8):
        ws[f"C{i}"] = Translator(f0, origin="C3").translate_formula(f"C{i}")
    ws["C8"] = "=SUM(C3:C7)"

# TVA 21% din atomul 7, cu $ si fara $
for name, f0 in (("tva_cu_dolar", "=B2*$B$1"), ("tva_fara_dolar", "=B2*B1")):
    ws = wb.create_sheet(name)
    for r in [["TVA", 0.21, "Formula TVA"], ["Caiet", 5], ["Pix", 3], ["Rigla", 8]]:
        ws.append(r)
    for i in range(2, 5):
        ws[f"C{i}"] = Translator(f0, origin="C2").translate_formula(f"C{i}")

# Ordinea operatiilor + intrebarea noua
ws = wb.create_sheet("ordine")
for k, f in enumerate(["=(2+3)*4", "=2+3*4", "=10-3+2", "=10-(3+2)", "=20-8/2", "=(20-8)/2", "=40-10/2", "=(40-10)/2", "=9+7+8/3", "=(9+7+8)/3", "=12/3*2"], start=1):
    ws[f"A{k}"] = "'" + f
    ws[f"B{k}"] = f

wb.save(S / "verificare.xlsx")

# Intrebarea mutata la atomul 4: formula copiata din E2 in E5
print("E2->E5:", Translator("=B2+C2+D2", origin="E2").translate_formula("E5"))
print("salvat", S / "verificare.xlsx")

if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1 and sys.argv[1] == "citeste":
        from openpyxl import load_workbook
        rb = load_workbook(S / "recalculat" / "verificare.xlsx", data_only=True)
        for w in rb:
            print("==", w.title)
            for row in w.iter_rows(values_only=True):
                print(row)
