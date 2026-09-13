# Verificator: execut Ex.2 exact cum e scris ACUM, independent de ex2_executa.py al reparatorului
import copy, sys
from pptx import Presentation
from pptx.util import Inches
sys.stdout.reconfigure(encoding="utf-8")
OUT = r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls6\lectia1-powerpoint-intro\verif\v_ex2.pptx"
NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

def citeste(prs):
    return [s.shapes.title.text for s in prs.slides]

prs = Presentation()
lay = prs.slide_layouts[1]
# 1-2: 8 slide-uri, numerele 1..8
for n in range(1, 9):
    s = prs.slides.add_slide(lay); s.shapes.title.text = str(n)
lst = prs.slides._sldIdLst
print("pas 2:", citeste(prs))
# 4: ordine inversa (drag-and-drop = reordonare sldIdLst)
ids = list(lst)
for e in ids: lst.remove(e)
for e in reversed(ids): lst.append(e)
print("pas 4:", citeste(prs))
# 5: click dreapta pe POZITIA 5 -> Dublare diapozitiv; copia imediat dupa original
src = prs.slides[4]
print("pas 5: pe pozitia 5 scrie", src.shapes.title.text)
dup = prs.slides.add_slide(src.slide_layout)
for sh in list(dup.shapes): sh._element.getparent().remove(sh._element)
for sh in src.shapes: dup.shapes._spTree.append(copy.deepcopy(sh._element))
new_id = lst[-1]; lst.remove(new_id); lst.insert(5, new_id)
print("pas 5 dupa:", citeste(prs), len(prs.slides))
# 6: sterge DEODATA pozitiile 2,4,6,8 (selectia se face inainte de orice stergere)
sel = [lst[i - 1] for i in (2, 4, 6, 8)]
print("pas 6 selectate (ce scrie pe ele):", [prs.slides[i - 1].shapes.title.text for i in (2, 4, 6, 8)])
for e in sel:
    prs.part.drop_rel(e.get(NS)); lst.remove(e)
print("pas 6 dupa:", citeste(prs))
prs.save(OUT)
back = Presentation(OUT)
print("REDESCHIS DE PE DISC:", citeste(back), "numar:", len(back.slides))
print("PROMIS 8,6,4,3,1 ->", citeste(back) == ["8", "6", "4", "3", "1"])
