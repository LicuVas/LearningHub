"""Executa Exercitiul 2 (lectia1-powerpoint-intro) in python-pptx pe cele trei citiri posibile.

Construieste un .pptx real cu 8 diapozitive numerotate, face operatiile pe lista de diapozitive (sldIdLst)
exact cum le face PowerPoint (dublarea se insereaza imediat dupa original), salveaza, REDESCHIDE de pe disc
si citeste numerele de pe diapozitive.
"""
import copy
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

OUT = Path(__file__).parent


def nou(n=8):
    p = Presentation()
    for k in range(1, n + 1):
        s = p.slides.add_slide(p.slide_layouts[5])  # Title Only
        s.shapes.title.text = str(k)
    return p


def ids(p):
    return p.slides._sldIdLst


def inverseaza(p):
    lst = ids(p)
    elems = list(lst)
    for e in elems:
        lst.remove(e)
    for e in reversed(elems):
        lst.append(e)


def dubleaza(p, poz):  # poz 1-based; copia vine imediat dupa original
    src = p.slides[poz - 1]
    s = p.slides.add_slide(src.slide_layout)
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    for sh in src.shapes:
        s.shapes._spTree.append(copy.deepcopy(sh._element))
    lst = ids(p)
    nou_el = lst[-1]
    lst.remove(nou_el)
    lst.insert(poz, nou_el)


def sterge_pozitii(p, pozitii):  # toate deodata (selectie cu Ctrl)
    lst = ids(p)
    elems = list(lst)
    for k in sorted(pozitii, reverse=True):
        if k <= len(elems):
            lst.remove(elems[k - 1])


def numere(p):
    return [s.shapes.title.text for s in p.slides]


def ruleaza(nume, pas_dublare, pas_stergere):
    p = nou()
    inverseaza(p)
    pas_dublare(p)
    jurnal = [f"dupa dublare: {numere(p)}"]
    pas_stergere(p, jurnal)
    f = OUT / f"ex2_{nume}.pptx"
    p.save(f)
    r = Presentation(f)
    print(f"{nume}: {' | '.join(jurnal)} -> REDESCHIS: {numere(r)} ({len(r.slides)} diapozitive)")


def dup_poz5(p):
    dubleaza(p, 5)


def dup_nr5(p):
    dubleaza(p, numere(p).index("5") + 1)


def st_poz_deodata(p, j):
    sterge_pozitii(p, [2, 4, 6, 8])


def st_nr_pare(p, j):
    poz = [i + 1 for i, t in enumerate(numere(p)) if int(t) % 2 == 0]
    sterge_pozitii(p, poz)


def st_pe_rand(p, j):
    for k in [2, 4, 6, 8]:
        if k > len(p.slides):
            j.append(f"pozitia {k} nu mai exista ({len(p.slides)} diapozitive)")
            break
        sterge_pozitii(p, [k])
        j.append(f"sters poz {k}: {numere(p)}")


p0 = nou()
inverseaza(p0)
print("dupa inversare:", numere(p0), "-> pe pozitia 5 scrie", numere(p0)[4])
ruleaza("A_pozitii_deodata", dup_poz5, st_poz_deodata)
ruleaza("B_numere_pare", dup_poz5, st_nr_pare)
ruleaza("C_pozitii_pe_rand", dup_poz5, st_pe_rand)
ruleaza("D_dublez_nr5_sterg_nr_pare", dup_nr5, st_nr_pare)
ruleaza("E_dublez_nr5_sterg_pozitii", dup_nr5, st_poz_deodata)
