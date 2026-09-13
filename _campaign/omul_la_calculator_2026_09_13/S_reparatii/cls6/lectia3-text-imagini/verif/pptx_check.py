from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE

F = r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls6\lectia3-text-imagini\produs\6A_Popescu_DespreMine.pptx"
p = Presentation(F)
for i, s in enumerate(p.slides, 1):
    pics = sum(1 for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    rows, maxw, minf, texts = 0, 0, 999, []
    for sh in s.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs).strip()
                if not t:
                    continue
                rows += 1
                maxw = max(maxw, len(t.split()))
                texts.append(t)
                for r in para.runs:
                    if r.font.size:
                        minf = min(minf, r.font.size.pt)
    fill = s.background.fill.type if s.follow_master_background is False else "master"
    print(i, "imagini=%d randuri=%d max_cuvinte=%d font_min=%s fundal=%s" % (pics, rows, maxw, minf, fill), texts)
