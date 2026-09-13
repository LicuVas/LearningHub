"""Reconstruieste Ex.1 'Despre mine' exact dupa rezolvarea-model NOUA, il citeste inapoi de pe disc si verifica
regulile obligatorii ale exercitiului: minim 1 imagine pe slide, max 6 randuri, max 6 cuvinte pe rand, font >= 18pt."""
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Cm, Pt

S = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\S_reparatii\cls6\lectia3-text-imagini")
IMG = Path(r"C:\00\Projects\LearningHub\_campaign\omul_la_calculator_2026_09_13\F_evaluari\cls6\lectia3-text-imagini\produs_elev\imagini")
OUT = S / "produs" / "6A_Popescu_DespreMine.pptx"
OUT.parent.mkdir(parents=True, exist_ok=True)

prs = Presentation()
prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
blank = prs.slide_layouts[6]


def fundal(slide, rgb, rgb2=None):
    f = slide.background.fill
    if rgb2:
        f.gradient()
        f.gradient_stops[0].color.rgb = RGBColor(*rgb)
        f.gradient_stops[1].color.rgb = RGBColor(*rgb2)
    else:
        f.solid()
        f.fore_color.rgb = RGBColor(*rgb)


def text(slide, x, y, w, h, randuri, size, bold=False, culori=None):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, r in enumerate(randuri):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = r
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(*(culori[i] if culori else (20, 30, 60)))
    return tb


# Slide 1 - Coperta
s = prs.slides.add_slide(blank)
fundal(s, (235, 242, 255))
text(s, 2, 2, 20, 3, ["Despre mine - Andrei"], 40, bold=True, culori=[(20, 60, 160)])
s.shapes.add_picture(str(IMG / "coperta.png"), Cm(22), Cm(5), height=Cm(9))
# Slide 2 - Despre mine: lista + O imagine mare pentru hobby-ul preferat, in dreapta listei
s = prs.slides.add_slide(blank)
fundal(s, (255, 255, 255))
text(s, 2, 1, 20, 2.5, ["Hobby-urile mele"], 36, bold=True)
text(s, 2, 4, 15, 10, ["• Fotbal cu prietenii", "• Desen in creion", "• Carti de aventuri", "• Lego"], 28)
s.shapes.add_picture(str(IMG / "hobby0.png"), Cm(18), Cm(4), height=Cm(11))
# Slide 3 - O familie imaginara: prenume inventate, imagine fara persoane reale
s = prs.slides.add_slide(blank)
fundal(s, (255, 255, 255))
text(s, 2, 1, 26, 2.5, ["O familie imaginara"], 36, bold=True)
text(s, 2, 4, 15, 8, ["Mama Ana", "Tata Radu", "Bunicul Ion"], 28, bold=True,
     culori=[(180, 30, 30), (20, 90, 40), (30, 60, 160)])
s.shapes.add_picture(str(IMG / "familie.png"), Cm(18), Cm(4), height=Cm(10))
# Slide 4 - Obiective: lista numerotata + imagine
s = prs.slides.add_slide(blank)
fundal(s, (255, 255, 255))
text(s, 2, 1, 26, 2.5, ["Obiectivele mele"], 36, bold=True)
text(s, 2, 4, 15, 10, ["1. Citesc 10 carti", "2. Invat sa inot", "3. Iau nota 10 la TIC"], 28,
     culori=[(20, 60, 160), (20, 90, 40), (150, 40, 120)])
s.shapes.add_picture(str(IMG / "obiective.png"), Cm(18), Cm(4), height=Cm(10))
# Slide 5 - Contact fictiv: forme + imagine mica + sursa imaginilor, fundal gradient
s = prs.slides.add_slide(blank)
fundal(s, (220, 235, 255), (255, 255, 255))
text(s, 2, 1, 26, 2.5, ["Contact (inventat)"], 36, bold=True)
for i, t in enumerate(["elev_test@scoala.ro", "0700-000-000"]):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(2), Cm(4.5 + i * 4), Cm(14), Cm(3))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(40, 90, 180)
    sh.text_frame.text = t
    sh.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
s.shapes.add_picture(str(IMG / "contact.png"), Cm(20), Cm(4.5), height=Cm(6))
text(s, 2, 16, 20, 1.5, ["Sursa imaginilor: Pexels"], 18)
prs.save(OUT)

# ---- citit inapoi de pe disc ----
p2 = Presentation(OUT)
rez = []
for n, sl in enumerate(p2.slides, 1):
    imagini = sum(1 for sh in sl.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    randuri, max_cuv, min_font = 0, 0, 999
    for sh in sl.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                t = "".join(r.text for r in p.runs).strip()
                if not t:
                    continue
                randuri += 1
                max_cuv = max(max_cuv, len(t.replace("•", "").split()))
                for r in p.runs:
                    if r.font.size:
                        min_font = min(min_font, r.font.size.pt)
    rez.append({"slide": n, "imagini": imagini, "randuri_text": randuri, "max_cuvinte_rand": max_cuv, "font_minim_pt": min_font})
ok = all(r["imagini"] >= 1 and r["randuri_text"] <= 6 and r["max_cuvinte_rand"] <= 6 and r["font_minim_pt"] >= 18 for r in rez)
iesire = {"fisier": str(OUT), "slide_uri": rez, "slide_uri_fara_imagine": [r["slide"] for r in rez if r["imagini"] < 1], "reguli_respectate": ok}
(S / "construieste_ex1_iesire.json").write_text(json.dumps(iesire, ensure_ascii=False, indent=1), encoding="utf-8")
print(json.dumps(iesire, ensure_ascii=False, indent=1))
