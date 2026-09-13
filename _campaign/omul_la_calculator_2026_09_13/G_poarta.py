"""Poarta mecanica pentru evaluarea „omul la calculator". Versiunea 2 - 13.09.2026.

v1 trecea evaluari fara munca (dovezi fara legatura, chiar C:\\Windows\\explorer.exe; exit 0 pe cale gresita).
v2 inchide gaurile gasite de verificatorul independent (E_verificare_adversariala.md: B1-B6, I1-I3, I11, I13, I15, I16).

Folosire:
  python G_poarta.py F_evaluari/cls7/lectia2-formatare-text
  python G_poarta.py F_evaluari --astept 25
Exit: 0 = toate trec · 1 = cel putin una pica · 2 = cale gresita / nimic de verificat.
"""
import argparse
import csv
import html
import json
import re
import sys
import zipfile
from pathlib import Path

CAMP = Path(__file__).resolve().parent
SITE = Path(r"C:\00\Projects\LearningHub")

CHECKS = {
    "U1": "Fac eu sarcina cap-coada",
    "U2": "Ma uit la ce vede elevul (randat, pe pasi)",
    "U3": "Verific pe a doua cale",
    "U4": "Ce mediu presupune lectia",
    "U5": "Reproduc defectul inainte sa-l afirm",
    "U7": "Starea finala reala (redeschid artefactul)",
    "U8": "Citesc inapoi ce promite lectia",
    "U9": "Checklistul de pe hartie (spec + R1-R6)",
    "U10": "Trei identificatori (fisier / titlu / ora din plan)",
    "U11": "Incepatorul real, cu constrangeri",
    "U12": "Cronometrez (calculat de poarta)",
    "U13": "Ce altceva ar mai putea fi",
    "U14": "Pre-mortem",
    "U15": "Citesc tot",
    "U16": "Norma din afara documentului",
    "U17": "Preconditia orei in sala reala",
}
# U6 („un lucru o data") e o regula de REPARATIE, nu de evaluare -> scos (I13).
STATUS = {"FACUT", "NEFACUT", "NESIGUR"}
NESIGUR_PERMIS = {"U7", "U11", "U17"}        # singurele care pot depinde de sala / de copii reali (I1)
GRAV = {"blocant", "important", "minor"}
H_VEDE = {"ecran_prima_vedere.png", "proiector_25.png", "innerText.txt", "innerText_vizibil.txt",
          "consola.json", "masuri.json", "parcurgere.json"}
DOSARE_OBLIGATORII = {"01_citit_inapoi.md": 300, "03_jurnal_sarcina.md": 300, "05_evaluare.md": 300, "09_checklist.md": 300}

ANEXA_COMUNA = ["fara_cont_email_stick", "puncte_de_verificare_vizibile", "sarcina_de_rezerva_retea_proiector",
                "tastarea_diacriticelor_pe_pc_elev", "blocaje_probabile_ipoteza", "unde_se_salveaza_fisierul"]
ANEXA_CLASA = {
    "cls5": ["ergonomie_contra_mobilierului_real", "afirmatii_hardware_actuale_2026"],
    "cls6": ["modul_prezentare_font_la_distanta", "o_idee_pe_diapozitiv", "denumiri_animatii_tranzitii_ro"],
    "cls7": ["stil_vs_formatare_manuala_ce_cere_lectia", "marcaje_paragraf_spatii_enter", "ghilimele_romanesti"],
    "cls8": ["numar_scris_ca_text_si_date", "copierea_formulei_in_jos", "lipirea_tabelului_starter", "separator_si_nume_functii"],
}
# Ritmuri PROVIZORII de incepator (de calibrat la clasa cu Vasile) - I11: agentul nu-si alege singur factorul.
RITM = {  # secunde/clic, caractere/minut tastate, cuvinte/minut citite
    "cls5": (20, 40, 90), "cls6": (20, 50, 100), "cls7": (15, 60, 115), "cls8": (15, 70, 125),
}
PORNIRE_MIN = 8


def inside(d: Path, p: str) -> bool:
    try:
        q = (d / p).resolve()
        return q.is_relative_to(d.resolve()) and q.is_file() and q.stat().st_size > 0
    except (OSError, ValueError):
        return False


def norm(t: str) -> str:
    t = re.sub(r"\[/?ASCUNS[^\]]*\]", " ", t)
    t = html.unescape(t).replace("\u00a0", " ")
    t = re.sub(r"[„”“\"'’‘«»]", "", t)
    return re.sub(r"\s+", " ", t).strip().lower()


def diacritice_la_1000(t: str) -> float:
    t = re.sub(r"\[/?ASCUNS[^\]]*\]", " ", t)
    lit = len(re.findall(r"[A-Za-zăâîșțşţĂÂÎȘȚŞŢ]", t))
    dia = len(re.findall(r"[ăâîșțşţĂÂÎȘȚŞŢ]", t))
    return round(1000 * dia / max(lit, 1), 1)


def inventar():
    rows = {}
    f = CAMP / "B_inventar.csv"
    if f.is_file():
        for r in csv.DictReader(f.open(encoding="utf-8")):
            rows[Path(r["path"]).stem + "|" + r["grade"]] = r
    return rows


def check_u1_artefact(d: Path, cls: str, ev: list, inv_row) -> list:
    errs = []
    own = [p for p in ev if inside(d, p) and Path(p).name not in H_VEDE and not p.lower().endswith(".md")]
    if not own:
        return [f"U1 FACUT fara artefact propriu (fisierele lui H_vede si .md nu conteaza): {ev}"]
    try:
        if cls == "cls6":
            pp = [p for p in own if p.lower().endswith(".pptx")]
            if not pp:
                errs.append("U1 la cls6 cere un .pptx construit")
            else:
                from pptx import Presentation
                if len(Presentation(str(d / pp[0])).slides) < 2:
                    errs.append("U1 cls6: .pptx cu mai putin de 2 diapozitive")
        elif cls == "cls7":
            dd = [p for p in own if p.lower().endswith(".docx")]
            if not dd:
                errs.append("U1 la cls7 cere un .docx construit")
            else:
                import docx
                if len([x for x in docx.Document(str(d / dd[0])).paragraphs if x.text.strip()]) < 3:
                    errs.append("U1 cls7: .docx cu mai putin de 3 paragrafe cu text")
        elif cls == "cls8":
            xx = [p for p in own if p.lower().endswith(".xlsx")]
            if len(xx) < 2:
                errs.append("U1 la cls8 cere .xlsx construit + copia RECALCULATA de LibreOffice (2 fisiere)")
            else:
                import openpyxl
                formule = 0
                valori = 0
                for p in xx:
                    wb = openpyxl.load_workbook(str(d / p))
                    wbv = openpyxl.load_workbook(str(d / p), data_only=True)
                    for ws in wb.worksheets:
                        for row in ws.iter_rows():
                            for c in row:
                                if isinstance(c.value, str) and c.value.startswith("="):
                                    formule += 1
                                    v = wbv[ws.title][c.coordinate].value
                                    if v is not None:
                                        valori += 1
                if formule < 3:
                    errs.append(f"U1 cls8: doar {formule} formule in fisierele xlsx")
                if valori < 3:
                    errs.append("U1 cls8: nicio copie recalculata cu valori (data_only) - a rulat LibreOffice?")
        elif cls == "cls5":
            jj = [p for p in own if p.lower().endswith((".json", ".pptx", ".docx", ".png", ".pdf"))]
            if not jj:
                errs.append("U1 cls5: cere raspunsurile elevului la exercitii (json) sau produsul proiectului")
            ex = int((inv_row or {}).get("exercises") or 0)
            js = [p for p in own if p.lower().endswith(".json")]
            if ex and js:
                try:
                    data = json.loads((d / js[0]).read_text(encoding="utf-8"))
                    n = len(data) if isinstance(data, list) else len(data.get("exercitii", []))
                    if n < ex:
                        errs.append(f"U1 cls5: {n} exercitii facute din {ex} cate are lectia (B_inventar.csv)")
                except Exception as e:  # noqa: BLE001
                    errs.append(f"U1 cls5: {js[0]} nu se poate citi: {e}")
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as e:
        errs.append(f"U1: artefactul nu se deschide ({e}) - e un fisier fals?")
    return errs


def check_pasi(d: Path, cls: str, data: dict, masuri: dict) -> list:
    errs = []
    f = d / "03_pasi.json"
    if not f.is_file():
        return ["lipseste 03_pasi.json (pasii elevului, tip rezultat/interfata)"]
    try:
        pasi = json.loads(f.read_text(encoding="utf-8"))
    except Exception as e:  # noqa: BLE001
        return [f"03_pasi.json invalid: {e}"]
    if not isinstance(pasi, list) or len(pasi) < 3:
        return ["03_pasi.json: sub 3 pasi"]
    rez = 0
    s_clic, cpm, wpm = RITM.get(cls, (20, 50, 100))
    sec = 0.0
    for i, p in enumerate(pasi):
        t = p.get("tip")
        if t == "rezultat":
            rez += 1
            if len(str(p.get("observat") or "")) < 5:
                errs.append(f"03_pasi[{i}] rezultat fara valoarea observata")
        elif t == "interfata":
            if p.get("gasit") not in ("da", "nu", "neclar"):
                errs.append(f"03_pasi[{i}] interfata fara gasit da/nu/neclar")
            elif p["gasit"] in ("da", "nu"):
                # pilot: un URL „plauzibil" necitit trecea. Acum: citatul copiat de pe pagina, salvat in L.
                sf = str(p.get("sursa_fisier") or "")
                if not sf or not inside(d, sf) or len((d / sf).read_text(encoding="utf-8", errors="replace")) < 80:
                    errs.append(f"03_pasi[{i}] interfata gasit={p['gasit']} fara sursa_fisier (citatul copiat de pe pagina deschisa, >=80 caractere)")
                elif not re.search(r"https?://", (d / sf).read_text(encoding="utf-8", errors="replace")):
                    errs.append(f"03_pasi[{i}] sursa_fisier nu contine adresa paginii citate")
            elif len(str(p.get("nota") or "")) < 40:
                errs.append(f"03_pasi[{i}] interfata neclar fara nota (unde ai cautat, de ce nu s-a gasit)")
        else:
            errs.append(f"03_pasi[{i}] tip invalid {t!r}")
        try:
            sec += int(p.get("clicuri", 0)) * s_clic + int(p.get("caractere", 0)) * 60 / cpm + int(p.get("cuvinte_citite", 0)) * 60 / wpm
        except (TypeError, ValueError):
            errs.append(f"03_pasi[{i}] clicuri/caractere/cuvinte_citite nu sunt numere")
    if rez < 1:
        errs.append("03_pasi.json: niciun pas executat efectiv (tip rezultat)")
    citire = (masuri.get("cuvinte_total") or 0) / wpm
    total = PORNIRE_MIN + citire + sec / 60
    verdict = "incape" if total <= 50 else "nu incape"
    total_dublu = PORNIRE_MIN + (citire + sec / 60) / 2
    print(f"  U12 {d.name}: pornire {PORNIRE_MIN} + citirea lectiei {citire:.1f} + sarcina {sec/60:.1f} = {total:.1f} min -> {verdict}"
          f" | la ritm DUBLU: {total_dublu:.1f} min (ritmuri {cls} provizorii)")
    if total_dublu <= 50:
        for i, f in enumerate(data.get("findings") or []):
            if f.get("check") == "U12" and f.get("gravitate") == "blocant":
                errs.append(f"finding[{i}] blocant pe timp, dar la ritm dublu ora incape ({total_dublu:.0f} min) - ritmurile sunt provizorii -> maxim important")
    tm = data.get("timp_estimat_minute") or {}
    if tm.get("verdict") != verdict:
        errs.append(f"U12: verdictul din log ({tm.get('verdict')!r}) difera de cel calculat: {verdict} "
                    f"(pornire {PORNIRE_MIN} + citirea lectiei {citire:.0f} + sarcina {sec/60:.0f} = {total:.0f} min; ritm {cls} provizoriu)")
    return errs


def check_lesson(d: Path, inv: dict) -> list:
    errs = []
    cls = d.parent.name
    log = d / "log.json"
    if not log.is_file():
        return ["lipseste log.json"]
    try:
        data = json.loads(log.read_text(encoding="utf-8"))
    except Exception as e:  # noqa: BLE001
        return [f"log.json nu e JSON valid: {e}"]

    lesson = SITE / str(data.get("lesson") or "")
    if not data.get("lesson") or not lesson.is_file():
        return [f"campul lesson nu indica un fisier existent: {data.get('lesson')!r}"]
    if lesson.stem != d.name:
        errs.append(f"folderul {d.name} nu corespunde lectiei {lesson.stem}")
    src = lesson.read_text(encoding="utf-8", errors="replace")
    mt = re.search(r"<title>(.*?)</title>", src, re.S)
    title = html.unescape(mt.group(1).strip()) if mt else ""

    it = d / "innerText.txt"
    ms = d / "masuri.json"
    if not it.is_file() or not ms.is_file():
        return errs + ["lipseste iesirea lui H_vede.py (innerText.txt / masuri.json) - rulati H_vede pe lectie"]
    text = it.read_text(encoding="utf-8", errors="replace")
    masuri = json.loads(ms.read_text(encoding="utf-8"))
    if norm(masuri.get("titlu", "")) != norm(title):
        errs.append("masuri.json nu provine din lectia declarata (titlul difera)")
    h1 = norm(masuri.get("h1", ""))[:40]
    if not h1 or h1 not in norm(text):
        errs.append("innerText.txt nu contine h1-ul lectiei - text fals sau alta lectie")
    if not list(d.glob("pas_*.png")) and not list(d.glob("dupa_atomi_*.png")):
        errs.append("lipsesc capturile pe pasi (pas_NN.png / dupa_atomi_NN.png) - H_vede v2")

    for fn, mn in DOSARE_OBLIGATORII.items():
        if not (d / fn).is_file() or len((d / fn).read_text(encoding="utf-8", errors="replace")) < mn:
            errs.append(f"{fn} lipseste sau are sub {mn} caractere")

    checks = data.get("checks") or {}
    notes = []
    for cid in CHECKS:
        c = checks.get(cid)
        if not isinstance(c, dict):
            errs.append(f"{cid} lipseste")
            continue
        st = c.get("status")
        if st not in STATUS:
            errs.append(f"{cid} status invalid: {st!r}")
            continue
        note = (c.get("note") or "").strip()
        ev = [str(p) for p in (c.get("evidence") or [])]
        bad = [p for p in ev if not inside(d, p)]
        if bad:
            errs.append(f"{cid}: dovezi in afara folderului lectiei sau inexistente: {bad}")
        good = [p for p in ev if inside(d, p)]
        if st == "FACUT":
            if cid == "U15":
                if "innerText.txt" not in [Path(p).name for p in good]:
                    errs.append("U15 FACUT cere innerText.txt ca dovada")
            elif cid == "U2":
                if not any(Path(p).name.startswith(("pas_", "dupa_atomi_")) for p in good):
                    errs.append("U2 FACUT cere cel putin o captura de pas / de exercitii, nu doar primul ecran")
            elif cid == "U1":
                errs += check_u1_artefact(d, cls, good, inv.get(d.name + "|" + cls))
            else:
                own = [p for p in good if Path(p).name not in H_VEDE]
                if not own:
                    errs.append(f"{cid} FACUT: dovada trebuie sa fie un fisier produs de tine, nu iesirea lui H_vede")
        elif cid not in NESIGUR_PERMIS:
            errs.append(f"{cid} e {st}, dar se poate face la evaluare - trebuie FACUT cu dovada")
        else:
            if len(note) < 80 or "?" not in note:
                errs.append(f"{cid} {st}: nota trebuie sa numeasca blocajul si intrebarea pentru Vasile (>=80 caractere, cu „?”)")
            notes.append(note)
    if len(set(notes)) != len(notes):
        errs.append("note NESIGUR identice la mai multe verificari - scuza copiata")

    # U3: a doua cale, structurat
    adc = d / "04_a_doua_cale.json"
    try:
        rows = json.loads(adc.read_text(encoding="utf-8")) if adc.is_file() else []
        okr = [r for r in rows if r.get("cale1") and r.get("cale2") and r["cale1"] != r["cale2"]
               and str(r.get("val1", "")).strip() != "" and str(r.get("val2", "")).strip() != ""]
        if len(okr) < 2:
            errs.append("04_a_doua_cale.json: sub 2 verificari {ce, cale1, val1, cale2, val2} cu cai diferite")
    except Exception as e:  # noqa: BLE001
        errs.append(f"04_a_doua_cale.json invalid: {e}")

    # U16: diacriticele recalculate
    real = diacritice_la_1000(text)
    decl = data.get("diacritice_la_1000")
    if not isinstance(decl, (int, float)) or abs(decl - real) > 0.5:
        errs.append(f"diacritice_la_1000: declarat {decl!r}, masurat de poarta {real}")

    errs += check_pasi(d, cls, data, masuri)

    # anexa de verificari umane care cazusera din v1 (I16)
    anexa = data.get("anexa") or {}
    for k in ANEXA_COMUNA + ANEXA_CLASA.get(cls, []):
        a = anexa.get(k)
        if not isinstance(a, dict) or a.get("status") not in ("FACUT", "INTREBARE_VASILE", "NU_SE_APLICA"):
            errs.append(f"anexa.{k} lipseste (FACUT / INTREBARE_VASILE / NU_SE_APLICA)")
            continue
        if a["status"] == "FACUT" and not any(inside(d, p) and Path(p).name not in H_VEDE for p in (a.get("evidence") or [])):
            errs.append(f"anexa.{k} FACUT fara fisier-dovada propriu in folder (iesirea lui H_vede nu conteaza)")
        if a["status"] == "INTREBARE_VASILE" and "?" not in (a.get("intrebare") or ""):
            errs.append(f"anexa.{k}: intrebarea pentru Vasile lipseste")
        if a["status"] == "NU_SE_APLICA" and len(a.get("de_ce") or "") < 30:
            errs.append(f"anexa.{k}: NU_SE_APLICA fara motiv")

    findings = data.get("findings")
    if not isinstance(findings, list):
        errs.append("findings lipseste sau nu e lista")
        findings = []
    if not findings and len((data.get("no_findings_reason") or "").strip()) < 80:
        errs.append("zero semnalari si fara no_findings_reason (>=80 caractere)")
    ntext = norm(text)
    for i, f in enumerate(findings):
        tag = f"finding[{i}]"
        if f.get("check") not in CHECKS and f.get("check") not in ANEXA_COMUNA + sum(ANEXA_CLASA.values(), []):
            errs.append(f"{tag} check invalid {f.get('check')!r}")
        g = f.get("gravitate")
        if g not in GRAV:
            errs.append(f"{tag} gravitate invalida {g!r}")
        if f.get("depinde_de_necunoscut") not in (True, False):
            errs.append(f"{tag} fara depinde_de_necunoscut true/false")
        elif f["depinde_de_necunoscut"] and g == "blocant":
            errs.append(f"{tag} blocant pe un fapt necunoscut - maxim important, cu formulare pe ambele variante (I15)")
        dt = f.get("dovada_tip")
        dov = (f.get("dovada") or "").strip()
        if dt == "citat":
            q = norm(dov)
            if len(q) < 8 or q not in ntext:
                errs.append(f"{tag} citatul nu exista in textul lectiei (innerText.txt): {dov[:80]!r}")
        elif dt == "observatie":
            if not f.get("dovada_fisier"):
                errs.append(f"{tag} observatie fara dovada_fisier")
        else:
            errs.append(f"{tag} dovada_tip trebuie sa fie citat sau observatie")
        fp = f.get("dovada_fisier")
        if fp and not inside(d, fp):
            errs.append(f"{tag} dovada_fisier inexistent sau in afara folderului: {fp}")
        if g == "blocant" and not fp:
            errs.append(f"{tag} blocant fara dovada_fisier (reproducerea - U5)")
        if cls == "cls8" and re.search(r"#VALUE!|Err:5\d\d", dov + str(f.get("ce_vede_omul", ""))) and ";" in dov:
            errs.append(f"{tag} #VALUE! pe formula cu ';' e artefact de format de fisier xlsx, nu defect al lectiei (B4)")
        if len((f.get("ce_trebuie_diferit") or "").strip()) < 15:
            errs.append(f"{tag} fara ce_trebuie_diferit")
        if f.get("rol") not in ("PROFESORUL", "SPECIALISTUL", "ELEVUL", "NORMA"):
            errs.append(f"{tag} rol invalid")

    if len((data.get("carry_forward") or "").strip()) < 30:
        errs.append("carry_forward lipseste")
    jur = d.parent / "JURNAL.md"
    if not jur.is_file() or f"## {d.name}" not in jur.read_text(encoding="utf-8", errors="replace"):
        errs.append(f"JURNAL.md al clasei nu are sectiunea '## {d.name}'")
    return errs


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("root")
    ap.add_argument("--astept", type=int)
    a = ap.parse_args()
    root = Path(a.root).resolve()
    if not root.is_dir():
        print(f"cale inexistenta: {root}")
        return 2
    dirs = [root] if (root / "log.json").is_file() else sorted({p.parent for p in root.rglob("log.json")})
    if not dirs:
        print(f"nicio evaluare (log.json) gasita sub {root}")
        return 2
    inv = inventar()
    ok = 0
    total = 0
    for d in dirs:
        e = check_lesson(d, inv)
        total += len(e)
        ok += not e
        print(("OK   " if not e else "PICA ") + str(d))
        for x in e:
            print("  - " + x)
    rc = 0 if ok == len(dirs) else 1
    if a.astept is not None and len(dirs) != a.astept:
        print(f"astept {a.astept} lectii evaluate, gasite {len(dirs)}")
        rc = 1
    print(f"\n{ok}/{len(dirs)} lectii trec; {total} probleme")
    return rc


if __name__ == "__main__":
    sys.exit(main())
